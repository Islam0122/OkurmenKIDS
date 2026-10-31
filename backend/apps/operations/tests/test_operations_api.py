"""Кабинет операционного директора: доступы ролей, интеграция с академией
(журналы, ДЗ, KPI, проблемные студенты), учётом прихода, продажами по
отчётам, маркетингом и HR, KPI-цели, отчёты сотрудников, поиск проблем и
выгрузки — на реальных данных в тестовой БД."""
from __future__ import annotations

import datetime as dt
from io import BytesIO
from zoneinfo import ZoneInfo

from django.utils import timezone

from apps.academy.models import Attendance, Course, CourseLessonPlan, Group, GroupTeacher, Lesson, Student
from apps.marketing.models import Campaign
from apps.operations.models import (
    EmployeeReport,
    IssueKind,
    KPITarget,
    OperationalIssue,
    OperationsSettings,
    Task,
)
from apps.operations.services import issues as issue_service
from apps.operations.services.reports import obligations
from apps.staff_attendance.models import AttendanceSession, AttendanceSettings
from apps.users.models import Subject, User

from .base import OpsTestCase, jwt_client

API = "/api/v1/operations"
TZ = ZoneInfo("Asia/Bishkek")


class AccessTests(OpsTestCase):
    ENDPOINTS = ["meta/", "dashboard/", "departments/", "departments/academy/", "employees/", "trainers/",
                 "team-leads/", "assistants/", "students/problems/", "sales/", "sales/head/", "marketing/", "hr/",
                 "attendance/", "daily/", "kpi/", "reports/", "activity/", "settings/", "issues/", "kpi/targets/"]

    def test_oversight_roles_open_every_page(self):
        for user in (self.od, self.director, self.admin):
            client = jwt_client(user.username)
            for endpoint in self.ENDPOINTS:
                response = client.get(f"{API}/{endpoint}")
                self.assertEqual(response.status_code, 200, f"{user.role} {endpoint}: {response.data}")

    def test_other_roles_are_denied_on_the_server(self):
        for user in (self.lead, self.trainer, self.assistant, self.manager, self.rop, self.marketer, self.hr,
                     self.accountant):
            client = jwt_client(user.username)
            for endpoint in ("dashboard/", "trainers/", "kpi/targets/", "settings/"):
                self.assertEqual(client.get(f"{API}/{endpoint}").status_code, 403, f"{user.role} {endpoint}")

    def test_operations_director_still_cannot_reach_payroll(self):
        client = jwt_client(self.od.username)
        self.assertEqual(client.get("/api/v1/accounting/my/").status_code, 403)
        self.assertEqual(client.get("/api/v1/students/").status_code, 403)

    def test_operations_director_gets_staff_attendance_reports(self):
        self.assertEqual(jwt_client(self.od.username).get("/api/v1/staff-attendance/admin/overview/").status_code, 200)


class AcademyIntegrationTests(OpsTestCase):
    def setUp(self):
        super().setUp()
        python, _ = Subject.objects.get_or_create(name="Python")
        course = Course.objects.create(name="Prog", count_lesson=20)
        course.subjects.set([python])
        self.group = Group.objects.create(name="Python 1", course=course, start_date=dt.date(2025, 1, 1))
        self.gt = GroupTeacher.objects.create(group=self.group, teacher=self.trainer.teacher_profile, subject=python)
        self.students = [Student.objects.create(first_name=f"S{i}", group=self.group, enrollment_date=dt.date(2025, 1, 1))
                         for i in range(3)]
        self.today = timezone.localdate()
        self.n = 0

    def lesson(self, date, status=Lesson.Status.SCHEDULED):
        self.n += 1
        plan = CourseLessonPlan.objects.create(course=self.group.course, lesson_number=self.n, subject=self.gt.subject,
                                               topic=f"Т{self.n}", homework_title="")
        return Lesson.objects.create(group=self.group, group_teacher=self.gt, subject=self.gt.subject, plan=plan,
                                     lesson_number=self.n, date=date, status=status, teacher=self.trainer.teacher_profile,
                                     start_time=dt.time(10), end_time=dt.time(11))

    def test_unfilled_past_lesson_is_a_warning_but_future_lesson_is_not(self):
        self.lesson(self.today - dt.timedelta(days=2))
        self.lesson(self.today + dt.timedelta(days=3))
        data = self.client_for(self.od).get(f"{API}/trainers/", {"period": "custom",
                                                                   "start_date": (self.today - dt.timedelta(days=6)).isoformat(),
                                                                   "end_date": (self.today + dt.timedelta(days=6)).isoformat()}).data
        row = next(r for r in data["trainers"] if r["id"] == self.trainer.teacher_profile.id)
        keys = {w["key"] for w in row["warnings"]}
        self.assertIn("journal_not_filled", keys)
        self.assertIn("attendance_not_updated", keys)
        self.assertEqual(row["journal"]["lessons_due"], 1)  # будущее занятие не учитывается как долг
        self.assertEqual(row["journal"]["upcoming"], 1)

    def test_only_future_lessons_produce_no_violation(self):
        self.lesson(self.today + dt.timedelta(days=1))
        data = self.client_for(self.od).get(f"{API}/trainers/", {"period": "this_month"}).data
        row = next(r for r in data["trainers"] if r["id"] == self.trainer.teacher_profile.id)
        self.assertNotIn("journal_not_filled", {w["key"] for w in row["warnings"]})

    def test_filled_lesson_clears_warnings_and_trainer_profile_opens(self):
        lesson = self.lesson(self.today - dt.timedelta(days=1), Lesson.Status.COMPLETED)
        for s in self.students:
            Attendance.objects.create(lesson=lesson, student=s, status=Attendance.Status.PRESENT)
        data = self.client_for(self.od).get(f"{API}/trainers/", {"period": "this_week"}).data
        row = next(r for r in data["trainers"] if r["id"] == self.trainer.teacher_profile.id)
        self.assertEqual(row["warnings"], [])
        detail = self.client_for(self.od).get(f"{API}/trainers/{self.trainer.teacher_profile.id}/",
                                              {"period": "this_week"})
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.data["control"]["summary"]["completed_lessons"], 1)

    def test_journal_issue_detected_once_and_auto_resolved(self):
        lesson = self.lesson(self.today - dt.timedelta(days=1))
        issue_service.scan()
        issue_service.scan()
        issues = OperationalIssue.objects.filter(kind=IssueKind.JOURNAL_NOT_FILLED)
        self.assertEqual(issues.count(), 1)
        self.assertEqual(issues.get().responsible, self.trainer)
        lesson.status = Lesson.Status.COMPLETED
        lesson.save()
        for s in self.students:
            Attendance.objects.create(lesson=lesson, student=s, status=Attendance.Status.PRESENT)
        issue_service.scan()
        issue = issues.get()
        self.assertEqual(issue.status, "resolved")
        # Снова не заполнил — та же запись открыта заново: повторное нарушение.
        self.lesson(self.today)
        Lesson.objects.filter(date=self.today).update(start_time=dt.time(0, 0), end_time=dt.time(0, 1))
        issue_service.scan()
        issue.refresh_from_db()
        self.assertEqual(issue.status, "open")
        self.assertEqual(issue.reopen_count, 1)

    def test_problem_students_from_assistant_activity(self):
        for days in (3, 2, 1):
            lesson = self.lesson(self.today - dt.timedelta(days=days), Lesson.Status.COMPLETED)
            for s in self.students:
                status = Attendance.Status.ABSENT if s == self.students[0] else Attendance.Status.PRESENT
                Attendance.objects.create(lesson=lesson, student=s, status=status)
        data = self.client_for(self.od).get(f"{API}/students/problems/", {"period": "this_month"}).data
        ids = [r["student_id"] for r in data["students"]]
        self.assertIn(self.students[0].pk, ids)
        self.assertNotIn(self.students[1].pk, ids)


class StaffAttendanceIntegrationTests(OpsTestCase):
    def setUp(self):
        super().setUp()
        config = AttendanceSettings.load()
        config.tracking_started_on = dt.date(2026, 1, 1)
        config.work_days = [0, 1, 2, 3, 4, 5, 6]
        config.save()
        User.objects.update(date_joined=dt.datetime(2026, 1, 1, tzinfo=TZ))

    def test_sessions_become_present_late_and_absent(self):
        day = timezone.localdate() - dt.timedelta(days=1)
        AttendanceSession.objects.create(user=self.manager, work_date=day, status="closed",
                                         checked_in_at=dt.datetime.combine(day, dt.time(8, 55), TZ),
                                         checked_out_at=dt.datetime.combine(day, dt.time(18, 5), TZ))
        AttendanceSession.objects.create(user=self.hr, work_date=day, status="closed", late_minutes=25,
                                         checked_in_at=dt.datetime.combine(day, dt.time(9, 35), TZ),
                                         checked_out_at=dt.datetime.combine(day, dt.time(18, 0), TZ))
        data = self.client_for(self.od).get(f"{API}/attendance/", {"period": "yesterday"}).data
        status = {r["user"]["id"]: r["status"] for r in data["rows"]}
        self.assertEqual(status[self.manager.pk], "present")
        self.assertEqual(status[self.hr.pk], "late")
        self.assertEqual(status[self.marketer.pk], "absent")
        self.assertEqual(data["summary"]["late_minutes"], 25)

    def test_absence_issue_only_for_finished_days(self):
        issue_service.scan()
        kinds = OperationalIssue.objects.filter(kind=IssueKind.STAFF_ABSENCE)
        today = timezone.localdate().isoformat()
        self.assertFalse(any(i.payload["date"] == today for i in kinds))
        self.assertTrue(kinds.exists())


class ReportsAndSalesTests(OpsTestCase):
    def report(self, user, metrics=None, day=None, **extra):
        body = {"kind": "daily", "period_start": (day or timezone.localdate()).isoformat(), "done": "Звонки",
                "metrics": metrics or {}, **extra}
        client = self.client_for(user)
        response = client.post("/api/v1/work-reports/", body, format="json")
        assert response.status_code == 201, response.data
        submitted = client.post(f"/api/v1/work-reports/{response.data['id']}/submit/")
        assert submitted.status_code == 200, submitted.data
        return submitted.data

    def test_report_review_return_requires_comment_and_scope(self):
        report = self.report(self.manager, {"calls": 10})
        self.assertEqual(self.client_for(self.marketer).post(f"/api/v1/work-reports/{report['id']}/review/",
                                                             {"decision": "accept"}).status_code, 404)
        self.assertEqual(self.client_for(self.rop).post(f"/api/v1/work-reports/{report['id']}/review/",
                                                        {"decision": "return"}).status_code, 400)
        response = self.client_for(self.rop).post(f"/api/v1/work-reports/{report['id']}/review/",
                                                  {"decision": "return", "comment": "Нет встреч"})
        self.assertEqual(response.data["status"], "returned")

    def test_unknown_or_negative_metrics_rejected(self):
        client = self.client_for(self.manager)
        bad = client.post("/api/v1/work-reports/", {"kind": "daily", "metrics": {"likes": 3}}, format="json")
        self.assertEqual(bad.status_code, 400)
        neg = client.post("/api/v1/work-reports/", {"kind": "daily", "metrics": {"calls": -1}}, format="json")
        self.assertEqual(neg.status_code, 400)

    def test_missing_report_only_when_rule_set_and_due_passed(self):
        today = timezone.localdate()
        self.assertEqual(obligations(today - dt.timedelta(days=3), today), [])  # правило не задано
        config = OperationsSettings.load()
        config.required_reports = {"sales_manager": ["daily"]}
        config.save()
        self.manager.date_joined = timezone.now() - dt.timedelta(days=30)
        self.manager.save()
        items = obligations(today - dt.timedelta(days=2), today)
        statuses = {o.period_start: o.status for o in items}
        self.assertEqual(statuses[today - dt.timedelta(days=1)], "missing")
        self.assertIn(statuses[today], ("pending", "missing"))
        self.report(self.manager, {"calls": 1}, day=today - dt.timedelta(days=1))
        statuses = {o.period_start: o.status for o in obligations(today - dt.timedelta(days=2), today)}
        self.assertEqual(statuses[today - dt.timedelta(days=1)], "submitted")

    def test_sales_funnel_conversion_and_plan_from_targets(self):
        self.report(self.manager, {"leads_processed": 20, "calls": 40, "trials_scheduled": 8, "trials_attended": 6,
                                   "sales": 5, "revenue": 50000})
        KPITarget.objects.create(metric="sales", department="sales", period="daily", target_value=10,
                                 valid_from=timezone.localdate() - dt.timedelta(days=1))
        data = self.client_for(self.od).get(f"{API}/sales/", {"period": "today"}).data
        self.assertFalse(data["crm_connected"])
        self.assertEqual(data["department"]["sales"], 5)
        self.assertEqual(data["department"]["conversion"], 25.0)
        self.assertEqual(data["department"]["trial_attendance"], 75.0)
        self.assertEqual(data["department"]["plan_achievement"], 50.0)
        head = self.client_for(self.od).get(f"{API}/sales/head/", {"period": "today"}).data
        self.assertEqual(head["heads"][0]["id"], self.rop.pk)


class MarketingBudgetTests(OpsTestCase):
    def setUp(self):
        super().setUp()
        Campaign.objects.create(name="Осенний набор", start_date=timezone.localdate(), budget=10000, spent=6000,
                                leads=30, sales=6, responsible=self.marketer)

    def test_budget_and_cpl_visible_to_operations_director_by_default(self):
        data = self.client_for(self.od).get(f"{API}/marketing/").data
        self.assertTrue(data["budget_visible"])
        self.assertEqual(data["summary"]["cost_per_lead"], 200.0)
        self.assertEqual(data["campaigns"][0]["conversion"], 20.0)

    def test_budget_hidden_when_setting_disabled(self):
        config = OperationsSettings.load()
        config.operations_director_sees_marketing_budget = False
        config.save()
        data = self.client_for(self.od).get(f"{API}/marketing/").data
        self.assertNotIn("cost_per_lead", data["summary"])
        self.assertNotIn("budget", data["campaigns"][0])
        self.assertEqual(self.client_for(self.director).get(f"{API}/marketing/").data["summary"]["cost_per_lead"],
                         200.0)
        kpi = self.client_for(self.od).get(f"{API}/kpi/", {"department": "marketing"}).data
        self.assertNotIn("cost_per_lead", [m["key"] for m in kpi["departments"][0]["metrics"]])

    def test_marketing_cabinet_campaign_crud(self):
        client = jwt_client(self.marketer.username)
        response = client.post("/api/v1/marketing/campaigns/", {"name": "Instagram Reels", "start_date": "2026-10-01",
                                                                "budget": "5000", "spent": "2500", "leads": 10,
                                                                "sales": 2}, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["cost_per_lead"], "250.00")
        bad = client.post("/api/v1/marketing/campaigns/", {"name": "x", "start_date": "2026-10-01", "leads": 1,
                                                           "sales": 3}, format="json")
        self.assertEqual(bad.status_code, 400)


class KPITargetsAndDashboardTests(OpsTestCase):
    def test_target_validation(self):
        client = self.client_for(self.od)
        ok = client.post(f"{API}/kpi/targets/", {"metric": "task_completion_rate", "department": "sales",
                                                 "period": "monthly", "target_value": "90"}, format="json")
        self.assertEqual(ok.status_code, 201, ok.data)
        wrong_dep = client.post(f"{API}/kpi/targets/", {"metric": "trainer_kpi", "department": "sales",
                                                        "period": "monthly", "target_value": "80"}, format="json")
        self.assertEqual(wrong_dep.status_code, 400)
        personal = client.post(f"{API}/kpi/targets/", {"metric": "sales", "employee": self.manager.pk,
                                                       "department": "hr", "period": "daily", "target_value": "3"},
                               format="json")
        self.assertEqual(personal.status_code, 201)
        self.assertEqual(personal.data["department"], "sales")  # отдел — по сотруднику
        nowhere = client.post(f"{API}/kpi/targets/", {"metric": "tasks_completed", "period": "monthly",
                                                      "target_value": "3"}, format="json")
        self.assertEqual(nowhere.status_code, 400)

    def test_dashboard_counts_real_tasks(self):
        done = self.create_task(self.od, assignee=self.trainer.pk)
        self.client_for(self.trainer).post(f"/api/v1/tasks/{done['id']}/submit/")
        late = self.create_task(self.od, assignee=self.manager.pk)
        Task.objects.filter(pk=late["id"]).update(deadline=timezone.now() - dt.timedelta(hours=3))
        self.create_task(self.od, assignee=self.hr.pk, requires_approval=True)
        data = self.client_for(self.od).get(f"{API}/dashboard/", {"period": "today"}).data
        self.assertEqual(data["metrics"]["tasks_completed"], 1)
        self.assertEqual(data["metrics"]["tasks_overdue"], 1)
        self.assertEqual(data["overdue_tasks"][0]["id"], late["id"])
        self.assertIsNone(data["metrics"]["kpi"]["achievement"])  # целей нет — норматив не выдумывается
        sales = self.client_for(self.od).get(f"{API}/dashboard/", {"period": "today", "department": "sales"}).data
        self.assertEqual(sales["metrics"]["tasks_completed"], 0)

    def test_overdue_task_issue_dedup_and_resolution(self):
        task = self.create_task(self.od, assignee=self.manager.pk)
        Task.objects.filter(pk=task["id"]).update(deadline=timezone.now() - dt.timedelta(hours=3))
        issue_service.scan()
        issue_service.scan()
        issues = OperationalIssue.objects.filter(kind=IssueKind.TASK_OVERDUE)
        self.assertEqual(issues.count(), 1)
        self.assertEqual(self.manager.ops_notifications.filter(kind="issue_detected").count(), 1)
        self.client_for(self.manager).post(f"/api/v1/tasks/{task['id']}/submit/")
        issue_service.scan()
        self.assertEqual(issues.get().status, "resolved")

    def test_issue_manual_workflow_and_history(self):
        client = self.client_for(self.od)
        created = client.post(f"{API}/issues/", {"title": "Нет таблицы", "department": "academy",
                                                 "responsible": self.lead.pk, "priority": "high"}, format="json")
        self.assertEqual(created.status_code, 201)
        iid = created.data["id"]
        self.assertEqual(client.post(f"{API}/issues/{iid}/status/", {"status": "dismissed"}).status_code, 400)
        client.post(f"{API}/issues/{iid}/comment/", {"text": "Связались"})
        resolved = client.post(f"{API}/issues/{iid}/status/", {"status": "resolved", "comment": "ok"})
        self.assertEqual(resolved.data["status"], "resolved")
        self.assertEqual([e["action"] for e in resolved.data["events"]], ["detected", "comment", "status"])

    def test_kpi_below_target_counts_only_complete_periods(self):
        KPITarget.objects.create(metric="tasks_completed", department="sales", period="daily", target_value=5,
                                 valid_from=timezone.localdate() - dt.timedelta(days=5))
        issue_service.scan()
        issue = OperationalIssue.objects.get(kind=IssueKind.KPI_BELOW_TARGET)
        yesterday = (timezone.localdate() - dt.timedelta(days=1)).isoformat()
        self.assertEqual(issue.payload["period_start"], yesterday)

    def test_kpi_overview_and_history(self):
        KPITarget.objects.create(metric="task_completion_rate", department="academy", period="monthly",
                                 target_value=80, valid_from=dt.date(2026, 1, 1))
        task = self.create_task(self.od, assignee=self.trainer.pk)
        Task.objects.filter(pk=task["id"]).update(deadline=timezone.now() - dt.timedelta(hours=1))
        data = self.client_for(self.od).get(f"{API}/kpi/", {"department": "academy"}).data
        metric = next(m for m in data["departments"][0]["metrics"] if m["key"] == "task_completion_rate")
        self.assertEqual(metric["actual"], 0.0)
        self.assertEqual(metric["target"], 80.0)
        self.assertEqual(metric["achievement"], 0.0)
        history = self.client_for(self.od).get(f"{API}/kpi/history/", {"metric": "task_completion_rate",
                                                                      "department": "academy"}).data
        self.assertEqual(len(history["history"]), 6)


class ExportAndSettingsTests(OpsTestCase):
    def test_csv_and_excel_exports_are_logged(self):
        self.create_task(self.od, assignee=self.trainer.pk, title="Экспорт")
        client = self.client_for(self.od)
        csv = client.get(f"{API}/reports/export/", {"kind": "tasks", "export_format": "csv"})
        self.assertEqual(csv.status_code, 200)
        self.assertIn("Экспорт", csv.content.decode("utf-8-sig"))
        xlsx = client.get(f"{API}/reports/export/", {"kind": "monthly", "export_format": "xlsx"})
        self.assertEqual(xlsx.status_code, 200)
        from openpyxl import load_workbook

        book = load_workbook(BytesIO(xlsx.content))
        self.assertIn("Задачи", book.sheetnames)
        history = client.get(f"{API}/reports/", {"kind": "department"}).data["history"]
        self.assertEqual(len(history), 2)

    def test_settings_validation(self):
        client = self.client_for(self.od)
        ok = client.patch(f"{API}/settings/", {"required_reports": {"teacher": ["daily"]}}, format="json")
        self.assertEqual(ok.status_code, 200)
        bad = client.patch(f"{API}/settings/", {"required_reports": {"wizard": ["daily"]}}, format="json")
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(self.client_for(self.lead).patch(f"{API}/settings/", {}, format="json").status_code, 403)

    def test_employee_profile_and_department_detail(self):
        self.create_task(self.od, assignee=self.manager.pk)
        profile = self.client_for(self.od).get(f"{API}/employees/{self.manager.pk}/").data
        self.assertEqual(profile["employee"]["department"], "sales")
        self.assertEqual(profile["tasks"]["open"], 1)
        rop = self.client_for(self.od).get(f"{API}/employees/{self.rop.pk}/").data
        self.assertIn(self.manager.pk, [s["id"] for s in rop["subordinates"]])
        dep = self.client_for(self.od).get(f"{API}/departments/sales/").data
        self.assertEqual(dep["department"]["employees"], 2)
        self.assertEqual(self.client_for(self.od).get(f"{API}/departments/nope/").status_code, 404)

    def test_team_lead_worklog_reports_count_as_submitted(self):
        from apps.worklog.models import TeamLeadReport

        config = OperationsSettings.load()
        config.required_reports = {"team_lead": ["daily"]}
        config.save()
        self.lead.date_joined = timezone.now() - dt.timedelta(days=10)
        self.lead.save()
        yesterday = timezone.localdate() - dt.timedelta(days=1)
        TeamLeadReport.objects.create(author=self.lead, kind="daily", date=yesterday, status="submitted")
        items = {o.period_start: o for o in obligations(yesterday, yesterday, users=[self.lead])}
        self.assertEqual(items[yesterday].status, "submitted")
        self.assertEqual(items[yesterday].source, "worklog")

    def test_hr_overview_has_no_salary_data(self):
        data = self.client_for(self.od).get(f"{API}/hr/").data
        self.assertFalse(any("salary" in key for key in data))
        self.assertGreaterEqual(data["active_employees"], 12)
        self.assertEqual(EmployeeReport.objects.count(), 0)
