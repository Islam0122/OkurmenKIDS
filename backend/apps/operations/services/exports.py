"""Отчёты кабинета: построение таблиц и выгрузка в CSV / Excel.

Виды: daily, weekly, monthly (операционный отчёт за период), department,
employee, kpi, attendance, tasks. Каждый вид — набор листов (заголовок +
строки) из тех же сервисов, что и страницы кабинета, поэтому выгрузка всегда
совпадает с экраном. Печать — на фронтенде (печатная вёрстка страницы)."""
from __future__ import annotations

import csv
import datetime as dt
import io

from django.utils import timezone

from ..departments import Department
from ..models import Task, TaskStatus
from . import staff
from .org import department_overview, employees
from .overview import kpi_overview
from .period import OpsFilters
from .reports import obligations

KINDS = [
    ("daily", "Ежедневный отчёт"),
    ("weekly", "Еженедельный отчёт"),
    ("monthly", "Ежемесячный отчёт"),
    ("department", "Отчёт по отделам"),
    ("employee", "Отчёт по сотрудникам"),
    ("kpi", "KPI-отчёт"),
    ("attendance", "Отчёт по посещаемости"),
    ("tasks", "Отчёт по задачам"),
]
STATUS_RU = {"submitted": "Сдан", "missing": "Не сдан", "pending": "Срок не наступил"}


def _fmt(value):
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return timezone.localtime(value).strftime("%d.%m.%Y %H:%M")
    if isinstance(value, dt.date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    return value


def _department_sheet(f):
    rows = [[d["label"], d["employees"], d["tasks"]["completed"], d["tasks"]["open"], d["tasks"]["overdue"],
             d["tasks"]["completion_rate"], d["attendance_rate"], d["reports"]["submitted"], d["reports"]["expected"],
             d["open_issues"], d["performance"]] for d in department_overview(f)]
    return ("Отделы", ["Отдел", "Сотрудники", "Выполнено задач", "Открыто", "Просрочено", "В срок, %",
                       "Посещаемость, %", "Отчётов сдано", "Отчётов ожидалось", "Открытых проблем",
                       "Сводный показатель"], rows)


def _employee_sheet(f):
    rows = []
    for e in employees(f):
        att = e["attendance"] or {}
        rows.append([e["full_name"], Department(e["department"]).label, e["role_label"], e["tasks"]["completed"],
                     e["tasks"]["open"], e["tasks"]["overdue"], e["reports"]["submitted"], e["reports"]["missing"],
                     att.get("attendance_rate"), att.get("late_minutes"), att.get("worked_hours")])
    return ("Сотрудники", ["Сотрудник", "Отдел", "Должность", "Выполнено задач", "Открыто", "Просрочено",
                           "Отчётов сдано", "Отчётов не сдано", "Посещаемость, %", "Опоздания, мин", "Часы"], rows)


def _tasks_sheet(f):
    now = timezone.now()
    qs = Task.objects.filter(created_at__date__lte=f.end).exclude(
        completed_at__date__lt=f.start).select_related("assignee", "created_by").order_by("department", "deadline")
    if f.department:
        qs = qs.filter(department=f.department)
    if f.employee_id:
        qs = qs.filter(assignee_id=f.employee_id)
    rows = [[t.title, t.get_department_display(), str(t.assignee or "—"), str(t.created_by), t.get_priority_display(),
             TaskStatus(t.effective_status(now)).label, _fmt(t.deadline), _fmt(t.completed_at),
             "Да" if t.completed_at and t.deadline and t.completed_at <= t.deadline else ""] for t in qs[:5000]]
    return ("Задачи", ["Задача", "Отдел", "Исполнитель", "Автор", "Приоритет", "Статус", "Дедлайн", "Выполнена",
                       "В срок"], rows)


def _attendance_sheet(f):
    users = staff.tracked_users(f.users_qs())
    rows = []
    for row in staff.rows(users, f.start, f.elapsed_end):
        r = staff.serialize_row(row)
        rows.append([r["user"]["full_name"], Department(r["user"]["department"]).label, r["user"]["role_label"],
                     _fmt(row.day), r["status_label"], _fmt(r["check_in"]), _fmt(r["check_out"]),
                     round(r["worked_minutes"] / 60, 2), r["late_minutes"], "Да" if r["geo_verified"] else "Нет"])
    return ("Посещаемость", ["Сотрудник", "Отдел", "Должность", "Дата", "Статус", "Приход", "Уход", "Часы",
                             "Опоздание, мин", "Геолокация подтверждена"], rows)


def _reports_sheet(f):
    rows = [[o.user.get_full_name() or o.user.username, o.user.get_role_display(), o.kind, _fmt(o.period_start),
             _fmt(o.due), STATUS_RU[o.status]] for o in obligations(f.start, f.end, users=list(f.users_qs()))]
    return ("Отчёты сотрудников", ["Сотрудник", "Должность", "Вид", "Период с", "Срок", "Статус"], rows)


def _kpi_sheet(f, can_see_finance):
    data = kpi_overview(f, can_see_finance=can_see_finance)
    rows = [[d["label"], m["label"], m["actual"], m["target"], m["achievement"], m["previous"],
             (m["trend"] or {}).get("change"), m["source"]] for d in data["departments"] for m in d["metrics"]]
    return ("KPI", ["Отдел", "Показатель", "Факт", "Цель", "Достижение, %", "Прошлый период", "Изменение, %",
                    "Источник"], rows)


def build(kind: str, f: OpsFilters, *, can_see_finance: bool) -> list[tuple[str, list, list]]:
    if kind in ("daily", "weekly", "monthly"):
        return [_department_sheet(f), _tasks_sheet(f), _reports_sheet(f), _attendance_sheet(f),
                _kpi_sheet(f, can_see_finance)]
    return {
        "department": lambda: [_department_sheet(f)],
        "employee": lambda: [_employee_sheet(f), _reports_sheet(f)],
        "kpi": lambda: [_kpi_sheet(f, can_see_finance)],
        "attendance": lambda: [_attendance_sheet(f)],
        "tasks": lambda: [_tasks_sheet(f)],
    }[kind]()


def to_csv(sheets) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    for index, (title, headers, rows) in enumerate(sheets):
        if index:
            writer.writerow([])
        writer.writerow([title])
        writer.writerow(headers)
        writer.writerows([[_fmt(v) for v in row] for row in rows])
    # BOM — чтобы Excel открыл кириллицу без вопросов.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def to_xlsx(sheets, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append(headers)
        for cell in ws[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="16A34A")
        for row in rows:
            ws.append([_fmt(v) for v in row])
        for i, header in enumerate(headers, start=1):
            width = max([len(str(header))] + [len(str(_fmt(r[i - 1]))) for r in rows[:200]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 48)
        ws.freeze_panes = "A3"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
