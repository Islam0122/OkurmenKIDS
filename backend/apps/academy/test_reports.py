"""Tests for Reports (services.reports, /api/v1/reports/, admin pages).

Absolute imports only — see the note at the top of apps/academy/tests.py.
"""
from __future__ import annotations

import datetime as dt
import io
import re

from django.db import connection
from django.test import Client as DjangoClient
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from apps.academy.models import (
    Attendance,
    Course,
    Group,
    GroupTeacher,
    Homework,
    HomeworkResult,
    Lesson,
    Student,
    StudentStatusEvent,
)
from apps.academy.services.academy_monthly_report import compute_academy_monthly_stats
from apps.academy.services.reports import (
    ReportFilterError,
    ReportFilters,
    build_full_report,
    build_group_detail,
    build_group_rows,
    build_overview,
    build_teacher_detail,
    build_teacher_rows,
    group_student_rows,
    kpi_level,
    overall_kpi,
)
from apps.academy.services.reports.excel import build_reports_excel
from apps.academy.services.reports.pdf import build_reports_pdf
from apps.academy.services.reports.table import paginate_groups
from apps.users.models import Subject, Teacher, User

TODAY = dt.date(2026, 9, 30)
SEPT = {"period": "custom", "start_date": "2026-09-01", "end_date": "2026-09-30"}


def make_teacher(username: str, *, first_name: str | None = None) -> Teacher:
    user = User.objects.create_user(
        username=username, email=f"{username}@okurmen.kg", password="Str0ngPassw0rd!",
        first_name=first_name or username.capitalize(), role=User.Role.TEACHER, is_verified=True,
    )
    return Teacher.objects.create(user=user)


class ReportsTestBase(TestCase):
    """Two groups, two teachers, one lesson plan of real attendance/homework
    in September 2026, plus a group with no teacher and nothing in it."""

    def setUp(self):
        self.admin = User.objects.create_superuser(
            username="admin", email="admin@okurmen.kg", password="Str0ngPassw0rd!", first_name="Admin"
        )
        self.python, _ = Subject.objects.get_or_create(name="Python")
        self.english, _ = Subject.objects.get_or_create(name="English")
        self.course = Course.objects.create(name="Python PRO", count_lesson=10)
        self.course.subjects.set([self.python, self.english])
        self.other_course = Course.objects.create(name="English Kids", count_lesson=10)
        self.other_course.subjects.set([self.english])

        self.aziza = make_teacher("aziza", first_name="Aziza")
        self.bek = make_teacher("bek", first_name="Bek")

        self.pro1 = Group.objects.create(name="PRO-01", course=self.course, start_date=dt.date(2026, 8, 1))
        self.eng1 = Group.objects.create(name="ENG-01", course=self.other_course, start_date=dt.date(2026, 8, 1))
        self.empty = Group.objects.create(name="EMPTY-01", course=self.course, start_date=dt.date(2026, 9, 1))

        self.gt_pro_py = GroupTeacher.objects.create(group=self.pro1, teacher=self.aziza, subject=self.python)
        self.gt_pro_en = GroupTeacher.objects.create(group=self.pro1, teacher=self.bek, subject=self.english)
        self.gt_eng = GroupTeacher.objects.create(group=self.eng1, teacher=self.bek, subject=self.english)

        self.s1 = Student.objects.create(first_name="Ali", group=self.pro1)
        self.s2 = Student.objects.create(first_name="Bota", group=self.pro1)
        self.s3 = Student.objects.create(first_name="Cholpon", group=self.eng1)
        self.inactive = Student.objects.create(
            first_name="Dana", group=self.pro1, is_active=False, status=Student.Status.WITHDRAWN
        )

        # PRO-01 / Aziza (Python): 2 completed lessons + 1 cancelled.
        self.l1 = self._lesson(self.gt_pro_py, dt.date(2026, 9, 2), Lesson.Status.COMPLETED)
        self.l2 = self._lesson(self.gt_pro_py, dt.date(2026, 9, 9), Lesson.Status.COMPLETED)
        self.l3 = self._lesson(self.gt_pro_py, dt.date(2026, 9, 16), Lesson.Status.CANCELLED)
        # PRO-01 / Bek (English): 1 completed lesson.
        self.l4 = self._lesson(self.gt_pro_en, dt.date(2026, 9, 3), Lesson.Status.COMPLETED)
        # ENG-01 / Bek: 1 completed lesson.
        self.l5 = self._lesson(self.gt_eng, dt.date(2026, 9, 4), Lesson.Status.COMPLETED)

        # Attendance: Aziza's lessons 3 attended of 4; Bek's PRO lesson 1 of 2; ENG 1 of 1.
        self._attend(self.l1, self.s1, Attendance.Status.PRESENT)
        self._attend(self.l1, self.s2, Attendance.Status.LATE)
        self._attend(self.l2, self.s1, Attendance.Status.PRESENT)
        self._attend(self.l2, self.s2, Attendance.Status.ABSENT)
        self._attend(self.l4, self.s1, Attendance.Status.PRESENT)
        self._attend(self.l4, self.s2, Attendance.Status.EXCUSED)
        self._attend(self.l5, self.s3, Attendance.Status.PRESENT)

        # Homework on l1: s1 checked (score 8), s2 not submitted.
        hw = Homework.objects.create(lesson=self.l1, title="HW1")
        HomeworkResult.objects.create(homework=hw, student=self.s1, status=HomeworkResult.Status.CHECKED, score=8)
        HomeworkResult.objects.create(homework=hw, student=self.s2, status=HomeworkResult.Status.NOT_SUBMITTED)
        # Homework on l5: s3 submitted, score 10.
        hw2 = Homework.objects.create(lesson=self.l5, title="HW2")
        HomeworkResult.objects.create(homework=hw2, student=self.s3, status=HomeworkResult.Status.SUBMITTED, score=10)

        # Dana left PRO-01 in September.
        StudentStatusEvent.objects.create(
            student=self.inactive, event_type=StudentStatusEvent.EventType.DEACTIVATED,
            reason=StudentStatusEvent.Reason.RELOCATION, group=self.pro1, event_date=dt.date(2026, 9, 10),
        )

    _lesson_no = 0

    def _lesson(self, gt, date, lesson_status):
        self._lesson_no += 1
        return Lesson.objects.create(
            group=gt.group, group_teacher=gt, teacher=gt.teacher, subject=gt.subject, lesson_number=self._lesson_no,
            date=date, start_time=dt.time(10), end_time=dt.time(11), status=lesson_status,
        )

    def _attend(self, lesson, student, att_status):
        Attendance.objects.create(lesson=lesson, student=student, status=att_status)

    def filters(self, **params):
        return ReportFilters.from_query({**SEPT, **params}, today=TODAY)


class ReportFiltersTests(TestCase):
    def test_default_is_current_month(self):
        f = ReportFilters.from_query({}, today=TODAY)
        self.assertEqual((f.period, f.start, f.end), ("this_month", dt.date(2026, 9, 1), TODAY))

    def test_presets(self):
        cases = {
            "today": (TODAY, TODAY),
            "this_week": (dt.date(2026, 9, 28), dt.date(2026, 10, 4)),
            "last_month": (dt.date(2026, 8, 1), dt.date(2026, 8, 31)),
            "this_quarter": (dt.date(2026, 7, 1), TODAY),
        }
        for period, expected in cases.items():
            f = ReportFilters.from_query({"period": period}, today=TODAY)
            self.assertEqual((f.start, f.end), expected, period)

    def test_custom_range_and_swapped_dates(self):
        f = ReportFilters.from_query({"start_date": "2026-09-20", "end_date": "2026-09-05"}, today=TODAY)
        self.assertEqual((f.period, f.start, f.end), ("custom", dt.date(2026, 9, 5), dt.date(2026, 9, 20)))

    def test_invalid_input_raises_filter_error(self):
        for params in ({"start_date": "31.09.2026", "end_date": "2026-09-30"}, {"period": "decade"},
                       {"period": "custom", "start_date": "2026-09-01"}):
            with self.assertRaises(ReportFilterError):
                ReportFilters.from_query(params, today=TODAY)

    def test_as_query_round_trips(self):
        f = ReportFilters.from_query({**SEPT, "group": "3", "teacher": "4"}, today=TODAY)
        self.assertEqual(ReportFilters.from_query(f.as_query(), today=TODAY), f)


class KpiFormulaTests(TestCase):
    def test_equal_default_weights_average_available_components(self):
        self.assertEqual(overall_kpi({"attendance": 80, "homework": 90, "activity": 100, "progress": None}), 90.0)

    def test_no_components_means_no_kpi(self):
        self.assertIsNone(overall_kpi({"attendance": None, "homework": None, "activity": None, "progress": None}))

    @override_settings(REPORTS_KPI_WEIGHTS={"attendance": 0.4, "homework": 0.3, "activity": 0.2, "progress": 0.1})
    def test_configured_weights(self):
        value = overall_kpi({"attendance": 90, "homework": 80, "activity": 100, "progress": 70})
        self.assertEqual(value, round(90 * 0.4 + 80 * 0.3 + 100 * 0.2 + 70 * 0.1, 1))
        # A missing component's weight is redistributed, not counted as 0%.
        value = overall_kpi({"attendance": 90, "homework": 80, "activity": 100, "progress": None})
        self.assertEqual(value, round((90 * 0.4 + 80 * 0.3 + 100 * 0.2) / 0.9, 1))

    def test_levels(self):
        self.assertEqual([kpi_level(v) for v in (95, 90, 89.9, 75, 74.9, None)],
                         ["good", "good", "warning", "warning", "bad", "none"])


class ReportCalculationTests(ReportsTestBase):
    def test_group_rows_from_real_data(self):
        rows = {r["name"]: r for r in build_group_rows(self.filters())}
        pro = rows["PRO-01"]
        self.assertEqual(pro["students"], {"total": 3, "active": 2, "left": 1, "new": 3, "returned": 0})
        self.assertEqual(pro["lessons"]["total"], 4)
        self.assertEqual(pro["lessons"]["held"], 3)
        # (present+late) / all marks: l1 2/2, l2 1/2, l4 1/2 -> 4/6
        self.assertEqual(pro["attendance_rate"], 66.7)
        self.assertEqual(pro["homework_rate"], 50.0)
        self.assertEqual(pro["activity_rate"], 75.0)
        self.assertEqual(pro["progress_rate"], 80.0)
        self.assertEqual(pro["kpi"], round((66.7 + 50 + 75 + 80) / 4, 1))
        self.assertEqual(pro["teacher_names"], "Bek, Aziza")  # ordered by subject: English, Python
        self.assertEqual(pro["subjects"], ["English", "Python"])

    def test_empty_group_and_group_without_teacher(self):
        rows = {r["name"]: r for r in build_group_rows(self.filters())}
        empty = rows["EMPTY-01"]
        self.assertFalse(empty["has_teacher"])
        self.assertEqual(empty["teacher_names"], "Не назначен")
        self.assertEqual(empty["students"]["total"], 0)
        self.assertIsNone(empty["kpi"])
        self.assertEqual(empty["kpi_level"], "none")
        self.assertIsNone(empty["attendance_rate"])  # "no data", never a fabricated 0%

    def test_teacher_rows_are_isolated_to_own_lessons(self):
        rows = {r["name"]: r for r in build_teacher_rows(self.filters())}
        aziza, bek = rows["Aziza"], rows["Bek"]
        self.assertEqual([g["name"] for g in aziza["groups"]], ["PRO-01"])
        self.assertEqual(aziza["subjects"], ["Python"])
        self.assertEqual(aziza["attendance_rate"], 75.0)  # 3 of 4 on her own lessons only
        self.assertEqual(aziza["activity_rate"], 66.7)  # 2 held of 3 (one cancelled)
        self.assertEqual(sorted(g["name"] for g in bek["groups"]), ["ENG-01", "PRO-01"])
        self.assertEqual(bek["attendance_rate"], round(2 / 3 * 100, 1))
        self.assertEqual(bek["students"], {"total": 4, "active": 3, "left": 1})

    def test_teacher_without_groups(self):
        make_teacher("nogroups", first_name="Nurlan")
        rows = {r["name"]: r for r in build_teacher_rows(self.filters())}
        self.assertEqual(rows["Nurlan"]["groups_count"], 0)
        self.assertIsNone(rows["Nurlan"]["kpi"])
        detail = build_teacher_detail(Teacher.objects.get(user__username="nogroups"), self.filters())
        self.assertEqual(detail["groups"], [])
        self.assertIsNone(detail["kpi"]["overall"])

    def test_overview_totals(self):
        ov = build_overview(self.filters())
        self.assertEqual(ov["students"]["total"], 4)
        self.assertEqual(ov["students"]["active"], 3)
        self.assertEqual(ov["students"]["left"], 1)
        self.assertEqual(ov["students"]["retention_rate"], 75.0)
        self.assertEqual(ov["teachers"]["total"], 2)
        self.assertEqual(ov["groups"]["total"], 3)
        self.assertEqual(ov["groups"]["without_teacher"], 1)
        self.assertEqual(ov["attendance"]["total"], 7)
        self.assertEqual(ov["kpi"]["attendance"], round(5 / 7 * 100, 1))

    def test_overall_kpi_matches_existing_academy_kpi(self):
        """Same components, equal default weights -> the existing Academy
        Monthly Report's KPI for the same (fully past) month."""
        ov = build_overview(self.filters())
        existing = compute_academy_monthly_stats(2026, 9)["kpi"]
        self.assertEqual(ov["kpi"]["attendance"], existing["attendance"])
        self.assertEqual(ov["kpi"]["homework"], existing["homework"])
        self.assertEqual(ov["kpi"]["activity"], existing["lessons"])
        self.assertEqual(ov["kpi"]["progress"], existing["student_progress"])
        self.assertEqual(ov["kpi"]["overall"], existing["total"])

    def test_period_without_data(self):
        f = ReportFilters.from_query({"period": "custom", "start_date": "2026-01-01", "end_date": "2026-01-31"},
                                     today=TODAY)
        ov = build_overview(f)
        self.assertFalse(ov["has_data"])
        self.assertIsNone(ov["kpi"]["overall"])
        self.assertEqual(ov["students"]["left"], 0)
        # Build exports for an empty period without errors.
        report = build_full_report(f)
        self.assertTrue(build_reports_pdf(report).startswith(b"%PDF"))
        build_reports_excel(report, "—")

    def test_filters_narrow_every_figure(self):
        by_group = build_overview(self.filters(group=str(self.eng1.id)))
        self.assertEqual(by_group["groups"]["total"], 1)
        self.assertEqual(by_group["students"]["total"], 1)
        self.assertEqual(by_group["attendance"]["total"], 1)

        by_teacher = build_overview(self.filters(teacher=str(self.aziza.id)))
        self.assertEqual(by_teacher["attendance"]["total"], 4)
        self.assertEqual(by_teacher["teachers"]["total"], 1)

        by_subject = build_overview(self.filters(subject=str(self.english.id)))
        self.assertEqual(by_subject["attendance"]["total"], 3)  # l4 (2 marks) + l5 (1)

        by_program = build_group_rows(self.filters(program=str(self.other_course.id)))
        self.assertEqual([r["name"] for r in by_program], ["ENG-01"])

    def test_activity_ignores_lessons_not_yet_due(self):
        self._lesson(self.gt_eng, dt.date(2026, 9, 29), Lesson.Status.SCHEDULED)
        f = ReportFilters.from_query(SEPT, today=dt.date(2026, 9, 20))
        rows = {r["name"]: r for r in build_group_rows(f)}
        self.assertEqual(rows["ENG-01"]["activity_rate"], 100.0)

    def test_returned_student(self):
        StudentStatusEvent.objects.create(
            student=self.inactive, event_type=StudentStatusEvent.EventType.REACTIVATED,
            group=self.pro1, event_date=dt.date(2026, 9, 20),
        )
        self.assertEqual(build_overview(self.filters())["students"]["returned"], 1)

    def test_group_detail_and_student_rows(self):
        detail = build_group_detail(self.pro1, self.filters())
        self.assertEqual(detail["group"]["teacher_names"], "Bek, Aziza")
        self.assertEqual(detail["homework"]["assigned"], 1)
        self.assertEqual(detail["homework"]["not_submitted"], 1)
        self.assertEqual(detail["homework"]["checked"], 1)
        self.assertEqual(len(detail["teacher_breakdown"]), 2)
        students = {r["name"]: r for r in group_student_rows(self.pro1, self.filters())}
        self.assertEqual(students["Ali"]["attendance_rate"], 100.0)
        self.assertEqual(students["Bota"]["attendance_rate"], round(1 / 3 * 100, 1))
        self.assertIsNone(students["Dana"]["attendance_rate"])
        self.assertFalse(students["Dana"]["is_active"])

    def test_query_count_does_not_grow_with_groups(self):
        def count():
            with CaptureQueriesContext(connection) as ctx:
                build_full_report(self.filters())
            return len(ctx.captured_queries)

        baseline = count()
        for index in range(5):
            group = Group.objects.create(name=f"EXTRA-{index}", course=self.course, start_date=dt.date(2026, 9, 1))
            gt = GroupTeacher.objects.create(group=group, teacher=self.aziza, subject=self.python)
            student = Student.objects.create(first_name=f"S{index}", group=group)
            lesson = self._lesson(gt, dt.date(2026, 9, 5), Lesson.Status.COMPLETED)
            self._attend(lesson, student, Attendance.Status.PRESENT)
        self.assertEqual(count(), baseline)

    def test_search_sort_paginate(self):
        rows = build_group_rows(self.filters())
        page = paginate_groups(rows, {"sort": "-kpi"})
        self.assertEqual(page.rows[-1]["name"], "EMPTY-01")  # "no data" sinks to the bottom
        page = paginate_groups(rows, {"q": "aziza"})
        self.assertEqual([r["name"] for r in page.rows], ["PRO-01"])
        page = paginate_groups(rows, {"page_size": "2", "page": "2"})
        self.assertEqual((page.total, page.pages, len(page.rows)), (3, 2, 1))


class ReportExportTests(ReportsTestBase):
    def test_pdf(self):
        pdf = build_reports_pdf(build_full_report(self.filters()))
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertGreater(len(pdf), 5000)

    def test_excel_sheets_and_filters(self):
        f = self.filters(group=str(self.pro1.id))
        wb = load_workbook(io.BytesIO(build_reports_excel(build_full_report(f), "Группа: PRO-01")))
        self.assertEqual(
            wb.sheetnames, ["Обзор", "Группы", "Тренеры", "Студенты", "Посещаемость", "Домашние задания", "KPI"]
        )
        groups = [row[0] for row in wb["Группы"].iter_rows(min_row=5, values_only=True)]
        self.assertEqual(groups, ["PRO-01"])
        students = sorted(row[0] for row in wb["Студенты"].iter_rows(min_row=5, values_only=True))
        self.assertEqual(students, ["Ali", "Bota", "Dana"])
        # Percentages are numbers (0..1) with a percent format.
        cell = wb["Группы"].cell(row=5, column=14)
        self.assertAlmostEqual(cell.value, 0.667)
        self.assertEqual(cell.number_format, "0.0%")


class ReportsApiTests(ReportsTestBase):
    def setUp(self):
        super().setUp()
        self.api = APIClient()

    def _get(self, name, *args, user=None, **params):
        self.api.force_authenticate(user)
        return self.api.get(reverse(name, args=args), {**SEPT, **params})

    def test_admin_endpoints(self):
        for name, args in (
            ("reports-overview", ()), ("reports-filters", ()), ("reports-groups", ()),
            ("reports-group-detail", (self.pro1.id,)), ("reports-teachers", ()),
            ("reports-teacher-detail", (self.aziza.id,)),
        ):
            response = self._get(name, *args, user=self.admin)
            self.assertEqual(response.status_code, status.HTTP_200_OK, name)

    def test_groups_payload(self):
        data = self._get("reports-groups", user=self.admin, sort="-kpi").data
        self.assertEqual(data["count"], 3)
        # ENG-01: 100% on every component; PRO-01 lower; EMPTY-01 has no data.
        self.assertEqual([r["name"] for r in data["results"]], ["ENG-01", "PRO-01", "EMPTY-01"])
        self.assertEqual(data["filters"]["start_date"], dt.date(2026, 9, 1))

    def test_group_detail_students_list(self):
        data = self._get("reports-group-detail", self.pro1.id, user=self.admin).data
        self.assertEqual(data["students_list"]["count"], 3)
        self.assertEqual(data["students"]["left"], 1)

    def test_exports(self):
        pdf = self._get("reports-export-pdf", user=self.admin)
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertIn("okurmenkids-report-20260901-20260930.pdf", pdf["Content-Disposition"])
        xlsx = self._get("reports-export-excel", user=self.admin)
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn("spreadsheetml", xlsx["Content-Type"])

    def test_bad_filters_are_400(self):
        response = self._get("reports-overview", user=self.admin, start_date="bad")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_missing_objects_are_404(self):
        self.assertEqual(self._get("reports-group-detail", 9999, user=self.admin).status_code, 404)
        self.assertEqual(self._get("reports-teacher-detail", 9999, user=self.admin).status_code, 404)

    def test_teacher_and_anonymous_are_denied(self):
        for name in ("reports-overview", "reports-groups", "reports-teachers", "reports-export-pdf",
                     "reports-export-excel"):
            self.assertEqual(self._get(name, user=self.aziza.user).status_code, status.HTTP_403_FORBIDDEN, name)
            self.api.force_authenticate(None)
            self.assertIn(self.api.get(reverse(name)).status_code, (401, 403), name)


class ReportsAdminPageTests(ReportsTestBase):
    def setUp(self):
        super().setUp()
        self.web = DjangoClient()
        self.web.force_login(self.admin)

    def test_pages_render_with_real_numbers(self):
        pages = {
            "overview": reverse("admin:academy_reports_overview"),
            "groups": reverse("admin:academy_reports_groups"),
            "group": reverse("admin:academy_reports_group_detail", args=[self.pro1.id]),
            "teachers": reverse("admin:academy_reports_teachers"),
            "teacher": reverse("admin:academy_reports_teacher_detail", args=[self.bek.id]),
        }
        for key, url in pages.items():
            response = self.web.get(url, SEPT)
            self.assertEqual(response.status_code, 200, key)
        body = self.web.get(pages["groups"], SEPT).content.decode()
        self.assertIn("PRO-01", body)
        self.assertIn("66,7%", body)
        self.assertIn("Не назначен", body)
        body = self.web.get(pages["teacher"], SEPT).content.decode()
        self.assertIn("ENG-01", body)
        self.assertIn("English", body)

    def test_invalid_filters_show_error_state(self):
        response = self.web.get(reverse("admin:academy_reports_overview"), {"start_date": "x", "end_date": "y"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Не удалось построить отчёт")

    def test_empty_period_shows_empty_state(self):
        response = self.web.get(reverse("admin:academy_reports_overview"),
                                {"period": "custom", "start_date": "2026-01-01", "end_date": "2026-01-31"})
        self.assertContains(response, "Нет данных за выбранный период")

    def test_sidebar_section(self):
        body = self.web.get(reverse("admin:index")).content.decode()
        self.assertIn(reverse("admin:academy_reports_overview"), body)
        self.assertIn(reverse("admin:academy_reports_teachers"), body)

    def test_teacher_staff_account_is_forbidden(self):
        user = self.aziza.user
        user.is_staff = True
        user.save()
        web = DjangoClient()
        web.force_login(user)
        for name in ("admin:academy_reports_overview", "admin:academy_reports_groups", "admin:academy_reports_teachers"):
            self.assertEqual(web.get(reverse(name)).status_code, 403, name)
        self.assertNotIn(reverse("admin:academy_reports_overview"), web.get(reverse("admin:index")).content.decode())

    def test_anonymous_redirected_to_login(self):
        response = DjangoClient().get(reverse("admin:academy_reports_overview"))
        self.assertEqual(response.status_code, 302)


# Reports UI must be fully Russian and emoji-free (visible text, tooltips,
# aria-labels, placeholders — not just headings).
EMOJI = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B50\uFE0F]")
ENGLISH_UI = re.compile(
    r"\b(Reports?|Overview|Groups?|Teachers?|Students?|Attendance|Homework|Activity|Progress|Performance|"
    r"Download|Export|Refresh|Loading|Search|Filters?|No data|Excellent|Good|Needs attention|Overall|Total)\b"
)


def _report_ui_text(html: str) -> str:
    """Visible text plus title/aria-label/placeholder attributes of the
    Reports content area (the admin chrome around it is out of scope)."""
    start = html.index('<div class="okr">')
    end = html.index("<script>", start)
    body = re.sub(r"<(style|svg)\b.*?</\1>", " ", html[start:end], flags=re.S)
    attrs = re.findall(r'(?:title|aria-label|placeholder)="([^"]*)"', body)
    text = re.sub(r"<[^>]+>", " ", body)
    return " ".join([text, *attrs])


class ReportsLocalizationTests(ReportsTestBase):
    def setUp(self):
        super().setUp()
        self.web = DjangoClient()
        self.web.force_login(self.admin)

    def test_pages_are_russian_and_emoji_free(self):
        pages = [
            (reverse("admin:academy_reports_overview"), SEPT),
            (reverse("admin:academy_reports_groups"), SEPT),
            (reverse("admin:academy_reports_groups"), {**SEPT, "q": "нет-такой"}),
            (reverse("admin:academy_reports_group_detail", args=[self.pro1.id]), SEPT),
            (reverse("admin:academy_reports_group_detail", args=[self.empty.id]), SEPT),
            (reverse("admin:academy_reports_teachers"), SEPT),
            (reverse("admin:academy_reports_teacher_detail", args=[self.bek.id]), SEPT),
            (reverse("admin:academy_reports_overview"), {"period": "custom", "start_date": "2026-01-01",
                                                          "end_date": "2026-01-31"}),
            (reverse("admin:academy_reports_overview"), {"start_date": "x", "end_date": "y"}),
        ]
        for url, params in pages:
            html = self.web.get(url, params).content.decode()
            text = _report_ui_text(html)
            self.assertIsNone(EMOJI.search(html), url)
            self.assertIsNone(ENGLISH_UI.search(text), (url, params, ENGLISH_UI.findall(text)))

    def test_key_russian_labels(self):
        body = self.web.get(reverse("admin:academy_reports_overview"), SEPT).content.decode()
        for label in ("Отчёты академии", "Всего студентов", "Активные студенты", "Ушедшие студенты",
                      "Показатели эффективности", "Обновить", "Скачать PDF", "Экспорт в Excel", "Общий KPI"):
            self.assertIn(label, body)

    def test_sidebar_is_russian(self):
        body = self.web.get(reverse("admin:index")).content.decode()
        self.assertIn("bi bi-bar-chart-line ok-nav-section__icon", body)
        for label in ("отчёты", "Обзор"):
            self.assertIn(label, body)
        self.assertNotIn(">Overview<", body)

    def test_excel_is_russian_and_emoji_free(self):
        wb = load_workbook(io.BytesIO(build_reports_excel(build_full_report(self.filters()), "—")))
        for ws in wb:
            for row in ws.iter_rows(values_only=True):
                for value in row:
                    if isinstance(value, str):
                        self.assertIsNone(EMOJI.search(value), value)
                        self.assertIsNone(ENGLISH_UI.search(value), (ws.title, value))

