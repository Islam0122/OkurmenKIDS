"""Tests for Control (services.control, /api/v1/control/).

Absolute imports only — see the note at the top of apps/academy/tests.py.
"""
from __future__ import annotations

import datetime as dt

from django.urls import reverse
from rest_framework import status
from rest_framework.test import APIClient
from django.test import TestCase

from apps.academy.models import (
    Attendance,
    Course,
    Group,
    GroupTeacher,
    Homework,
    HomeworkResult,
    Lesson,
    Student,
)
from apps.academy.services.control import ControlQuery, ControlService
from apps.academy.services.reports import ReportFilters
from apps.users.models import Subject, Teacher, User

TODAY = dt.date(2026, 9, 30)
NOON = dt.time(12, 0)
SEPT = {"period": "custom", "start_date": "2026-09-01", "end_date": "2026-10-31"}
# API tests run against the real clock, so their lessons live safely in the past.
PAST = {"period": "custom", "start_date": "2025-03-01", "end_date": "2025-03-31"}


def make_teacher(username: str) -> Teacher:
    user = User.objects.create_user(
        username=username, email=f"{username}@okurmen.kg", password="Str0ngPassw0rd!",
        first_name=username.capitalize(), role=User.Role.TEACHER, is_verified=True,
    )
    return Teacher.objects.create(user=user)


class ControlTestBase(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(
            username="admin", email="admin@okurmen.kg", password="Str0ngPassw0rd!", first_name="Admin"
        )
        self.python, _ = Subject.objects.get_or_create(name="Python")
        self.soft, _ = Subject.objects.get_or_create(name="Soft Skills")
        self.course = Course.objects.create(name="Prog", count_lesson=20)
        self.course.subjects.set([self.python, self.soft])

        self.aizhan = make_teacher("aizhan")
        self.almaz = make_teacher("almaz")

        self.g1 = Group.objects.create(name="Prog Soft 2", course=self.course, start_date=dt.date(2025, 1, 1))
        self.g2 = Group.objects.create(name="Next 3", course=self.course, start_date=dt.date(2025, 1, 1))
        self.gt1 = GroupTeacher.objects.create(group=self.g1, teacher=self.aizhan, subject=self.python)
        self.gt1_soft = GroupTeacher.objects.create(group=self.g1, teacher=self.aizhan, subject=self.soft)
        self.gt2 = GroupTeacher.objects.create(group=self.g2, teacher=self.almaz, subject=self.python)

        self.s1 = Student.objects.create(first_name="Сумая", last_name="Камалидинова", group=self.g1)
        self.s2 = Student.objects.create(first_name="Али", last_name="Азамат уулу", group=self.g1)
        self.s3 = Student.objects.create(first_name="Артём", last_name="Иванов", group=self.g2)

    _n = 0

    def lesson(self, gt, date, lesson_status=Lesson.Status.COMPLETED, **extra) -> Lesson:
        self._n += 1
        extra.setdefault("teacher", gt.teacher)
        extra.setdefault("start_time", dt.time(10))
        extra.setdefault("end_time", dt.time(11))
        return Lesson.objects.create(
            group=gt.group, group_teacher=gt, subject=gt.subject, lesson_number=self._n,
            date=date, status=lesson_status, **extra,
        )

    def fill(self, lesson, *, attendance=True, homework=True, grades=True, deadline=None):
        """Make `lesson` fully filled in (or leave the given parts out)."""
        students = list(lesson.group.students.filter(is_active=True))
        if attendance:
            for s in students:
                Attendance.objects.create(lesson=lesson, student=s, status=Attendance.Status.PRESENT)
        if homework:
            hw = Homework.objects.create(lesson=lesson, title="HW", deadline=deadline)
            if grades:
                for s in students:
                    HomeworkResult.objects.create(homework=hw, student=s, status=HomeworkResult.Status.CHECKED, score=8)
            return hw
        return None

    def build(self, *, status=None, restrict=None, **params):
        filters = ReportFilters.from_query({**SEPT, **params}, today=TODAY)
        service = ControlService(ControlQuery(filters=filters, status=status, now_time=NOON), restrict_teacher=restrict)
        return service.build()

    def row(self, data, teacher, group):
        for item in data["items"]:
            if item["teacher"] and item["teacher"]["id"] == teacher.id and item["group"]["id"] == group.id:
                return item
        self.fail(f"no row for {teacher} / {group}")


class ControlRulesTests(ControlTestBase):
    def test_fully_completed_lesson_is_ok(self):
        self.fill(self.lesson(self.gt1, dt.date(2026, 9, 28)))
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["status"], "ok")
        self.assertEqual(row["lessons"]["total"], 1)
        self.assertEqual(row["attendance"], {"completed": 1, "total": 1, "percent": 100.0, "level": "ok"})
        self.assertEqual(row["homework"]["completed"], 1)
        self.assertEqual(row["grades"]["completed"], 1)
        self.assertEqual(row["issues"], [])

    def test_missing_attendance_is_attention(self):
        self.fill(self.lesson(self.gt1, dt.date(2026, 9, 28)), attendance=False)
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["status"], "attention")
        self.assertEqual((row["attendance"]["completed"], row["attendance"]["total"]), (0, 1))
        self.assertEqual(row["homework"]["level"], "ok")

    def test_missing_homework_is_attention(self):
        self.fill(self.lesson(self.gt1, dt.date(2026, 9, 28)), homework=False)
        data = self.build()
        row = self.row(data, self.aizhan, self.g1)
        self.assertEqual(row["status"], "attention")
        self.assertEqual((row["homework"]["completed"], row["homework"]["total"]), (0, 1))
        # No homework means nothing to score — not a second, separate failure.
        self.assertEqual(row["grades"]["total"], 0)
        self.assertIn("ДЗ не выдано: 1 занятие", row["issues"])

    def test_unchecked_homework_is_attention(self):
        lesson = self.lesson(self.gt1, dt.date(2026, 9, 28))
        hw = self.fill(lesson, grades=False)
        HomeworkResult.objects.create(homework=hw, student=self.s1, status=HomeworkResult.Status.SUBMITTED, score=7)
        HomeworkResult.objects.create(homework=hw, student=self.s2, status=HomeworkResult.Status.CHECKED, score=9)
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["status"], "attention")
        self.assertEqual(row["homework"]["completed"], 0)
        self.assertEqual(row["grades"]["completed"], 1)

    def test_missing_grades_is_problem_and_lists_students(self):
        lesson = self.lesson(self.gt1, dt.date(2026, 9, 28))
        hw = self.fill(lesson, grades=False)
        HomeworkResult.objects.create(homework=hw, student=self.s2, status=HomeworkResult.Status.CHECKED, score=9)
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["status"], "problem")
        self.assertEqual((row["grades"]["completed"], row["grades"]["total"]), (0, 1))
        self.assertEqual(row["grades"]["students_missing"], 1)

        detail = ControlService(
            ControlQuery(filters=ReportFilters.from_query(SEPT, today=TODAY), now_time=NOON)
        ).build_detail(group_id=self.g1.id, teacher_id=self.aizhan.id)
        grades = detail["lessons"][0]["grades"]
        self.assertEqual((grades["given"], grades["total"], grades["missing"]), (1, 2, 1))
        self.assertEqual([s["name"] for s in grades["missing_students"]], ["Камалидинова Сумая"])

    def test_not_submitted_counts_as_graded(self):
        lesson = self.lesson(self.gt1, dt.date(2026, 9, 28))
        hw = self.fill(lesson, grades=False)
        HomeworkResult.objects.create(homework=hw, student=self.s1, status=HomeworkResult.Status.NOT_SUBMITTED)
        HomeworkResult.objects.create(homework=hw, student=self.s2, status=HomeworkResult.Status.CHECKED, score=9)
        self.assertEqual(self.row(self.build(), self.aizhan, self.g1)["status"], "ok")

    def test_several_missing_components_is_problem(self):
        self.lesson(self.gt1, dt.date(2026, 9, 28), Lesson.Status.SCHEDULED)
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["status"], "problem")
        self.assertEqual(row["lessons"]["not_closed"], 1)

    def test_small_gap_is_warning_level(self):
        for day in range(1, 9):
            lesson = self.lesson(self.gt1, dt.date(2026, 9, day))
            self.fill(lesson, attendance=day > 2)
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual((row["attendance"]["completed"], row["attendance"]["total"]), (6, 8))
        self.assertEqual(row["attendance"]["level"], "warning")
        self.assertEqual(row["status"], "attention")

    def test_future_lesson_is_upcoming_not_unfilled(self):
        self.lesson(self.gt1, dt.date(2026, 10, 2), Lesson.Status.SCHEDULED)
        # Today's lesson whose slot hasn't ended yet is still upcoming.
        self.lesson(self.gt1, TODAY, Lesson.Status.SCHEDULED, start_time=dt.time(15), end_time=dt.time(16))
        data = self.build()
        row = self.row(data, self.aizhan, self.g1)
        self.assertEqual(row["status"], "upcoming")
        self.assertEqual(row["lessons"]["total"], 0)
        self.assertEqual(row["lessons"]["upcoming"], 2)
        self.assertEqual(data["summary"]["attention_count"], 0)

    def test_today_lesson_past_its_slot_is_due(self):
        self.lesson(self.gt1, TODAY, Lesson.Status.SCHEDULED)  # 10:00–11:00, now is 12:00
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["lessons"]["total"], 1)
        self.assertEqual(row["status"], "problem")

    def test_homework_before_deadline_is_waiting(self):
        lesson = self.lesson(self.gt1, dt.date(2026, 9, 28))
        self.fill(lesson, grades=False, deadline=dt.date(2026, 10, 3))
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["status"], "ok")
        self.assertEqual((row["homework"]["total"], row["grades"]["total"]), (0, 0))

    def test_cancelled_and_not_required(self):
        self.lesson(self.gt1, dt.date(2026, 9, 21), Lesson.Status.CANCELLED)
        lesson = self.lesson(self.gt1, dt.date(2026, 9, 28), homework_not_required=True)
        self.fill(lesson, homework=False)
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["status"], "ok")
        self.assertEqual(row["lessons"]["cancelled"], 1)
        self.assertEqual(row["homework"]["total"], 0)

    def test_no_lessons_is_no_data(self):
        data = self.build()
        self.assertEqual(data["items"], [])
        self.assertEqual(data["summary"]["total_lessons"], 0)
        self.assertIsNone(data["summary"]["attendance_completion"])
        self.lesson(self.gt1, dt.date(2026, 9, 21), Lesson.Status.CANCELLED)
        self.assertEqual(self.row(self.build(), self.aizhan, self.g1)["status"], "no_data")

    def test_group_without_students(self):
        empty = Group.objects.create(name="Empty", course=self.course, start_date=dt.date(2025, 1, 1))
        gt = GroupTeacher.objects.create(group=empty, teacher=self.aizhan, subject=self.python)
        lesson = self.lesson(gt, dt.date(2026, 9, 28))
        Homework.objects.create(lesson=lesson, title="HW")
        row = self.row(self.build(), self.aizhan, empty)
        self.assertEqual(row["status"], "ok")
        self.assertEqual((row["attendance"]["total"], row["grades"]["total"]), (0, 0))

    def test_inactive_student_not_required(self):
        lesson = self.lesson(self.gt1, dt.date(2026, 9, 28))
        self.fill(lesson)
        Student.objects.create(first_name="Дана", group=self.g1, is_active=False, status=Student.Status.WITHDRAWN)
        row = self.row(self.build(), self.aizhan, self.g1)
        self.assertEqual(row["status"], "ok")
        self.assertEqual(row["grades"]["students_missing"], 0)

    def test_substitute_teacher_is_responsible(self):
        lesson = self.lesson(self.gt1, dt.date(2026, 9, 28), teacher=self.almaz)
        data = self.build()
        row = self.row(data, self.almaz, self.g1)
        self.assertEqual(row["lessons"]["total"], 1)
        self.assertFalse(any(i["teacher"]["id"] == self.aizhan.id for i in data["items"]))
        detail = ControlService(
            ControlQuery(filters=ReportFilters.from_query(SEPT, today=TODAY), now_time=NOON)
        ).build_detail(group_id=self.g1.id, teacher_id=self.almaz.id)
        self.assertEqual(detail["lessons"][0]["id"], lesson.id)
        self.assertEqual(detail["lessons"][0]["planned_teacher"]["id"], self.aizhan.id)

    def test_multiple_groups_and_trainers(self):
        g3 = Group.objects.create(name="Next 2", course=self.course, start_date=dt.date(2025, 1, 1))
        gt3 = GroupTeacher.objects.create(group=g3, teacher=self.aizhan, subject=self.python)
        Student.objects.create(first_name="Нурсултан", group=g3)
        self.fill(self.lesson(self.gt1, dt.date(2026, 9, 28)))
        self.fill(self.lesson(gt3, dt.date(2026, 9, 28)), attendance=False)
        self.lesson(self.gt2, dt.date(2026, 9, 28), Lesson.Status.SCHEDULED)

        data = self.build()
        self.assertEqual(len(data["items"]), 3)
        self.assertEqual(self.row(data, self.aizhan, self.g1)["status"], "ok")
        self.assertEqual(self.row(data, self.aizhan, g3)["status"], "attention")
        self.assertEqual(self.row(data, self.almaz, self.g2)["status"], "problem")
        # Problems first.
        self.assertEqual([i["status"] for i in data["items"]], ["problem", "attention", "ok"])
        summary = data["summary"]
        self.assertEqual((summary["total_lessons"], summary["completed_lessons"], summary["not_closed_lessons"]), (3, 2, 1))
        self.assertEqual(summary["attention_count"], 2)
        self.assertEqual(summary["attendance_completion"], 33.3)

    def test_filters(self):
        self.fill(self.lesson(self.gt1, dt.date(2026, 9, 28)))
        self.fill(self.lesson(self.gt1_soft, dt.date(2026, 9, 29)), attendance=False)
        self.lesson(self.gt2, dt.date(2026, 9, 28), Lesson.Status.SCHEDULED)
        self.lesson(self.gt2, dt.date(2026, 8, 28), Lesson.Status.SCHEDULED)  # outside the period

        self.assertEqual(len(self.build()["items"]), 2)
        only_almaz = self.build(teacher=str(self.almaz.id))
        self.assertEqual([i["teacher"]["id"] for i in only_almaz["items"]], [self.almaz.id])
        self.assertEqual(only_almaz["items"][0]["lessons"]["total"], 1)
        self.assertEqual([i["group"]["id"] for i in self.build(group=str(self.g1.id))["items"]], [self.g1.id])
        soft = self.build(subject=str(self.soft.id))
        self.assertEqual(soft["items"][0]["lessons"]["total"], 1)
        self.assertEqual(soft["items"][0]["status"], "attention")
        by_status = self.build(status="problem")
        self.assertEqual([i["teacher"]["id"] for i in by_status["items"]], [self.almaz.id])
        # The summary still describes the whole scope, not just the filtered rows.
        self.assertEqual(by_status["summary"]["total_lessons"], 3)
        self.assertEqual(self.build(period="custom", start_date="2026-08-01", end_date="2026-08-31")["summary"]["total_lessons"], 1)


class ControlApiTests(ControlTestBase):
    def setUp(self):
        super().setUp()
        self.client = APIClient()
        self.l1 = self.lesson(self.gt1, dt.date(2025, 3, 10))
        self.fill(self.l1, attendance=False)
        self.l2 = self.lesson(self.gt2, dt.date(2025, 3, 11))
        self.fill(self.l2)

    def get(self, name, user, params=None, **kwargs):
        self.client.force_authenticate(user)
        return self.client.get(reverse(name, kwargs=kwargs or None), params or {})

    def test_requires_authentication(self):
        self.assertEqual(self.client.get(reverse("control-overview")).status_code, status.HTTP_401_UNAUTHORIZED)

    def test_admin_sees_everything(self):
        response = self.get("control-overview", self.admin, PAST)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["items"]), 2)
        self.assertEqual(response.data["summary"]["total_lessons"], 2)
        self.assertEqual(len(response.data["options"]["teachers"]), 2)

    def test_trainer_sees_only_own_lessons(self):
        response = self.get("control-overview", self.aizhan.user, {**PAST, "teacher": self.almaz.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([i["teacher"]["id"] for i in response.data["items"]], [self.aizhan.id])
        self.assertEqual([t["id"] for t in response.data["options"]["teachers"]], [self.aizhan.id])

    def test_account_without_role_is_forbidden(self):
        user = User.objects.create_user(username="nobody", email="n@okurmen.kg", password="x", first_name="N")
        self.assertEqual(self.get("control-overview", user, PAST).status_code, status.HTTP_403_FORBIDDEN)

    def test_detail(self):
        params = {**PAST, "group": self.g1.id, "teacher": self.aizhan.id}
        response = self.get("control-detail", self.admin, params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["row"]["status"], "attention")
        lesson = response.data["lessons"][0]
        self.assertEqual(lesson["id"], self.l1.id)
        self.assertEqual(lesson["attendance"]["marked"], 0)
        self.assertEqual(len(lesson["attendance"]["missing_students"]), 2)
        self.assertIn("Посещаемость не отмечена", lesson["problems"])
        # A trainer can't open a colleague's row — they always get their own.
        other = self.get("control-detail", self.almaz.user, params)
        self.assertEqual(other.status_code, 404)

    def test_lesson_check_and_scoping(self):
        response = self.get("control-lesson", self.admin, pk=self.l1.id)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["teacher"]["id"], self.aizhan.id)
        self.assertEqual(response.data["status"], "attention")
        self.assertEqual(self.get("control-lesson", self.aizhan.user, pk=self.l1.id).status_code, 200)
        self.assertEqual(self.get("control-lesson", self.almaz.user, pk=self.l1.id).status_code, 404)

    def test_invalid_params(self):
        self.assertEqual(self.get("control-overview", self.admin, {"status": "bogus"}).status_code, 400)
        self.assertEqual(self.get("control-overview", self.admin, {"period": "custom"}).status_code, 400)
        self.assertEqual(self.get("control-detail", self.admin, PAST).status_code, 400)


class ControlTeachersTests(ControlTestBase):
    """The trainer level — the Admin Panel's «Контроль тренеров»."""

    def teachers(self, *, sort="status", status=None, restrict=None, **params):
        filters = ReportFilters.from_query({**SEPT, **params}, today=TODAY)
        service = ControlService(ControlQuery(filters=filters, status=status, now_time=NOON), restrict_teacher=restrict)
        return service.build_teachers(sort=sort)

    def trow(self, data, teacher):
        return next(r for r in data["items"] if r["teacher"]["id"] == teacher.id)

    def test_trainer_statuses(self):
        cases = [
            ({}, "ok"),
            ({"attendance": False}, "attention"),
            ({"grades": False, "unchecked": True}, "attention"),
            ({"grades": False}, "problem"),
            ({"attendance": False, "homework": False}, "problem"),
        ]
        for options, expected in cases:
            with self.subTest(options=options):
                Lesson.objects.all().delete()
                lesson = self.lesson(self.gt1, dt.date(2026, 9, 28))
                unchecked = options.pop("unchecked", False)
                hw = self.fill(lesson, **options)
                if unchecked:
                    for s in (self.s1, self.s2):
                        HomeworkResult.objects.create(homework=hw, student=s, status=HomeworkResult.Status.SUBMITTED, score=7)
                self.assertEqual(self.trow(self.teachers(), self.aizhan)["status"], expected)

    def test_future_and_cancelled_lessons_are_not_unfilled(self):
        self.lesson(self.gt1, dt.date(2026, 10, 2), Lesson.Status.SCHEDULED)
        self.lesson(self.gt2, dt.date(2026, 9, 21), Lesson.Status.CANCELLED)
        data = self.teachers()
        self.assertEqual(self.trow(data, self.aizhan)["status"], "upcoming")
        self.assertEqual(self.trow(data, self.almaz)["status"], "no_data")
        self.assertEqual((data["summary"]["attention"], data["summary"]["problem"]), (0, 0))

    def test_all_trainers_listed_with_groups_and_summary(self):
        idle = make_teacher("idle")
        self.fill(self.lesson(self.gt1, dt.date(2026, 9, 28)))
        self.fill(self.lesson(self.gt1_soft, dt.date(2026, 9, 29)))
        self.lesson(self.gt2, dt.date(2026, 9, 28), Lesson.Status.SCHEDULED)
        data = self.teachers()
        self.assertEqual({r["teacher"]["id"] for r in data["items"]}, {self.aizhan.id, self.almaz.id, idle.id})
        aizhan = self.trow(data, self.aizhan)
        self.assertEqual((aizhan["groups_count"], aizhan["lessons"]["total"]), (1, 2))
        self.assertEqual(self.trow(data, idle)["status"], "no_data")
        self.assertEqual([r["status"] for r in data["items"]], ["problem", "ok", "no_data"])
        summary = data["summary"]
        self.assertEqual((summary["teachers"], summary["ok"], summary["problem"], summary["no_data"]), (3, 1, 1, 1))

    def test_today_summary(self):
        self.fill(self.lesson(self.gt1, TODAY))  # 10:00–11:00, closed and filled
        self.lesson(self.gt2, TODAY, Lesson.Status.SCHEDULED)  # over, nothing filled
        self.lesson(self.gt2, TODAY, Lesson.Status.SCHEDULED, start_time=dt.time(15), end_time=dt.time(16))
        self.lesson(self.gt2, TODAY, Lesson.Status.CANCELLED)
        today = self.teachers(period="last_month")["summary"]["today"]
        self.assertEqual(today, {"total": 3, "closed": 1, "not_filled": 1, "upcoming": 1})

    def test_filters(self):
        self.fill(self.lesson(self.gt1, dt.date(2026, 9, 28)))
        self.fill(self.lesson(self.gt1_soft, dt.date(2026, 9, 29)), attendance=False)
        self.lesson(self.gt2, dt.date(2026, 8, 20), Lesson.Status.SCHEDULED)

        self.assertEqual(self.trow(self.teachers(), self.almaz)["status"], "no_data")  # period
        august = self.teachers(period="custom", start_date="2026-08-01", end_date="2026-08-31")
        self.assertEqual(self.trow(august, self.almaz)["status"], "problem")
        self.assertEqual([r["teacher"]["id"] for r in self.teachers(teacher=str(self.almaz.id))["items"]], [self.almaz.id])
        by_group = self.teachers(group=str(self.g1.id))
        self.assertEqual([r["teacher"]["id"] for r in by_group["items"]], [self.aizhan.id])
        self.assertEqual(self.teachers(subject=str(self.python.id))["items"][0]["status"], "ok")
        self.assertEqual(self.teachers(subject=str(self.soft.id))["items"][0]["status"], "attention")
        self.assertEqual([r["teacher"]["id"] for r in self.teachers(status="attention")["items"]], [self.aizhan.id])

    def test_sorting(self):
        zara = make_teacher("zara")
        g3 = Group.objects.create(name="Z-1", course=self.course, start_date=dt.date(2025, 1, 1))
        gt3 = GroupTeacher.objects.create(group=g3, teacher=zara, subject=self.python)
        Student.objects.create(first_name="Z", group=g3)
        self.fill(self.lesson(self.gt1, dt.date(2026, 9, 28)))
        self.fill(self.lesson(self.gt1_soft, dt.date(2026, 9, 28)), attendance=False)
        self.fill(self.lesson(gt3, dt.date(2026, 9, 28)))
        self.lesson(self.gt2, dt.date(2026, 9, 28), Lesson.Status.SCHEDULED)

        def ids(sort):
            return [r["teacher"]["id"] for r in self.teachers(sort=sort)["items"]]

        self.assertEqual(ids("status"), [self.almaz.id, self.aizhan.id, zara.id])
        self.assertEqual(ids("name"), [self.aizhan.id, self.almaz.id, zara.id])
        self.assertEqual(ids("-grades"), [self.aizhan.id, zara.id, self.almaz.id])  # no scores due -> last
        self.assertEqual(ids("-unfilled"), [self.aizhan.id, self.almaz.id, zara.id])
        self.assertEqual(ids("attendance"), [self.almaz.id, self.aizhan.id, zara.id])
        self.assertEqual(ids("bogus"), ids("status"))

    def test_detail_lists_groups_and_unfilled_lessons(self):
        filled = self.lesson(self.gt1, dt.date(2026, 9, 27))
        self.fill(filled)
        gap = self.lesson(self.gt1_soft, dt.date(2026, 9, 29))
        hw = self.fill(gap, grades=False)
        HomeworkResult.objects.create(homework=hw, student=self.s1, status=HomeworkResult.Status.CHECKED, score=9)
        self.lesson(self.gt1, dt.date(2026, 10, 3), Lesson.Status.SCHEDULED)
        service = ControlService(ControlQuery(filters=ReportFilters.from_query(SEPT, today=TODAY), now_time=NOON))
        detail = service.build_teacher_detail(self.aizhan)
        self.assertEqual(detail["row"]["lessons"]["total"], 2)
        self.assertEqual([g["group"]["id"] for g in detail["groups"]], [self.g1.id])
        self.assertEqual([p["id"] for p in detail["problems"]], [gap.id])
        problem = detail["problems"][0]
        self.assertEqual(problem["teacher"]["id"], self.aizhan.id)
        self.assertEqual(problem["attendance"]["label"], "Заполнено")
        self.assertEqual(problem["grades"]["label"], "Выставлено 1 из 2")
        self.assertEqual([s["name"] for s in problem["grades"]["missing_students"]], ["Азамат уулу Али"])
        self.assertEqual(len(detail["upcoming"]), 1)


class ControlAdminPanelTests(ControlTestBase):
    """/admin/academy/control/... — pages, permissions, exports."""

    def setUp(self):
        super().setUp()
        from django.contrib.auth.models import Permission
        from django.test import Client

        self.web = Client()
        self.l1 = self.lesson(self.gt1, dt.date(2025, 3, 10))
        self.fill(self.l1, grades=False)
        self.l2 = self.lesson(self.gt2, dt.date(2025, 3, 11))
        self.fill(self.l2)
        self.manager = User.objects.create_user(
            username="manager", email="m@okurmen.kg", password="x", first_name="Manager", is_staff=True,
        )
        self.manager.user_permissions.add(Permission.objects.get(codename="view_academymonthlyreport"))
        self.aizhan.user.is_staff = True
        self.aizhan.user.save()

    def get(self, user, name, params=None, **kwargs):
        self.web.force_login(user)
        return self.web.get(reverse(name, kwargs=kwargs or None), params or PAST)

    def test_admin_sees_every_trainer(self):
        response = self.get(self.admin, "admin:academy_control_teachers")
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn("Контроль тренеров", body)
        self.assertIn("Aizhan", body)
        self.assertIn("Almaz", body)
        self.assertEqual(response.context["data"]["summary"]["problem"], 1)
        self.assertIn(reverse("admin:academy_control_teachers"), self.web.get(reverse("admin:index")).content.decode())

    def test_manager_sees_permitted_trainers(self):
        response = self.get(self.manager, "admin:academy_control_teachers")
        self.assertEqual(response.status_code, 200)
        self.assertEqual({r["teacher"]["id"] for r in response.context["data"]["items"]}, {self.aizhan.id, self.almaz.id})

    def test_trainer_sees_only_themselves(self):
        response = self.get(self.aizhan.user, "admin:academy_control_teachers", {**PAST, "teacher": self.almaz.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([r["teacher"]["id"] for r in response.context["data"]["items"]], [self.aizhan.id])
        self.assertNotIn("Almaz", response.content.decode())
        other = self.get(self.aizhan.user, "admin:academy_control_teacher_detail", teacher_id=self.almaz.id)
        self.assertEqual(other.status_code, 404)

    def test_staff_without_access_is_forbidden(self):
        nobody = User.objects.create_user(username="staff", email="s@okurmen.kg", password="x", first_name="S", is_staff=True)
        self.assertEqual(self.get(nobody, "admin:academy_control_teachers").status_code, 403)

    def test_detail_page_links_to_the_lesson(self):
        response = self.get(self.admin, "admin:academy_control_teacher_detail", teacher_id=self.aizhan.id)
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn("Открыть урок", body)
        self.assertIn(reverse("admin:academy_lesson_change", args=[self.l1.id]), body)
        self.assertIn("Баллы не выставлены", body)

    def test_exports(self):
        from io import BytesIO

        from openpyxl import load_workbook

        excel = self.get(self.admin, "admin:academy_control_export_excel")
        self.assertEqual(excel.status_code, 200)
        wb = load_workbook(BytesIO(excel.content))
        self.assertEqual(wb.sheetnames, ["Тренеры", "По группам", "Незаполненные занятия"])
        header = [c.value for c in wb["По группам"][4]]
        self.assertEqual(header, ["Тренер", "Группа", "Период", "Уроков", "Посещаемость", "ДЗ", "Баллы", "Статус",
                                  "Незаполненные данные"])
        cells = [c.value for c in wb["По группам"][5]]
        self.assertEqual(cells[:2], ["Aizhan", "Prog Soft 2"])
        self.assertEqual(cells[7], "Проблема")
        pdf = self.get(self.admin, "admin:academy_control_export_pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertEqual(self.get(self.admin, "admin:academy_control_export_pdf", {"status": "x"}).status_code, 400)
