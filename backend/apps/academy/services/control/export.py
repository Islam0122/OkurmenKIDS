"""«Контроль тренеров» as Excel and PDF — built on the Reports exports'
own helpers (table/sheet writers, fonts, page layout; see
services.reports.excel / services.reports.pdf), not a second export system.

Every figure comes from ControlService for the same filters as the screen.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from apps.academy.services.monthly_report_pdf import BRAND, INK, INK_MUTED, INK_SECONDARY, _BOLD, _REGULAR, _ensure_fonts, _wrap_text
from apps.academy.services.reports import excel as report_excel
from apps.academy.services.reports import pdf as report_pdf

from . import rules

STATUS_FILLS = {
    rules.STATUS_OK: PatternFill("solid", fgColor="E7F3EC"),
    rules.STATUS_ATTENTION: PatternFill("solid", fgColor="FBF1DF"),
    rules.STATUS_PROBLEM: PatternFill("solid", fgColor="FBEAE7"),
}


def _ratio(component: dict) -> str:
    return f"{component['completed']}/{component['total']}" if component["total"] else "—"


def _period(filters) -> str:
    return f"{filters.start:%d.%m.%Y} — {filters.end:%d.%m.%Y}"


def export_rows(group_rows: list[dict], filters) -> list[list]:
    """Тренер, Группа, Период, Уроков, Посещаемость, ДЗ, Баллы, Статус,
    Незаполненные данные — one line per trainer×group."""
    return [
        [
            (row["teacher"] or {}).get("name", "Тренер не назначен"),
            row["group"]["name"],
            _period(filters),
            row["lessons"]["total"],
            _ratio(row["attendance"]),
            _ratio(row["homework"]),
            _ratio(row["grades"]),
            row["status_label"],
            "; ".join(row["issues"]) or "—",
        ]
        for row in group_rows
    ]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _fill_status(ws, first_row: int, col: int, statuses: list[str]) -> None:
    for offset, status in enumerate(statuses):
        if status in STATUS_FILLS:
            ws.cell(row=first_row + offset, column=col).fill = STATUS_FILLS[status]


def build_control_excel(data: dict) -> bytes:
    """`data` — {"filters": ReportFilters, "teachers": [...], "groups": [...], "problems": [...]}."""
    filters = data["filters"]
    report = {"filters": filters}
    wb = Workbook()

    ws = report_excel._sheet(wb, "Тренеры", "Контроль тренеров", report, first=True)
    headers = [
        ("Тренер", 28, None), ("Групп", 8, None), ("Уроков", 9, None), ("Закрыто", 10, None),
        ("Посещаемость", 14, None), ("ДЗ", 10, None), ("Баллы", 10, None), ("Статус", 14, None),
        ("Незаполненные данные", 60, None),
    ]
    teachers = data["teachers"]
    rows = [
        [t["teacher"]["name"], t["groups_count"], t["lessons"]["total"], _ratio(t["lessons"]), _ratio(t["attendance"]),
         _ratio(t["homework"]), _ratio(t["grades"]), t["status_label"], "; ".join(t["issues"]) or "—"]
        for t in teachers
    ]
    report_excel._write_table(ws, 4, headers, rows)
    _fill_status(ws, 5, 8, [t["status"] for t in teachers])
    ws.freeze_panes = "B5"

    ws = report_excel._sheet(wb, "По группам", "Контроль по группам", report)
    headers = [
        ("Тренер", 26, None), ("Группа", 26, None), ("Период", 24, None), ("Уроков", 9, None),
        ("Посещаемость", 14, None), ("ДЗ", 10, None), ("Баллы", 10, None), ("Статус", 14, None),
        ("Незаполненные данные", 60, None),
    ]
    report_excel._write_table(ws, 4, headers, export_rows(data["groups"], filters))
    _fill_status(ws, 5, 8, [g["status"] for g in data["groups"]])
    ws.freeze_panes = "C5"

    ws = report_excel._sheet(wb, "Незаполненные занятия", "Что осталось заполнить", report)
    headers = [
        ("Дата", 12, report_excel.DATE), ("Тренер", 26, None), ("Группа", 24, None), ("Предмет", 18, None),
        ("Посещаемость", 22, None), ("ДЗ", 28, None), ("Баллы", 26, None), ("Статус", 14, None),
        ("Что не заполнено", 50, None),
    ]
    problems = data["problems"]
    rows = [
        [p["date"], (p["teacher"] or {}).get("name", "—"), p["group"]["name"], (p["subject"] or {}).get("name", "—"),
         p["attendance"]["label"], p["homework"]["label"], p["grades"]["label"], p["status_label"],
         "; ".join(p["problems"])]
        for p in problems
    ]
    report_excel._write_table(ws, 4, headers, rows)
    _fill_status(ws, 5, 8, [p["status"] for p in problems])
    ws.freeze_panes = "B5"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def excel_filename(filters) -> str:
    return f"okurmenkids-control-{filters.start:%Y%m%d}-{filters.end:%Y%m%d}.xlsx"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def build_control_pdf(data: dict) -> bytes:
    _ensure_fonts()
    filters = data["filters"]
    summary = data["summary"]
    buffer = io.BytesIO()
    doc = report_pdf._Doc(buffer, f"OkurmenKIDS — контроль тренеров {_period(filters)}")
    report_pdf._page_title(doc, "OKURMENKIDS · КОНТРОЛЬ", "Контроль тренеров",
                           f"{filters.period_label} · {_period(filters)}")
    report_pdf._stat_cards(doc, [
        (str(summary["teachers"]), "Тренеров", None),
        (str(summary["ok"]), "Работают нормально", None),
        (str(summary["attention"]), "Требуют внимания", None),
        (str(summary["problem"]), "Есть проблемы", None),
    ])

    report_pdf._section_label(doc, "Тренеры")
    columns = [
        {"title": "Тренер", "width": 0.26}, {"title": "Групп", "width": 0.08, "align": "right"},
        {"title": "Уроки", "width": 0.1, "align": "right"}, {"title": "Посещаемость", "width": 0.14, "align": "right"},
        {"title": "ДЗ", "width": 0.1, "align": "right"}, {"title": "Баллы", "width": 0.1, "align": "right"},
        {"title": "Статус", "width": 0.22},
    ]
    report_pdf._table(doc, columns, [
        [t["teacher"]["name"], t["groups_count"], _ratio(t["lessons"]), _ratio(t["attendance"]),
         _ratio(t["homework"]), _ratio(t["grades"]), t["status_label"]]
        for t in data["teachers"]
    ], empty="Нет тренеров для выбранных фильтров.")

    report_pdf._section_label(doc, "По группам")
    columns = [
        {"title": "Тренер", "width": 0.2}, {"title": "Группа", "width": 0.2},
        {"title": "Уроков", "width": 0.09, "align": "right"}, {"title": "Посещ.", "width": 0.1, "align": "right"},
        {"title": "ДЗ", "width": 0.09, "align": "right"}, {"title": "Баллы", "width": 0.09, "align": "right"},
        {"title": "Статус", "width": 0.23},
    ]
    groups = data["groups"]
    report_pdf._table(doc, columns, [
        [row[0], row[1], row[3], row[4], row[5], row[6], row[7]] for row in export_rows(groups, filters)
    ], empty="Нет занятий за выбранный период.")

    with_issues = [g for g in groups if g["issues"]]
    if with_issues:
        report_pdf._section_label(doc, "Незаполненные данные")
        width = report_pdf.CONTENT_W
        for row in with_issues:
            lines = _wrap_text("; ".join(row["issues"]), _REGULAR, 8, width - 12)
            doc.ensure_space(16 + 11 * len(lines))
            doc.text(report_pdf.MARGIN, doc.y,
                     f"{(row['teacher'] or {}).get('name', 'Тренер не назначен')} · {row['group']['name']} · {row['status_label']}",
                     font=_BOLD, size=8.6, color=INK)
            doc.y -= 12
            for line in lines:
                doc.text(report_pdf.MARGIN + 10, doc.y, line, size=8, color=INK_SECONDARY)
                doc.y -= 11
            doc.y -= 5
    else:
        doc.ensure_space(20)
        doc.text(report_pdf.MARGIN, doc.y, "Все обязательные данные заполнены.", size=9, color=BRAND)
        doc.y -= 14
    doc.text(report_pdf.MARGIN, max(doc.y - 6, report_pdf.MARGIN + report_pdf.FOOTER_H),
             "Предстоящие и отменённые занятия не считаются незаполненными.", size=7.3, color=INK_MUTED)
    doc.c.save()
    return buffer.getvalue()


def pdf_filename(filters) -> str:
    return f"okurmenkids-control-{filters.start:%Y%m%d}-{filters.end:%Y%m%d}.pdf"
