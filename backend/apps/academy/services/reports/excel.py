"""Export Excel — the Reports section as a multi-sheet .xlsx workbook.

Sheets: Overview, Groups, Teachers, Students, Attendance, Homework, KPI.
Every figure comes from services.reports.service for the same filters the
screens and the PDF use. Percentages are written as real numbers with a
percent number format (so they sort/filter/sum in Excel); a figure with no
underlying data is left blank rather than written as a fabricated 0%.
"""
from __future__ import annotations

import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .kpi import COMPONENT_LABELS, LEVEL_LABELS
from .service import build_all_student_rows

BRAND = "2F8F5B"
BRAND_DARK = "1F5F3C"
LEVEL_FILLS = {
    "good": PatternFill("solid", fgColor="E7F3EC"),
    "warning": PatternFill("solid", fgColor="FBF1DF"),
    "bad": PatternFill("solid", fgColor="FBEAE7"),
}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor=BRAND)
_TITLE_FONT = Font(bold=True, size=14, color=BRAND_DARK)
_BOLD = Font(bold=True)
_THIN = Side(style="thin", color="E1E7E2")
_BORDER = Border(bottom=_THIN)
PCT = "0.0%"
DATE = "DD.MM.YYYY"


def _pct(value):
    return None if value is None else round(value / 100, 4)


def _write_table(ws, start_row: int, headers: list[tuple[str, int, str | None]], rows: list[list],
                 *, level_col: int | None = None, levels: list[str] | None = None) -> int:
    """headers: (title, width, number_format). Returns the next free row."""
    for col, (title, width, _fmt) in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)
    ws.row_dimensions[start_row].height = 30
    for r, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = _BORDER
            fmt = headers[col - 1][2]
            if fmt:
                cell.number_format = fmt
        if level_col and levels and levels[r - start_row - 1] in LEVEL_FILLS:
            ws.cell(row=r, column=level_col).fill = LEVEL_FILLS[levels[r - start_row - 1]]
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
    else:
        ws.cell(row=start_row + 1, column=1, value="Нет данных для выбранных фильтров.").font = Font(
            italic=True, color="89938D"
        )
    return start_row + max(len(rows), 1) + 2


def _sheet(wb, title: str, heading: str, report: dict, *, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    filters = report["filters"]
    ws["A1"] = heading
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Период: {filters.start:%d.%m.%Y} — {filters.end:%d.%m.%Y} ({filters.period_label})"
    ws["A2"].font = Font(color="68736C")
    return ws


def _overview_sheet(wb, report: dict, filter_labels: str) -> None:
    ov = report["overview"]
    ws = _sheet(wb, "Overview", "OKURMENKIDS — Academy Performance Report", report, first=True)
    ws["A3"] = f"Фильтры: {filter_labels}"
    ws["A4"] = f"Сформирован: {timezone.localtime():%d.%m.%Y %H:%M}"
    for cell in ("A3", "A4"):
        ws[cell].font = Font(color="89938D")
    st = ov["students"]
    rows = [
        ["Всего студентов", st["total"], None],
        ["Активные студенты", st["active"], None],
        ["Ушли за период", st["left"], None],
        ["Новые за период", st["new"], None],
        ["Вернулись за период", st["returned"], None],
        ["На паузе", st["paused"], None],
        ["Удержание", _pct(st["retention_rate"]), PCT],
        ["Тренеры", ov["teachers"]["total"], None],
        ["Группы", ov["groups"]["total"], None],
        ["Групп без тренера", ov["groups"]["without_teacher"], None],
        ["Занятий в периоде", ov["lessons"]["total"], None],
        ["Проведено занятий", ov["lessons"]["held"], None],
        ["Посещаемость", _pct(ov["attendance"]["rate"]), PCT],
        ["Выполнение ДЗ", _pct(ov["homework"]["completion_rate"]), PCT],
        ["Attendance KPI", _pct(ov["kpi"]["attendance"]), PCT],
        ["Homework KPI", _pct(ov["kpi"]["homework"]), PCT],
        ["Activity KPI", _pct(ov["kpi"]["activity"]), PCT],
        ["Progress KPI", _pct(ov["kpi"]["progress"]), PCT],
        ["Overall KPI", _pct(ov["kpi"]["overall"]), PCT],
    ]
    headers = [("Показатель", 34, None), ("Значение", 16, None)]
    _write_table(ws, 6, headers, [r[:2] for r in rows])
    for index, (_label, _value, fmt) in enumerate(rows, start=7):
        if fmt:
            ws.cell(row=index, column=2).number_format = fmt
    ws.cell(row=6 + len(rows), column=1).font = _BOLD
    ws.cell(row=6 + len(rows), column=2).font = _BOLD
    ws.auto_filter.ref = None
    ws.freeze_panes = "A7"


def _groups_sheet(wb, report: dict) -> None:
    ws = _sheet(wb, "Groups", "Результаты групп", report)
    headers = [
        ("Группа", 30, None), ("Программа", 24, None), ("Тренер(ы)", 28, None), ("Предметы", 22, None),
        ("Статус", 13, None), ("Дата начала", 13, DATE), ("Всего", 9, None), ("Активные", 10, None),
        ("Ушли", 8, None), ("Новые", 8, None), ("Вернулись", 10, None), ("Занятий", 9, None),
        ("Проведено", 10, None), ("Посещаемость", 13, PCT), ("ДЗ", 10, PCT), ("Активность", 12, PCT),
        ("Прогресс", 11, PCT), ("KPI", 10, PCT), ("Статус", 16, None),
    ]
    rows = [
        [g["name"], g["program"], g["teacher_names"], ", ".join(g["subjects"]), g["status_display"], g["start_date"],
         g["students"]["total"], g["students"]["active"], g["students"]["left"], g["students"]["new"],
         g["students"]["returned"], g["lessons"]["total"], g["lessons"]["held"], _pct(g["attendance_rate"]),
         _pct(g["homework_rate"]), _pct(g["activity_rate"]), _pct(g["progress_rate"]), _pct(g["kpi"]),
         LEVEL_LABELS[g["kpi_level"]]]
        for g in report["groups"]
    ]
    _write_table(ws, 4, headers, rows, level_col=18, levels=[g["kpi_level"] for g in report["groups"]])
    ws.freeze_panes = "B5"


def _teachers_sheet(wb, report: dict) -> None:
    ws = _sheet(wb, "Teachers", "Результаты тренеров", report)
    headers = [
        ("Тренер", 26, None), ("Должность", 22, None), ("Групп", 8, None), ("Группы", 40, None),
        ("Предметы", 26, None), ("Студентов", 10, None), ("Активные", 10, None), ("Ушли", 8, None),
        ("Занятий", 9, None), ("Проведено", 10, None), ("Посещаемость", 13, PCT), ("ДЗ", 10, PCT),
        ("Активность", 12, PCT), ("Прогресс", 11, PCT), ("KPI", 10, PCT), ("Статус", 16, None),
    ]
    rows = [
        [t["name"], t["position"], t["groups_count"], ", ".join(g["name"] for g in t["groups"]) or "Нет групп",
         ", ".join(t["subjects"]), t["students"]["total"], t["students"]["active"], t["students"]["left"],
         t["lessons"]["total"], t["lessons"]["held"], _pct(t["attendance_rate"]), _pct(t["homework_rate"]),
         _pct(t["activity_rate"]), _pct(t["progress_rate"]), _pct(t["kpi"]), LEVEL_LABELS[t["kpi_level"]]]
        for t in report["teachers"]
    ]
    _write_table(ws, 4, headers, rows, level_col=15, levels=[t["kpi_level"] for t in report["teachers"]])
    ws.freeze_panes = "B5"


def _students_sheet(wb, report: dict) -> None:
    ws = _sheet(wb, "Students", "Студенты", report)
    students = build_all_student_rows(report["filters"], report["group_ids"])
    headers = [
        ("Студент", 28, None), ("Группа", 30, None), ("Статус", 18, None), ("Отметок", 9, None),
        ("Посетил", 9, None), ("Посещаемость", 13, PCT), ("Результатов ДЗ", 12, None), ("Сдано", 8, None),
        ("ДЗ", 10, PCT), ("Средний балл", 12, "0.0"), ("Прогресс", 11, PCT),
    ]
    rows = [
        [s["name"], s["group"], s["status_display"], s["attendance_total"], s["attended"], _pct(s["attendance_rate"]),
         s["homework_results"], s["homework_submitted"], _pct(s["homework_rate"]), s["average_score"],
         _pct(s["progress_rate"])]
        for s in students
    ]
    _write_table(ws, 4, headers, rows)
    ws.freeze_panes = "B5"


def _attendance_sheet(wb, report: dict) -> None:
    ws = _sheet(wb, "Attendance", "Посещаемость по группам", report)
    headers = [
        ("Группа", 30, None), ("Тренер(ы)", 28, None), ("Отметок", 10, None), ("Присутствовали", 14, None),
        ("Опоздали", 10, None), ("Отсутствовали", 13, None), ("Уваж. причина", 13, None),
        ("Посещаемость", 13, PCT), ("Пропуски", 11, PCT),
    ]
    rows = [
        [g["name"], g["teacher_names"], a["total"], a["present"], a["late"], a["absent"], a["excused"],
         _pct(a["rate"]), _pct(a["absence_rate"])]
        for g in report["groups"]
        for a in [g["attendance"]]
    ]
    total = report["overview"]["attendance"]
    if rows:
        rows.append(["ИТОГО", "", total["total"], total["present"], total["late"], total["absent"], total["excused"],
                     _pct(total["rate"]), _pct(total["absence_rate"])])
    _write_table(ws, 4, headers, rows)
    if rows:
        for col in range(1, len(headers) + 1):
            ws.cell(row=4 + len(rows), column=col).font = _BOLD
    ws.freeze_panes = "B5"


def _homework_sheet(wb, report: dict) -> None:
    ws = _sheet(wb, "Homework", "Домашние задания по группам", report)
    headers = [
        ("Группа", 30, None), ("Тренер(ы)", 28, None), ("Выдано ДЗ", 11, None), ("Результатов", 12, None),
        ("Сдано", 8, None), ("Не сдано", 10, None), ("Проверено", 11, None), ("С опозданием", 13, None),
        ("Средний балл", 12, "0.0"), ("Выполнение", 12, PCT),
    ]
    rows = [
        [g["name"], g["teacher_names"], h["assigned"], h["results"], h["submitted"], h["not_submitted"],
         h["checked"], h["late"], h["average_score"], _pct(h["completion_rate"])]
        for g in report["groups"]
        for h in [g["homework"]]
    ]
    total = report["overview"]["homework"]
    if rows:
        rows.append(["ИТОГО", "", total["assigned"], total["results"], total["submitted"], total["not_submitted"],
                     total["checked"], total["late"], total["average_score"], _pct(total["completion_rate"])])
    _write_table(ws, 4, headers, rows)
    if rows:
        for col in range(1, len(headers) + 1):
            ws.cell(row=4 + len(rows), column=col).font = _BOLD
    ws.freeze_panes = "B5"


def _kpi_sheet(wb, report: dict) -> None:
    ov = report["overview"]
    ws = _sheet(wb, "KPI", "KPI", report)
    ws["A4"] = "Формула: Overall KPI = Σ (компонент × вес) / Σ весов компонентов, по которым есть данные."
    ws["A4"].font = Font(italic=True, color="68736C")
    next_row = _write_table(
        ws, 6, [("Компонент", 36, None), ("Вес", 10, "0.0%"), ("KPI академии", 14, PCT)],
        [[w["label"], w["weight"] / 100, _pct(ov["kpi"][w["key"]])] for w in ov["kpi_weights"]]
        + [["Overall KPI", 1, _pct(ov["kpi"]["overall"])]],
    )
    ws.auto_filter.ref = None
    headers = [
        ("Группа / тренер", 36, None), ("Тип", 10, None), (COMPONENT_LABELS["attendance"], 14, PCT),
        (COMPONENT_LABELS["homework"], 14, PCT), (COMPONENT_LABELS["activity"], 16, PCT),
        (COMPONENT_LABELS["progress"], 16, PCT), ("Overall KPI", 12, PCT), ("Статус", 16, None),
    ]
    rows, levels = [], []
    for kind, items in (("Группа", report["groups"]), ("Тренер", report["teachers"])):
        for item in items:
            rows.append([item["name"], kind, _pct(item["attendance_rate"]), _pct(item["homework_rate"]),
                         _pct(item["activity_rate"]), _pct(item["progress_rate"]), _pct(item["kpi"]),
                         LEVEL_LABELS[item["kpi_level"]]])
            levels.append(item["kpi_level"])
    _write_table(ws, next_row, headers, rows, level_col=7, levels=levels)


def build_reports_excel(report: dict, filter_labels: str) -> bytes:
    wb = Workbook()
    _overview_sheet(wb, report, filter_labels)
    _groups_sheet(wb, report)
    _teachers_sheet(wb, report)
    _students_sheet(wb, report)
    _attendance_sheet(wb, report)
    _homework_sheet(wb, report)
    _kpi_sheet(wb, report)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def excel_filename(filters) -> str:
    return f"okurmenkids-report-{filters.start:%Y%m%d}-{filters.end:%Y%m%d}.xlsx"
