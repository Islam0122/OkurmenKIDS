"""Downloadable Excel templates for the "Импорт данных" admin page.

Every template's "Данные" sheet header row uses the *exact* column keys
the corresponding importer reads (`apps.academy.services.import_export`,
`apps.users.import_export.teachers`, `apps.academy.services.
group_import_export`) — the spec is explicit that "template must match
importer expectations exactly", so nothing here decorates a header cell's
text (no "*", no "(required)" suffix): required columns are instead
flagged with a fill colour, and every column gets a cell comment with its
own description/example, plus a full legend on the second sheet.
"""
from __future__ import annotations

import io

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL_REQUIRED = PatternFill("solid", fgColor="2F8F5B")
_HEADER_FILL_OPTIONAL = PatternFill("solid", fgColor="6B9080")
_EXAMPLE_FONT = Font(italic=True, color="6B7280")
_TITLE_FONT = Font(bold=True, size=13, color="1F2937")
_SECTION_FONT = Font(bold=True, size=11, color="2F8F5B")
_WRAP = Alignment(wrap_text=True, vertical="top")


class UnknownTemplateType(Exception):
    """Raised for an unrecognized ``type`` query parameter."""


class TemplateColumn:
    __slots__ = ("key", "required", "example", "note")

    def __init__(self, key: str, required: bool, example: str, note: str):
        self.key = key
        self.required = required
        self.example = example
        self.note = note


STUDENT_COLUMNS = [
    TemplateColumn(
        "id", False, "",
        "Внутренний номер студента. Оставьте пустым, чтобы создать нового студента. "
        "Укажите id существующего студента (см. список студентов), чтобы обновить его данные.",
    ),
    TemplateColumn("first_name", True, "Айгерим", "Имя студента."),
    TemplateColumn("last_name", False, "Нурланова", "Фамилия студента."),
    TemplateColumn(
        "phone", False, "+996700123456",
        "Телефон студента. Разрешены цифры, пробелы, +, -, (), от 5 до 30 символов.",
    ),
    TemplateColumn(
        "parent_phone", False, "+996700654321",
        "Телефон родителя, тот же формат, что и phone.",
    ),
    TemplateColumn(
        "group", False, "Python-01",
        "Точное название уже существующей группы. Группа не создаётся автоматически — "
        "если группа с таким названием не найдена, строка будет отклонена с ошибкой.",
    ),
    TemplateColumn(
        "is_active", False, "да",
        "Активен ли студент: да/нет, true/false, yes/no, 1/0. Пусто — считается «да».",
    ),
]

TEACHER_COLUMNS = [
    TemplateColumn("username", True, "aigerim.n", "Логин тренера для входа в систему. Должен быть уникальным."),
    TemplateColumn("email", True, "aigerim@okurmenkids.kg", "Электронная почта тренера. Должна быть уникальной."),
    TemplateColumn("first_name", True, "Айгерим", "Имя тренера."),
    TemplateColumn("last_name", False, "Нурланова", "Фамилия тренера."),
    TemplateColumn("phone", False, "+996700123456", "Телефон тренера, разрешены цифры, пробелы, +, -, ()."),
    TemplateColumn("position", False, "Тренер", "Должность. Пусто — подставится «Тренер»."),
    TemplateColumn("experience_years", False, "3", "Стаж работы в годах, целое число от 0 до 60."),
    TemplateColumn("bio", False, "Опыт преподавания программирования детям.", "Краткая биография."),
    TemplateColumn("hire_date", False, "2024-09-01", "Дата приёма на работу в формате ГГГГ-ММ-ДД."),
    TemplateColumn(
        "subjects", False, "Python, Scratch",
        "Список предметов через запятую. Каждый предмет должен уже существовать в системе.",
    ),
    TemplateColumn("is_active", False, "да", "Активен ли аккаунт: да/нет, true/false, 1/0. Пусто — «да»."),
    TemplateColumn(
        "is_verified", False, "",
        "Подтверждён ли аккаунт администратором: да/нет. Пусто — не меняется "
        "(при создании — «нет»).",
    ),
    TemplateColumn(
        "password", False, "",
        "Пароль для нового аккаунта. Пусто — сгенерируется автоматически. "
        "Не используется при обновлении существующего тренера и никогда не выгружается при экспорте.",
    ),
]

GROUP_COLUMNS = [
    TemplateColumn(
        "name", True, "Python-01",
        "Название группы. Уникально — используется, чтобы найти существующую группу и обновить её "
        "(иначе создаётся новая).",
    ),
    TemplateColumn(
        "course", True, "Python для детей",
        "Точное название уже существующего курса. Курс не создаётся автоматически.",
    ),
    TemplateColumn(
        "status", False, "active",
        "Статус группы: active/paused/completed/cancelled (или по-русски: Активна/Приостановлена/"
        "Завершена/Отменена). Пусто — «active».",
    ),
    TemplateColumn("start_date", True, "2025-09-01", "Дата начала группы в формате ГГГГ-ММ-ДД."),
    TemplateColumn("end_date", False, "2026-05-31", "Дата окончания группы в формате ГГГГ-ММ-ДД (необязательно)."),
    TemplateColumn("max_students", False, "12", "Максимальное количество студентов, целое число от 1."),
    TemplateColumn("description", False, "Группа выходного дня.", "Свободное описание группы."),
]

_TEMPLATES = {
    "student": ("Студенты", "students_template", STUDENT_COLUMNS),
    "teacher": ("Тренеры", "teachers_template", TEACHER_COLUMNS),
    "group": ("Группы", "groups_template", GROUP_COLUMNS),
}

_STRATEGY_NOTE = {
    "student": (
        "Стратегия импорта: создание или обновление. Ключ обновления — колонка id. Пустой id — "
        "создаётся новый студент; заполненный id существующего студента — обновляются его данные. "
        "Импорт выполняется в одной транзакции: если хотя бы одна строка не проходит проверку, "
        "не сохраняется ни одна запись."
    ),
    "teacher": (
        "Стратегия импорта: создание или обновление. Ключ обновления — username или email "
        "(оба уникальны). Если пользователь с таким username или email уже существует, его данные "
        "обновляются; иначе создаётся новый тренер. Импорт выполняется в одной транзакции."
    ),
    "group": (
        "Стратегия импорта: создание или обновление. Ключ обновления — name (название группы "
        "уникально). Если группа с таким названием уже существует, её данные обновляются; иначе "
        "создаётся новая группа. Импорт выполняется в одной транзакции."
    ),
}


def _autosize_columns(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 60)


def _build_data_sheet(workbook: Workbook, columns: list[TemplateColumn]) -> None:
    sheet = workbook.active
    sheet.title = "Данные"

    for col_index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_index, value=column.key)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL_REQUIRED if column.required else _HEADER_FILL_OPTIONAL
        cell.alignment = Alignment(vertical="center")
        marker = "Обязательное поле. " if column.required else "Необязательное поле. "
        example_note = f" Пример: {column.example}" if column.example else ""
        cell.comment = Comment(f"{marker}{column.note}{example_note}", "OkurmenKIDS")

        example_cell = sheet.cell(row=2, column=col_index, value=column.example)
        example_cell.font = _EXAMPLE_FONT

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    _autosize_columns(sheet, [max(len(c.key), len(c.example)) for c in columns])


def _build_instructions_sheet(workbook: Workbook, entity: str, columns: list[TemplateColumn]) -> None:
    sheet = workbook.create_sheet("Инструкция")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 70

    title, _, _ = _TEMPLATES[entity]
    row = 1
    sheet.cell(row=row, column=1, value=f"Шаблон импорта — {title}").font = _TITLE_FONT
    row += 2

    sheet.cell(row=row, column=1, value="Как использовать").font = _SECTION_FONT
    row += 1
    for line in (
        "1. Заполните лист «Данные», не меняя названия колонок в первой строке.",
        "2. Строка 2 листа «Данные» — пример заполнения; замените её своими значениями или удалите.",
        "3. Сохраните файл в формате .xlsx и загрузите на странице «Импорт данных».",
        "4. Сначала нажмите «Предпросмотр», чтобы увидеть ошибки без сохранения данных.",
        _STRATEGY_NOTE[entity],
    ):
        cell = sheet.cell(row=row, column=1, value=line)
        cell.alignment = _WRAP
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
    row += 1

    sheet.cell(row=row, column=1, value="Описание колонок").font = _SECTION_FONT
    row += 1
    headers = ["Колонка", "Обязательна", "Пример", "Описание"]
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E5E7EB")
    row += 1

    for column in columns:
        sheet.cell(row=row, column=1, value=column.key)
        required_cell = sheet.cell(row=row, column=2, value="Да" if column.required else "Нет")
        required_cell.font = Font(bold=column.required, color="C7402E" if column.required else "6B7280")
        sheet.cell(row=row, column=3, value=column.example or "—")
        note_cell = sheet.cell(row=row, column=4, value=column.note)
        note_cell.alignment = _WRAP
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Общие форматы").font = _SECTION_FONT
    row += 1
    for line in (
        "Даты — формат ГГГГ-ММ-ДД (например, 2025-09-01).",
        "Логические значения — да/нет, true/false, yes/no, 1/0 (регистр не важен).",
        "Телефон — цифры, пробелы, +, -, (), от 5 до 30 символов.",
        "Кодировка файла — UTF-8; поддерживаются форматы .xlsx и .csv.",
    ):
        cell = sheet.cell(row=row, column=1, value=line)
        cell.alignment = _WRAP
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1


def build_template(entity: str) -> HttpResponse:
    if entity not in _TEMPLATES:
        raise UnknownTemplateType(f'Неизвестный тип шаблона «{entity}».')

    _, base_filename, columns = _TEMPLATES[entity]
    workbook = Workbook()
    _build_data_sheet(workbook, columns)
    _build_instructions_sheet(workbook, entity, columns)

    buffer = io.BytesIO()
    workbook.save(buffer)
    timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{base_filename}_{timestamp}.xlsx"'
    return response
