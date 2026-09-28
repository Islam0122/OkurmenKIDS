"""Payment audit on the seed_scholarships mock data: the money flow end to
end, checked against plain ORM queries rather than against the UI.

Mock August period: 20 × 1 500 = 30 000, 7 paid (10 500), 13 pending (19 500).
"""
from __future__ import annotations

import io
import threading
from decimal import Decimal
from unittest import mock, skipUnless

from django.core.management import call_command
from django.db import connection
from django.db.models import Sum
from django.db.models.query import QuerySet
from django.test import TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.scholarships.demo import mock_payments as seed
from apps.scholarships.models import PaymentStatus, ScholarshipAward, ScholarshipPeriod, ScholarshipRunLog
from apps.scholarships.services.payments import pay_awards
from apps.users.models import User

from .base import ScholarshipFixture, make_admin, make_user

D = Decimal
PAID, UNPAID = PaymentStatus.PAID, PaymentStatus.UNPAID


def orm_money(awards) -> tuple:
    """(total, paid, pending, remaining) straight from the database."""
    total = awards.aggregate(s=Sum("amount"))["s"] or D(0)
    paid = awards.filter(payment_status=PAID).aggregate(s=Sum("paid_amount"))["s"] or D(0)
    pending = awards.filter(payment_status=UNPAID).aggregate(s=Sum("amount"))["s"] or D(0)
    return total, paid, pending, total - paid


class MockMoneyFixture(ScholarshipFixture):
    def setUp(self):
        super().setUp()
        call_command("seed_scholarships", stdout=io.StringIO())
        self.periods = {f"{p.period_start:%d.%m}": p for p in ScholarshipPeriod.objects.filter(seed.mock_period_q())}
        self.august = self.periods["01.08"]
        self.web = self.client_class()
        self.web.force_login(self.admin)
        self.pay_url = reverse("admin:scholarships_pay")
        self.list_url = reverse("admin:scholarships_scholarshipaward_changelist")

    def unpaid(self, period=None):
        return (period or self.august).awards.filter(payment_status=UNPAID).order_by("rank")

    def pay(self, ids, method="cash", **extra):
        return self.web.post(self.pay_url, {"award": ids, "payment_method": method, "next": self.list_url, **extra})


class PaymentAuditTests(MockMoneyFixture):
    def test_single_payment(self):
        award = self.unpaid().first()
        other = make_admin("other_admin")
        self.assertEqual(orm_money(self.august.awards), (D(30000), D(10500), D(19500), D(19500)))
        # Everything but the id and the method is decided by the server.
        self.pay(award.pk, "bank", amount="150000", paid_amount="150000", paid_by=other.pk, paid_at="2020-01-01")
        award.refresh_from_db()
        self.assertEqual((award.payment_status, award.amount, award.paid_amount), (PAID, D(1500), D(1500)))
        self.assertEqual((award.paid_by, award.payment_method), (self.admin, "bank"))
        self.assertLess(abs((timezone.now() - award.paid_at).total_seconds()), 60)
        self.assertEqual(orm_money(self.august.awards), (D(30000), D(12000), D(18000), D(18000)))
        log = ScholarshipRunLog.objects.filter(action=ScholarshipRunLog.Action.PAYMENT).latest("pk")
        self.assertEqual(log.triggered_by, self.admin)
        self.assertIn(award.evaluation.student_name, log.message)

    def test_duplicate_payment(self):
        award = self.unpaid().first()
        self.pay(award.pk)
        award.refresh_from_db()
        first = (award.paid_at, award.payment_method)
        money = orm_money(self.august.awards)
        # Straight POSTs, the same id several times, another method.
        self.pay([award.pk, award.pk, award.pk], "bank")
        response = self.web.post(self.pay_url, {"award": award.pk, "payment_method": "other", "next": self.list_url},
                                 follow=True)
        self.assertContains(response, "повторно не выплачены")
        award.refresh_from_db()
        self.assertEqual((award.paid_at, award.payment_method), first)
        self.assertEqual(orm_money(self.august.awards), money)
        self.assertEqual(pay_awards([award.pk], method="cash", user=self.admin).paid, [])

    def test_mass_payment(self):
        ids = list(self.unpaid().values_list("pk", flat=True)[:5])
        paid_before = self.august.awards.filter(payment_status=PAID).count()
        _, paid, pending, _ = orm_money(self.august.awards)
        self.pay(ids)
        self.assertEqual(self.august.awards.filter(payment_status=PAID).count(), paid_before + 5)
        _, paid_after, pending_after, _ = orm_money(self.august.awards)
        self.assertEqual((paid_after - paid, pending - pending_after), (D(7500), D(7500)))
        self.assertEqual(set(ScholarshipAward.objects.filter(pk__in=ids).values_list("paid_amount", flat=True)), {D(1500)})

    def test_mass_payment_is_atomic(self):
        ids = list(self.unpaid(self.periods["01.07"]).values_list("pk", flat=True)[:10])
        original = QuerySet.update

        def failing_update(qs, **fields):
            if "payment_status" in fields:
                raise RuntimeError("database error")
            return original(qs, **fields)

        with mock.patch.object(QuerySet, "update", failing_update), self.assertRaises(RuntimeError):
            pay_awards(ids, method="cash", user=self.admin)
        self.assertFalse(ScholarshipAward.objects.filter(pk__in=ids, payment_status=PAID).exists())

    def test_payment_totals(self):
        for key, period in self.periods.items():
            awards = period.awards
            total, paid, pending, remaining = orm_money(awards)
            self.assertEqual(paid + pending, total, key)
            self.assertEqual(
                awards.filter(payment_status=PAID).count() + awards.filter(payment_status=UNPAID).count(),
                awards.count(), key,
            )
            summary = self.web.get(self.list_url, {"period__id__exact": period.pk}).context["summary"]
            self.assertEqual((summary["accrued"] or 0, summary["paid_sum"] or 0), (total, paid), key)

    def test_pending_totals(self):
        for key, period in self.periods.items():
            _, _, pending, _ = orm_money(period.awards)
            summary = self.web.get(self.list_url, {"period__id__exact": period.pk}).context["summary"]
            self.assertEqual(summary["remaining"] or 0, pending, key)
            self.assertEqual(summary["unpaid"], period.awards.filter(payment_status=UNPAID).count(), key)

    def test_remaining_amount(self):
        total, paid, pending, remaining = orm_money(ScholarshipAward.objects.filter(period__in=self.periods.values()))
        self.assertEqual((total, paid, pending, remaining), (D(327500), D(158000), D(169500), D(169500)))
        self.pay(list(self.unpaid().values_list("pk", flat=True)[:2]))
        total2, paid2, pending2, remaining2 = orm_money(ScholarshipAward.objects.filter(period__in=self.periods.values()))
        self.assertEqual((total2, paid2, remaining2), (total, paid + 3000, remaining - 3000))
        self.assertEqual(pending2, remaining2)

    def test_accounting_report(self):
        for key, period in self.periods.items():
            report = self.web.get(reverse("admin:scholarships_report"), {"period": period.pk}).context["report"]
            t, awards = report.rows_totals, period.awards
            total, paid, pending, _ = orm_money(awards)
            self.assertEqual(
                (t.students, t.paid, t.unpaid, t.amount, t.paid_amount, t.remaining),
                (awards.values("student").distinct().count(), awards.filter(payment_status=PAID).count(),
                 awards.filter(payment_status=UNPAID).count(), total, paid, pending),
                key,
            )

    def test_excel_export(self):
        september = self.periods["01.09"]
        response = self.web.get(reverse("admin:scholarships_report_xlsx"), {"period": september.pk})
        rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
        data = [r for r in rows[4:] if r[0] and r[5] in ("Выдано", "Не выдано")]
        db = {a.evaluation.student_name: a for a in september.awards.select_related("evaluation", "paid_by")}
        self.assertEqual((len(data), len({r[0] for r in data})), (50, 50))
        for row in data:
            award = db[row[0]]
            self.assertEqual((D(row[4]), row[5] == "Выдано"), (award.amount, award.is_paid), row[0])
            if award.is_paid:
                self.assertEqual(
                    (row[6], row[7], row[8]),
                    (f"{timezone.localtime(award.paid_at):%d.%m.%Y %H:%M}", award.get_payment_method_display(),
                     str(award.paid_by)),
                )
            else:
                self.assertEqual((row[6], row[7], row[8]), (None, None, None))
        totals = {r[3]: r[4] for r in rows if r[3] in ("Всего начислено", "Всего выплачено", "Остаток")}
        self.assertEqual(totals, {"Всего начислено": 150000, "Всего выплачено": 90000, "Остаток": 60000})

    def test_permissions(self):
        award = self.unpaid().first()
        self.assertEqual(self.client_class().post(self.pay_url, {"award": award.pk, "payment_method": "cash"}).status_code, 302)
        clerk = make_user("clerk", role=User.Role.TEACHER)
        clerk.is_staff = True
        clerk.save()
        web = self.client_class()
        web.force_login(clerk)
        self.assertEqual(web.post(self.pay_url, {"award": award.pk, "payment_method": "cash"}).status_code, 403)
        csrf = self.client_class(enforce_csrf_checks=True)
        csrf.force_login(self.admin)
        self.assertEqual(csrf.post(self.pay_url, {"award": award.pk, "payment_method": "cash"}).status_code, 403)
        api = APIClient()
        api.force_authenticate(self.admin)
        url = reverse("scholarship-award-detail", args=[award.pk])
        self.assertEqual(api.patch(url, {"payment_status": "paid", "paid_amount": "150000"}, format="json").status_code, 405)
        # A draft period's award can't be paid either.
        draft = self.periods["01.11"].awards.first()
        self.pay(draft.pk)
        award.refresh_from_db()
        draft.refresh_from_db()
        self.assertEqual((award.payment_status, draft.payment_status), (UNPAID, UNPAID))


@skipUnless(connection.vendor == "postgresql", "needs real concurrent transactions (PostgreSQL)")
class PaymentConcurrencyTests(TransactionTestCase):
    """Two admins press «Выдать» at the same moment. A real TransactionTestCase
    (ScholarshipFixture is a TestCase: its data would never be committed,
    so the other connections couldn't see it)."""

    def setUp(self):
        self.admin = make_admin()
        call_command("seed_scholarships", stdout=io.StringIO())
        self.august = ScholarshipPeriod.objects.get(seed.mock_period_q(), period_start="2026-08-01")

    def race(self, *id_lists):
        barrier = threading.Barrier(len(id_lists))
        results, errors = [], []

        def worker(ids):
            try:
                barrier.wait(timeout=10)
                results.append(pay_awards(ids, method="cash", user=self.admin))
            except Exception as exc:  # surfaced by the assertions below
                errors.append(exc)
            finally:
                connection.close()

        threads = [threading.Thread(target=worker, args=(ids,)) for ids in id_lists]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=30)
        self.assertEqual(errors, [])
        return results

    def test_concurrent_payment(self):
        award = self.august.awards.filter(payment_status=UNPAID).first()
        results = self.race([award.pk], [award.pk])
        self.assertEqual(sorted(len(r.paid) for r in results), [0, 1])
        self.assertEqual(sorted(len(r.already_paid) for r in results), [0, 1])
        self.assertEqual(ScholarshipRunLog.objects.filter(action=ScholarshipRunLog.Action.PAYMENT).count(), 1)
        self.assertEqual(orm_money(self.august.awards)[1], D(12000))

    def test_concurrent_overlapping_mass_payments(self):
        ids = list(self.august.awards.filter(payment_status=UNPAID).order_by("rank").values_list("pk", flat=True)[:8])
        results = self.race(ids[:5], ids[3:])
        self.assertEqual(sum(len(r.paid) for r in results), 8)
        self.assertEqual(sum(len(r.already_paid) for r in results), 2)
        self.assertEqual(orm_money(self.august.awards)[1], D(10500) + 8 * D(1500))
