"""Scholarship payments: «Выдать», «Выдать выбранным», protection from a
second payment, filters, totals and the accounting exports."""
from __future__ import annotations

import csv
import datetime as dt
import io
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.scholarships.models import PaymentStatus, ScholarshipAward, ScholarshipRunLog
from apps.scholarships.services.generation import approve_period, create_period
from apps.scholarships.services.payments import cancel_payment, pay_awards
from apps.scholarships.services.report import ReportFilters, build_report
from apps.scholarships.services.report_pdf import render_report_pdf
from apps.users.models import User

from .base import OCT_1, SEP_1, ScholarshipFixture, make_user

SEP_30 = dt.date(2026, 9, 30)
NBSP = " "


class PaymentFixture(ScholarshipFixture):
    def setUp(self):
        super().setUp()
        self.configure(require_complete_feedback=False, award_amount=Decimal("1500"))
        for index, name in enumerate(("Aibek", "Nurai", "Azamat")):
            student = self.student(name)
            self.study(student, self.python, self.t_python, attended=8 - index, missed=index)
        self.period = create_period(period_start=SEP_1, period_end=SEP_30, max_recipients=3, today=OCT_1)
        approve_period(self.period, user=self.admin)
        self.awards = {a.evaluation.student_name.split()[0]: a for a in self.period.awards.select_related("evaluation")}
        self.web = self.client_class()
        self.web.force_login(self.admin)
        self.list_url = reverse("admin:scholarships_scholarshipaward_changelist")
        self.pay_url = reverse("admin:scholarships_pay")

    def reload(self, award) -> ScholarshipAward:
        return ScholarshipAward.objects.get(pk=award.pk)


class PayServiceTests(PaymentFixture):
    def test_pay_one(self):
        award = self.awards["Aibek"]
        result = pay_awards([award.pk], method="cash", comment="В кассе", user=self.admin)
        self.assertEqual([a.pk for a in result.paid], [award.pk])
        award = self.reload(award)
        self.assertEqual(award.payment_status, PaymentStatus.PAID)
        self.assertEqual((award.paid_by, award.paid_amount, award.payment_method, award.payment_comment),
                         (self.admin, Decimal("1500"), "cash", "В кассе"))
        self.assertIsNotNone(award.paid_at)
        # «Начислено» is untouched — the payment is a separate status.
        self.assertEqual(award.status, ScholarshipAward.Status.APPROVED)
        log = ScholarshipRunLog.objects.filter(action=ScholarshipRunLog.Action.PAYMENT).get()
        self.assertIn("Aibek Test", log.message)
        self.assertIn("наличные", log.message)

    def test_never_paid_twice(self):
        award = self.awards["Aibek"]
        pay_awards([award.pk], method="cash", user=self.admin)
        first = self.reload(award)
        result = pay_awards([award.pk, self.awards["Nurai"].pk], method="bank", user=self.admin)
        self.assertEqual([a.pk for a in result.already_paid], [award.pk])
        self.assertEqual([a.pk for a in result.paid], [self.awards["Nurai"].pk])
        again = self.reload(award)
        self.assertEqual((again.paid_at, again.payment_method), (first.paid_at, "cash"))

    def test_draft_period_cannot_be_paid(self):
        other = self.student("Bakyt")
        self.study(other, self.python, self.t_python, attended=5)
        draft = create_period(period_start=dt.date(2026, 8, 31), period_end=SEP_30, max_recipients=None, today=OCT_1)
        pending = draft.awards.first()
        result = pay_awards([pending.pk], method="cash", user=self.admin)
        self.assertEqual((result.paid, [a.pk for a in result.not_approved]), ([], [pending.pk]))
        self.assertEqual(self.reload(pending).payment_status, PaymentStatus.UNPAID)
        # …and the database refuses it too.
        with self.assertRaises(IntegrityError), transaction.atomic():
            ScholarshipAward.objects.filter(pk=pending.pk).update(payment_status=PaymentStatus.PAID, paid_at=timezone.now())

    def test_validation(self):
        with self.assertRaisesMessage(ValidationError, "Выберите способ выплаты."):
            pay_awards([self.awards["Aibek"].pk], method="bitcoin", user=self.admin)
        with self.assertRaisesMessage(ValidationError, "Не выбрано ни одной стипендии."):
            pay_awards([], method="cash", user=self.admin)
        self.assertFalse(ScholarshipAward.objects.filter(payment_status=PaymentStatus.PAID).exists())

    def test_cancel(self):
        award = self.awards["Aibek"]
        pay_awards([award.pk], method="other", user=self.admin)
        cancel_payment(award, user=self.admin)
        award = self.reload(award)
        self.assertEqual((award.payment_status, award.paid_at, award.payment_method), (PaymentStatus.UNPAID, None, ""))
        with self.assertRaisesMessage(ValidationError, "ещё не выдана"):
            cancel_payment(award, user=self.admin)


class PayAdminTests(PaymentFixture):
    def test_pay_one_from_the_list(self):
        award = self.awards["Aibek"]
        page = self.web.get(self.list_url)
        self.assertContains(page, "data-ok-pay-one", count=3)
        self.assertContains(page, "Выдача стипендии")
        self.assertContains(page, "Подтвердить выплату")

        next_url = f"{self.list_url}?period__id__exact={self.period.pk}"
        response = self.web.post(self.pay_url, {"award": award.pk, "payment_method": "bank", "next": next_url})
        # Back to the same filtered list, scrolled to the paid row.
        self.assertRedirects(response, f"{next_url}#aw-{award.pk}", fetch_redirect_response=False)
        self.assertTrue(self.reload(award).is_paid)

        page = self.web.get(next_url)
        self.assertContains(page, "стипендия выдана")
        self.assertContains(page, "data-ok-pay-one", count=2)
        self.assertContains(page, "Подробнее", count=1)
        self.assertEqual(page.context["summary"]["paid_sum"], Decimal("1500"))
        self.assertEqual(page.context["summary"]["remaining"], Decimal("3000"))

    def test_bulk_pay_and_repeat(self):
        ids = [a.pk for a in self.awards.values()]
        response = self.web.post(self.pay_url, {"award": ids, "payment_method": "cash", "next": self.list_url}, follow=True)
        self.assertContains(response, f"Выдано стипендий: 3 на сумму 4{NBSP}500{NBSP}сом.")
        self.assertEqual(ScholarshipAward.objects.filter(payment_status=PaymentStatus.PAID).count(), 3)
        self.assertNotContains(response, "data-ok-pay-row")

        response = self.web.post(self.pay_url, {"award": ids[:1], "payment_method": "cash", "next": self.list_url}, follow=True)
        self.assertContains(response, "Уже были выданы — повторно не выплачены")

    def test_bad_input_pays_nothing(self):
        award = self.awards["Aibek"]
        response = self.web.post(self.pay_url, {"award": award.pk, "payment_method": "", "next": "https://evil.example/"})
        self.assertRedirects(response, self.list_url, fetch_redirect_response=False)
        self.assertFalse(self.reload(award).is_paid)
        self.assertEqual(self.web.get(self.pay_url).status_code, 302)

    def test_permissions(self):
        staff = make_user("clerk", role=User.Role.TEACHER)
        staff.is_staff = True
        staff.save()
        web = self.client_class()
        web.force_login(staff)
        response = web.post(self.pay_url, {"award": self.awards["Aibek"].pk, "payment_method": "cash"})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(self.reload(self.awards["Aibek"]).is_paid)

    def test_cancel_from_admin(self):
        award = self.awards["Aibek"]
        pay_awards([award.pk], method="cash", user=self.admin)
        url = reverse("admin:scholarships_cancel_payment", args=[award.pk])
        self.web.post(url, {"next": self.list_url})
        self.assertFalse(self.reload(award).is_paid)

    def test_filters(self):
        pay_awards([self.awards["Aibek"].pk], method="cash", user=self.admin)
        pay_awards([self.awards["Nurai"].pk], method="bank", user=self.admin)
        paid = self.web.get(self.list_url, {"payment_status__exact": "paid"})
        self.assertEqual(paid.context["summary"]["total"], 2)
        self.assertNotContains(paid, "Azamat Test")
        cash = self.web.get(self.list_url, {"payment_method__exact": "cash"})
        self.assertEqual([a.evaluation.student_name for a in cash.context["cl"].result_list], ["Aibek Test"])
        unpaid = self.web.get(self.list_url, {"payment_status__exact": "unpaid", "q": "aza"})
        self.assertEqual(unpaid.context["summary"]["total"], 1)

    def test_work_queue_order(self):
        """Who can be paid now comes first, a draft period's awards next,
        the paid ones last."""
        other = self.student("Bakyt")
        self.study(other, self.python, self.t_python, attended=5)
        create_period(period_start=dt.date(2026, 8, 31), period_end=SEP_30, max_recipients=1, today=OCT_1)
        pay_awards([self.awards["Aibek"].pk], method="cash", user=self.admin)
        page = self.web.get(self.list_url)
        states = [("paid" if a.is_paid else "payable" if a.is_payable else "pending") for a in page.context["cl"].result_list]
        self.assertEqual(states, ["payable", "payable", "pending", "paid"])

    def test_period_list_counts_payments(self):
        pay_awards([self.awards["Aibek"].pk], method="cash", user=self.admin)
        response = self.web.get(reverse("admin:scholarships_scholarshipperiod_changelist"))
        period = response.context["cl"].result_list[0]
        self.assertEqual((period.paid_count, period.unpaid_count), (1, 2))

    def test_dashboard_shows_payment_state(self):
        pay_awards([self.awards["Aibek"].pk], method="cash", user=self.admin)
        page = self.web.get(reverse("admin:scholarships_scholarshipperiod_change", args=[self.period.pk]))
        self.assertContains(page, "Выдано")
        self.assertContains(page, "Не выдано")
        self.assertContains(page, "Выплаты")


class ExportTests(PaymentFixture):
    def setUp(self):
        super().setUp()
        pay_awards([self.awards["Aibek"].pk], method="cash", comment="Касса", user=self.admin)
        self.query = {"period": self.period.pk}

    def test_xlsx(self):
        response = self.web.get(reverse("admin:scholarships_report_xlsx"), self.query)
        self.assertEqual(response.status_code, 200)
        self.assertIn("scholarship-payments-2026-09-01-2026-09-30.xlsx", response["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(response.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[3][:10], (
            "Ученик", "Группа", "Программа", "Период", "Сумма, сом", "Статус", "Дата выплаты",
            "Способ выплаты", "Ответственный", "Комментарий",
        ))
        aibek = rows[4]
        self.assertEqual((aibek[0], aibek[1], aibek[2], aibek[5], aibek[7], aibek[8], aibek[9]),
                         ("Aibek Test", "Prog SOFT 1", "Prog SOFT", "Выдано", "Наличные", str(self.admin), "Касса"))
        self.assertEqual(aibek[4], 1500)
        totals = {row[3]: row[4] for row in rows if row[3] in {"Всего начислено", "Всего выплачено", "Остаток", "Количество учеников"}}
        self.assertEqual(totals, {"Всего начислено": 4500, "Всего выплачено": 1500, "Остаток": 3000, "Количество учеников": 3})

    def test_xlsx_follows_filters(self):
        response = self.web.get(reverse("admin:scholarships_report_xlsx"), {**self.query, "status": "unpaid"})
        sheet = load_workbook(io.BytesIO(response.content)).active
        names = [row[0] for row in sheet.iter_rows(min_row=5, values_only=True) if row[0] and row[0].endswith("Test")]
        self.assertEqual(sorted(names), ["Azamat Test", "Nurai Test"])

    def test_csv(self):
        response = self.web.get(reverse("admin:scholarships_report_csv"), {**self.query, "method": "cash"})
        self.assertEqual(response.status_code, 200)
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(rows[0][0], "Ученик")
        self.assertEqual(rows[1][:2], ["Aibek Test", "Prog SOFT 1"])
        self.assertIn(["Всего выплачено", "1500"], rows)
        self.assertIn(["Остаток", "0"], rows)

    def test_pdf_with_payments(self):
        response = self.web.get(reverse("admin:scholarships_report_pdf"), self.query)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))
        pdf = render_report_pdf(build_report(self.period, ReportFilters(status="paid")))
        self.assertTrue(pdf.startswith(b"%PDF"))

    def test_report_page(self):
        page = self.web.get(reverse("admin:scholarships_report"), {**self.query, "method": "cash"})
        self.assertEqual(page.context["report"].rows_totals.paid_amount, Decimal("1500"))
        self.assertContains(page, "Касса")  # comment as a tooltip
        self.assertNotContains(page, "Nurai Test")

    def test_no_periods(self):
        ScholarshipAward.objects.all().delete()
        self.period.delete()
        for name in ("scholarships_report_xlsx", "scholarships_report_csv"):
            self.assertEqual(self.web.get(reverse(f"admin:{name}")).status_code, 302)
