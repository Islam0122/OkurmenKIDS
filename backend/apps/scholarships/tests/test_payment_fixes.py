"""Regressions for the SCHOLARSHIP FUNCTIONAL AUDIT fixes, on the
seed_scholarships mock data (see test_payment_audit.MockMoneyFixture).

1. deleting a student never erases scholarship history;
2. every award has an amount ≥ 0, and only a positive amount is paid;
3. who paid survives the payer's deactivation / deletion attempt;
4. cancelling a payment keeps the original payment in the journal;
5. payment and journal entry are one transaction;
+  the database constraints behind all of the above.
"""
from __future__ import annotations

import io
import threading
from decimal import Decimal
from unittest import mock, skipUnless

from django.core.management import call_command
from django.db import IntegrityError, connection, transaction
from django.db.models import ProtectedError, RestrictedError, Sum
from django.test import TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.academy.models import Student
from apps.scholarships.demo import mock_payments as seed
from apps.scholarships.models import PaymentStatus, ScholarshipAward, ScholarshipPeriod, ScholarshipRunLog
from apps.scholarships.services.payments import cancel_payment, pay_awards
from apps.scholarships.services.report import ReportFilters, build_report
from apps.users.models import User

from .base import make_admin
from .test_payment_audit import MockMoneyFixture, orm_money

D = Decimal
PAID, UNPAID = PaymentStatus.PAID, PaymentStatus.UNPAID


def all_money():
    return orm_money(ScholarshipAward.objects.all())


def payment_logs():
    return ScholarshipRunLog.objects.filter(action=ScholarshipRunLog.Action.PAYMENT)


class StudentDeletionTests(MockMoneyFixture):
    def test_cannot_delete_student_with_scholarship_history(self):
        award = self.august.awards.filter(payment_status=PAID).select_related("student").first()
        student = award.student
        history = list(student.scholarship_awards.values_list("pk", "amount", "payment_status", "paid_amount"))
        money = all_money()

        api = APIClient()
        api.force_authenticate(self.admin)
        response = api.delete(reverse("student-detail", args=[student.pk]))
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["detail"], "Нельзя удалить ученика: существует история стипендий.")

        # The database refuses it too, whatever the entry point.
        with self.assertRaises(RestrictedError), transaction.atomic():
            Student.objects.filter(pk=student.pk).delete()
        with self.assertRaises(RestrictedError), transaction.atomic():
            student.scholarship_evaluations.all().delete()

        self.assertTrue(Student.objects.filter(pk=student.pk).exists())
        self.assertEqual(
            list(student.scholarship_awards.values_list("pk", "amount", "payment_status", "paid_amount")), history,
        )
        self.assertEqual(all_money(), money)

    def test_student_without_scholarships_can_still_be_deleted(self):
        plain = self.student("Plain")
        api = APIClient()
        api.force_authenticate(self.admin)
        self.assertEqual(api.delete(reverse("student-detail", args=[plain.pk])).status_code, 204)

    def test_draft_period_deletion_still_takes_its_awards(self):
        november = self.periods["01.11"]
        self.assertEqual(november.awards.filter(payment_status=PAID).count(), 0)
        november.delete()
        self.assertFalse(ScholarshipAward.objects.filter(period_id=november.pk).exists())


class AmountRulesTests(MockMoneyFixture):
    def test_cannot_pay_null_amount(self):
        award = self.unpaid().first()
        with self.assertRaises(IntegrityError), transaction.atomic():
            ScholarshipAward.objects.filter(pk=award.pk).update(amount=None)
        # A period without «Сумма стипендии» now gives amount 0, never NULL.
        from apps.scholarships.services.generation import award_amount_of
        self.assertEqual(award_amount_of(ScholarshipPeriod(award_amount=None)), D(0))

    def test_cannot_pay_zero_amount(self):
        award = self.unpaid().first()
        ScholarshipAward.objects.filter(pk=award.pk).update(amount=0)
        money = all_money()
        result = pay_awards([award.pk], method="cash", user=self.admin)
        self.assertEqual(([a.pk for a in result.no_amount], result.paid), ([award.pk], []))
        response = self.web.post(self.pay_url, {"award": award.pk, "payment_method": "cash", "next": self.list_url},
                                 follow=True)
        self.assertContains(response, "Сумма стипендии 0 сом — выдать нельзя")
        award.refresh_from_db()
        self.assertEqual(award.payment_status, UNPAID)
        self.assertEqual(all_money(), money)
        # …and the database won't hold a paid award with no money either.
        with self.assertRaises(IntegrityError), transaction.atomic():
            ScholarshipAward.objects.filter(pk=award.pk).update(
                payment_status=PAID, paid_at=timezone.now(), paid_by=self.admin, paid_by_name="x",
                paid_amount=0, payment_method="cash",
            )

    def test_cannot_pay_negative_amount(self):
        award = self.unpaid().first()
        with self.assertRaises(IntegrityError), transaction.atomic():
            ScholarshipAward.objects.filter(pk=award.pk).update(amount=D("-1500"))
        paid = self.august.awards.filter(payment_status=PAID).first()
        with self.assertRaises(IntegrityError), transaction.atomic():
            ScholarshipAward.objects.filter(pk=paid.pk).update(paid_amount=D("-1"))
        with self.assertRaises(IntegrityError), transaction.atomic():
            ScholarshipAward.objects.filter(pk=paid.pk).update(paid_amount=paid.amount + 1)


class PayerHistoryTests(MockMoneyFixture):
    def test_payment_history_survives_user_deactivation_or_deletion(self):
        islam = make_admin("islam")
        islam.first_name, islam.last_name = "Islam", "Duishobaev"
        islam.save()
        award = self.unpaid().first()
        pay_awards([award.pk], method="bank", user=islam)

        # Deletion is refused while the user has payments…
        with self.assertRaises(ProtectedError):
            islam.delete()
        # …and deactivation / renaming never rewrites history.
        islam.is_active = False
        islam.first_name = "Renamed"
        islam.save()
        award.refresh_from_db()
        self.assertEqual((award.paid_by_id, award.paid_by_name, award.paid_by_label), (islam.pk, "Islam Duishobaev", "Islam Duishobaev"))
        row = next(r for r in build_report(self.august, ReportFilters()).rows if r.award_id == award.pk)
        self.assertEqual(row.paid_by, "Islam Duishobaev")
        sheet = load_workbook(io.BytesIO(self.web.get(
            reverse("admin:scholarships_report_xlsx"), {"period": self.august.pk}).content)).active
        self.assertIn("Islam Duishobaev", [r[8] for r in sheet.iter_rows(values_only=True)])
        self.assertContains(self.web.get(self.list_url, {"period__id__exact": self.august.pk}),
                            'data-paid-by="Islam Duishobaev"')
        # The admin «delete user» page explains instead of deleting.
        page = self.web.post(reverse("admin:users_user_delete", args=[islam.pk]), {"post": "yes"})
        self.assertTrue(User.objects.filter(pk=islam.pk).exists())
        self.assertIn(page.status_code, (200, 403))


class CancelHistoryTests(MockMoneyFixture):
    def setUp(self):
        super().setUp()
        self.award = self.august.awards.filter(payment_status=PAID).select_related("evaluation").first()
        self.original = (self.award.paid_amount, self.award.paid_by_id, self.award.paid_by_label,
                         self.award.paid_at, self.award.payment_method)
        cancel_payment(self.award, reason="Ошибка при выплате", user=self.admin)
        self.log = payment_logs().latest("pk")

    def test_cancel_payment_preserves_original_amount(self):
        self.award.refresh_from_db()
        self.assertEqual((self.award.payment_status, self.award.paid_amount), (UNPAID, None))
        self.assertEqual(D(self.log.details["original_amount"]), self.original[0])
        self.assertEqual(self.log.details["event"], "payment_cancelled")
        self.assertIn("Сумма: 1 500 сом", self.log.message)

    def test_cancel_payment_preserves_original_paid_by(self):
        d = self.log.details
        self.assertEqual((d["original_paid_by_id"], d["original_paid_by"]), (self.original[1], self.original[2]))
        self.assertEqual((d["cancelled_by_id"], d["cancelled_by"], d["reason"]), (self.admin.pk, str(self.admin), "Ошибка при выплате"))
        self.assertEqual(self.log.triggered_by, self.admin)

    def test_cancel_payment_preserves_original_paid_at(self):
        d = self.log.details
        self.assertEqual(d["original_paid_at"], self.original[3].isoformat())
        self.assertEqual(d["original_payment_method"], self.original[4])
        self.assertIsNotNone(d["cancelled_at"])
        self.assertIn(f"{timezone.localtime(self.original[3]):%d.%m.%Y %H:%M}", self.log.message)


class JournalAtomicityTests(MockMoneyFixture):
    def test_payment_and_journal_are_atomic(self):
        ids = list(self.unpaid().values_list("pk", flat=True)[:3])
        logs = payment_logs().count()
        pay_awards(ids, method="cash", user=self.admin)
        self.assertEqual(payment_logs().count(), logs + 1)
        log = payment_logs().latest("pk")
        self.assertEqual(sorted(a["award_id"] for a in log.details["awards"]), sorted(ids))
        self.assertEqual((D(log.details["total"]), log.details["paid_by_id"]), (D(4500), self.admin.pk))
        self.assertIn("4 500 сом", log.message)

    def test_journal_failure_rolls_back_payment(self):
        ids = list(self.unpaid().values_list("pk", flat=True)[:5])
        money, logs = all_money(), payment_logs().count()
        with mock.patch.object(ScholarshipRunLog.objects, "create", side_effect=RuntimeError("journal is down")):
            with self.assertRaises(RuntimeError):
                pay_awards(ids, method="cash", user=self.admin)
        self.assertFalse(ScholarshipAward.objects.filter(pk__in=ids, payment_status=PAID).exists())
        self.assertEqual((all_money(), payment_logs().count()), (money, logs))
        # Same for a cancellation.
        paid = self.august.awards.filter(payment_status=PAID).first()
        with mock.patch.object(ScholarshipRunLog.objects, "create", side_effect=RuntimeError("journal is down")):
            with self.assertRaises(RuntimeError):
                cancel_payment(paid, reason="x", user=self.admin)
        paid.refresh_from_db()
        self.assertTrue(paid.is_paid)


class ConstraintTests(MockMoneyFixture):
    def setUp(self):
        super().setUp()
        self.paid = self.august.awards.filter(payment_status=PAID).first()
        self.unpaid_award = self.unpaid().first()

    def assertRefused(self, pk, **fields):
        with self.assertRaises(IntegrityError), transaction.atomic():
            ScholarshipAward.objects.filter(pk=pk).update(**fields)

    def test_paid_requires_amount(self):
        self.assertRefused(self.paid.pk, paid_amount=None)
        self.assertRefused(self.paid.pk, amount=0)

    def test_paid_requires_paid_at(self):
        self.assertRefused(self.paid.pk, paid_at=None)

    def test_paid_requires_paid_by(self):
        self.assertRefused(self.paid.pk, paid_by=None)
        self.assertRefused(self.paid.pk, paid_by_name="")
        self.assertRefused(self.paid.pk, payment_method="")

    def test_pending_has_no_payment_data(self):
        for field, value in (("paid_at", timezone.now()), ("paid_by", self.admin), ("paid_amount", D(1500)),
                             ("payment_method", "cash"), ("paid_by_name", "x")):
            self.assertRefused(self.unpaid_award.pk, **{field: value})
        self.assertRefused(self.unpaid_award.pk, payment_status="refunded")
        self.assertFalse(ScholarshipAward.objects.filter(payment_status=UNPAID).exclude(
            paid_at=None).exists())


class FinancialRegressionTests(MockMoneyFixture):
    def test_financial_totals_unchanged_after_failed_operations(self):
        money = all_money()
        self.assertEqual(money, (D(327500), D(158000), D(169500), D(169500)))
        paid = self.august.awards.filter(payment_status=PAID).select_related("student").first()
        unpaid = self.unpaid().first()
        api = APIClient()
        api.force_authenticate(self.admin)
        api.delete(reverse("student-detail", args=[paid.student_id]))                  # 409
        self.pay([paid.pk])                                                            # already paid
        self.pay([unpaid.pk], "bitcoin")                                               # bad method
        self.pay([self.periods["01.11"].awards.first().pk])                            # not approved
        self.pay([999999])                                                             # no such award
        self.web.post(reverse("admin:scholarships_cancel_payment", args=[paid.pk]), {}) # no reason
        with mock.patch.object(ScholarshipRunLog.objects, "create", side_effect=RuntimeError):
            with self.assertRaises(RuntimeError):
                pay_awards([unpaid.pk], method="cash", user=self.admin)                # journal down
        self.assertEqual(all_money(), money)

    def test_duplicate_payment_still_protected(self):
        award = self.unpaid().first()
        self.pay(award.pk, amount="150000", paid_by=make_admin("x").pk, paid_at="2020-01-01")
        award.refresh_from_db()
        first = (award.paid_at, award.paid_amount, award.paid_by_id)
        self.assertEqual(first[1:], (D(1500), self.admin.pk))
        for _ in range(3):
            self.pay([award.pk, award.pk], "bank")
        award.refresh_from_db()
        self.assertEqual((award.paid_at, award.paid_amount, award.paid_by_id), first)
        self.assertEqual(orm_money(self.august.awards)[1], D(12000))


@skipUnless(connection.vendor == "postgresql", "needs real concurrent transactions (PostgreSQL)")
class ConcurrencyStillProtectedTests(TransactionTestCase):
    def setUp(self):
        self.admin = make_admin()
        call_command("seed_scholarships", stdout=io.StringIO())
        self.august = ScholarshipPeriod.objects.get(seed.mock_period_q(), period_start="2026-08-01")

    def test_concurrent_payment_still_protected(self):
        award = self.august.awards.filter(payment_status=UNPAID).first()
        barrier = threading.Barrier(2)
        results, errors = [], []

        def worker():
            try:
                barrier.wait(timeout=10)
                results.append(pay_awards([award.pk], method="cash", user=self.admin))
            except Exception as exc:  # surfaced below
                errors.append(exc)
            finally:
                connection.close()

        threads = [threading.Thread(target=worker) for _ in range(2)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=30)
        self.assertEqual(errors, [])
        self.assertEqual(sorted(len(r.paid) for r in results), [0, 1])
        # One payment, one journal entry — the journal is in the same transaction.
        self.assertEqual(payment_logs().filter(details__event="payment").count(), 1)
        self.assertEqual(
            self.august.awards.filter(payment_status=PAID).aggregate(s=Sum("paid_amount"))["s"], D(12000),
        )
