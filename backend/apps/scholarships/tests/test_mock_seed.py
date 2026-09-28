"""`seed_scholarships`: mock payment data — every payment state, idempotent,
and --clear/--remove touch only mock rows."""
from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

from django.core.management import call_command
from django.urls import reverse
from openpyxl import load_workbook
from django.db.models import Q, Sum

from apps.academy.models import Group, Student
from apps.scholarships.demo import mock_payments as mock
from apps.scholarships.models import PaymentMethod, PaymentStatus, ScholarshipAward, ScholarshipPeriod
from apps.scholarships.services.generation import create_period
from apps.scholarships.services.report import ReportFilters, build_report

from .base import OCT_1, ScholarshipFixture


def run(*args) -> str:
    out = io.StringIO()
    call_command("seed_scholarships", *args, stdout=out)
    return out.getvalue()


class MockSeedTests(ScholarshipFixture):
    def setUp(self):
        super().setUp()
        self.real_student = self.student("Real")
        self.study(self.real_student, self.python, self.t_python, attended=4)
        self.real_period = create_period(
            period_start=dt.date(2026, 8, 31), period_end=dt.date(2026, 9, 30), max_recipients=None, today=OCT_1,
        )

    def periods(self):
        return {p.period_start.strftime("%d.%m") + "-" + p.period_end.strftime("%d.%m"): p
                for p in ScholarshipPeriod.objects.filter(mock.mock_period_q())}

    def test_financial_scenarios(self):
        output = run()
        self.assertIn("SCHOLARSHIP MOCK FINANCIAL DATA", output)
        self.assertIn("Суммы сверены", output)
        periods = self.periods()
        self.assertEqual(len(periods), 8)
        # (awards, total, paid, paid amount, pending amount) — spelled out, not derived from the plan.
        expected = {
            "01.05-31.05": (20, "30000", 20, "30000", "0"),       # fully paid
            "01.06-30.06": (20, "35000", 8, "14000", "21000"),    # 5×2500 + 5×2000 + 5×1500 + 5×1000
            "01.07-31.07": (20, "30000", 0, "0", "30000"),        # nothing paid
            "01.08-31.08": (20, "30000", 7, "10500", "19500"),    # partly paid
            "15.08-15.09": (0, "0", 0, "0", "0"),                 # no scholarships
            "01.09-30.09": (50, "150000", 30, "90000", "60000"),  # 50 × 3000
            "01.10-31.10": (15, "22500", 9, "13500", "9000"),
            "01.11-30.11": (20, "30000", 0, "0", "30000"),        # not approved
        }
        for key, (count, total, paid, paid_amount, pending_amount) in expected.items():
            t = build_report(periods[key], ReportFilters()).rows_totals
            self.assertEqual(
                (t.awards, t.amount, t.paid, t.paid_amount, t.remaining, t.amount - t.paid_amount),
                (count, Decimal(total), paid, Decimal(paid_amount), Decimal(pending_amount), Decimal(pending_amount)),
                key,
            )
        june = periods["01.06-30.06"].awards
        by_amount = sorted(june.values_list("amount", flat=True))
        self.assertEqual({a: by_amount.count(a) for a in by_amount},
                         {Decimal(v): 5 for v in ("1000", "1500", "2000", "2500")})
        self.assertEqual(len(set(june.filter(payment_status=PaymentStatus.PAID).values_list("amount", flat=True))), 4)
        self.assertEqual(periods["15.08-15.09"].max_recipients, 10)
        self.assertTrue(periods["01.11-30.11"].is_draft)

        awards = ScholarshipAward.objects.filter(period__in=periods.values())
        self.assertFalse(awards.filter(amount__isnull=True).exists())
        paid = awards.filter(payment_status=PaymentStatus.PAID)
        self.assertEqual(paid.filter(paid_by=self.admin).exclude(paid_at=None).exclude(payment_method="").count(), paid.count())
        self.assertFalse(paid.exclude(status=ScholarshipAward.Status.APPROVED).exists())
        unpaid = awards.filter(payment_status=PaymentStatus.UNPAID)
        self.assertFalse(unpaid.filter(Q(paid_at__isnull=False) | Q(paid_by__isnull=False) | ~Q(payment_method="")).exists())
        cash = paid.filter(payment_method=PaymentMethod.CASH).count()
        self.assertTrue(0.6 <= cash / paid.count() <= 0.8)
        self.assertGreater(len({a.paid_at.date() for a in paid}), 4)

    def test_verification_catches_a_wrong_number(self):
        run()
        award = ScholarshipAward.objects.filter(payment_status=PaymentStatus.PAID).first()
        ScholarshipAward.objects.filter(pk=award.pk).update(paid_amount=award.amount - 1)
        result = mock.SeedResult()
        result.checks = [mock.PeriodCheck(period=award.period, plan=next(
            p for p in mock.PLANS if p.start == award.period.period_start and p.end == award.period.period_end
        ), figures={})]
        with self.assertRaises(mock.FinanceMismatch):
            mock.verify_finances(result.checks)

    def test_idempotent(self):
        run()
        first = (Student.objects.count(), ScholarshipAward.objects.count(), sorted(p.pk for p in self.periods().values()))
        paid_sum = ScholarshipAward.objects.aggregate(s=Sum("paid_amount"))["s"]
        run()
        self.assertEqual(
            (Student.objects.count(), ScholarshipAward.objects.count(), sorted(p.pk for p in self.periods().values())),
            first,
        )
        self.assertEqual(ScholarshipAward.objects.aggregate(s=Sum("paid_amount"))["s"], paid_sum)

    def test_mock_students_stay_out_of_real_periods(self):
        run()
        self.assertFalse(self.real_period.evaluations.filter(mock.mock_student_q("student__")).exists())
        from apps.scholarships.services.generation import recalculate_period
        recalculate_period(self.real_period, user=self.admin, today=OCT_1)
        self.assertEqual(list(self.real_period.evaluations.values_list("student_id", flat=True)), [self.real_student.pk])

    def test_clear_and_remove_touch_only_mock_rows(self):
        real_awards = self.real_period.awards.count()
        run()
        run("--clear")
        self.assertEqual(len(self.periods()), 8)
        run("--remove")
        self.assertEqual(self.periods(), {})
        self.assertFalse(Student.objects.filter(mock.mock_student_q()).exists())
        self.assertFalse(Group.objects.filter(name__startswith=mock.GROUP_PREFIX).exists())
        self.assertTrue(Student.objects.filter(pk=self.real_student.pk).exists())
        self.assertTrue(Group.objects.filter(pk=self.group.pk).exists())
        self.assertEqual(self.real_period.awards.count(), real_awards)

    def test_real_period_with_same_dates_is_kept(self):
        manual = create_period(period_start=dt.date(2026, 5, 1), period_end=dt.date(2026, 5, 31),
                               max_recipients=5, today=OCT_1)
        output = run()
        self.assertIn("уже есть реальный период", output)
        self.assertEqual(len(self.periods()), 7)
        manual.refresh_from_db()
        self.assertFalse(manual.title.startswith(mock.TITLE_PREFIX))


class MockMoneyFlowTests(ScholarshipFixture):
    """The money checks from the brief, on the seeded August period:
    20 × 1 500 = 30 000, 7 paid (10 500), 13 pending (19 500)."""

    def setUp(self):
        super().setUp()
        run()
        self.period = ScholarshipPeriod.objects.get(mock.mock_period_q(), period_start=dt.date(2026, 8, 1))
        self.web = self.client_class()
        self.web.force_login(self.admin)
        self.list_url = reverse("admin:scholarships_scholarshipaward_changelist")
        self.pay_url = reverse("admin:scholarships_pay")

    def summary(self):
        return self.web.get(self.list_url, {"period__id__exact": self.period.pk}).context["summary"]

    def report(self):
        return build_report(self.period, ReportFilters()).rows_totals

    def unpaid(self):
        return self.period.awards.filter(payment_status=PaymentStatus.UNPAID).order_by("rank")

    def test_money_flow(self):
        # 1–3. total, paid, remaining
        s = self.summary()
        self.assertEqual((s["total"], s["accrued"], s["paid"], s["paid_sum"], s["unpaid"], s["remaining"]),
                         (20, Decimal("30000"), 7, Decimal("10500"), 13, Decimal("19500")))

        # 4–6. pay one → paid_amount +1 500, remaining −1 500
        one = self.unpaid().first()
        self.web.post(self.pay_url, {"award": one.pk, "payment_method": "cash", "next": self.list_url})
        s = self.summary()
        self.assertEqual((s["paid"], s["paid_sum"], s["remaining"]), (8, Decimal("12000"), Decimal("18000")))
        t = self.report()
        self.assertEqual((t.paid, t.paid_amount, t.remaining), (8, Decimal("12000"), Decimal("18000")))

        # 7. bulk: 3 more
        ids = list(self.unpaid().values_list("pk", flat=True)[:3])
        self.web.post(self.pay_url, {"award": ids, "payment_method": "bank", "next": self.list_url})
        s = self.summary()
        self.assertEqual((s["paid"], s["paid_sum"], s["unpaid"], s["remaining"]), (11, Decimal("16500"), 9, Decimal("13500")))

        # 8. paying the same four again changes nothing
        page = self.web.post(self.pay_url, {"award": [one.pk, *ids], "payment_method": "cash", "next": self.list_url},
                             follow=True)
        self.assertContains(page, "повторно не выплачены")
        self.assertEqual(self.summary()["paid_sum"], Decimal("16500"))

        # 9. accounting report page
        page = self.web.get(reverse("admin:scholarships_report"), {"period": self.period.pk})
        t = page.context["report"].rows_totals
        self.assertEqual((t.amount, t.paid_amount, t.remaining), (Decimal("30000"), Decimal("16500"), Decimal("13500")))
        self.assertContains(page, "16\u00a0500\u00a0сом")

        # 10. Excel: every row has its amount, totals match
        response = self.web.get(reverse("admin:scholarships_report_xlsx"), {"period": self.period.pk})
        rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
        amounts = [row[4] for row in rows[4:] if row[0] and row[5] in ("Выдано", "Не выдано")]
        self.assertEqual((len(amounts), set(amounts)), (20, {1500}))
        totals = {row[3]: row[4] for row in rows if row[3] in ("Всего начислено", "Всего выплачено", "Остаток")}
        self.assertEqual(totals, {"Всего начислено": 30000, "Всего выплачено": 16500, "Остаток": 13500})

        # 11. PDF is built from the same report object
        response = self.web.get(reverse("admin:scholarships_report_pdf"), {"period": self.period.pk})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))
