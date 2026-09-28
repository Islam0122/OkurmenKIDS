"""Accounting exports of «Отчёты по стипендиям» — Excel and CSV.

Both are built from the very `ScholarshipReport` the page and the PDF
render (same period, same filters), so the accountant's file and the
admin's screen can never disagree. One row per award, then the totals.
"""
from __future__ import annotations

import csv
import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .report import NO_GROUP, ScholarshipReport

COLUMNS = (
    ("Ученик", 30),
    ("Группа", 16),
    ("Программа", 22),
    ("Период", 24),
    ("Сумма, сом", 13),
    ("Статус", 15),
    ("Дата выплаты", 17),
    ("Способ выплаты", 20),
    ("Ответственный", 22),
    ("Комментарий", 36),
)
AMOUNT_COLUMN = 5  # 1-based, «Сумма, сом»


def _paid_at(row) -> str:
    return f"{timezone.localtime(row.paid_at):%d.%m.%Y %H:%M}" if row.paid_at else ""


def _values(report: ScholarshipReport, row) -> list:
    return [
        row.student_name,
        row.group_name or NO_GROUP,
        row.program,
        report.range_label,
        row.amount,
        row.status_label,
        _paid_at(row),
        row.method_label,
        row.paid_by,
        row.comment,
    ]


def summary_lines(report: ScholarshipReport) -> list[tuple[str, object]]:
    t = report.rows_totals
    return [
        ("Всего начислено", t.amount),
        ("Всего выплачено", t.paid_amount),
        ("Остаток", t.remaining),
        ("Количество учеников", t.students),
    ]


def filters_label(report: ScholarshipReport) -> str:
    f = report.filters
    parts = [
        f"группа {f.group}" if f.group else "",
        f"программа {f.program}" if f.program else "",
        f"статус «{f.status_label}»" if f.status else "",
        f"способ «{f.method_label}»" if f.method else "",
        f"ученик «{f.student}»" if f.student else "",
    ]
    return ", ".join(p for p in parts if p)


def filename(report: ScholarshipReport, extension: str) -> str:
    return f"scholarship-payments-{report.period.period_start}-{report.period.period_end}.{extension}"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_BRAND = "2F8F5B"
_BRAND_SOFT = "E7F3EC"
_BORDER = Side(style="thin", color="E1E7E2")
_STATUS_FILL = {"Выдано": "E7F3EC", "Не выдано": "FBF1DF"}


def export_xlsx(report: ScholarshipReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Стипендии"
    last_col = get_column_letter(len(COLUMNS))

    ws["A1"] = f"Отчёт по стипендиям · {report.range_label}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(f"A1:{last_col}1")
    meta = f"{report.period.title} · сформирован {timezone.localtime():%d.%m.%Y %H:%M}"
    if filters_label(report):
        meta += f" · фильтр: {filters_label(report)}"
    ws["A2"] = meta
    ws["A2"].font = Font(color="68736C", size=10)
    ws.merge_cells(f"A2:{last_col}2")

    header_row = 4
    for index, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=index, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_BRAND)
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[header_row].height = 20

    row_index = header_row
    for row in report.rows:
        row_index += 1
        for column, value in enumerate(_values(report, row), start=1):
            cell = ws.cell(row=row_index, column=column, value=value)
            cell.border = Border(bottom=_BORDER)
            if column == AMOUNT_COLUMN:
                cell.number_format = "#,##0"
            if column == 6 and value in _STATUS_FILL:
                cell.fill = PatternFill("solid", fgColor=_STATUS_FILL[value])
    if report.rows:
        ws.auto_filter.ref = f"A{header_row}:{last_col}{row_index}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    row_index += 2
    for label, value in summary_lines(report):
        label_cell = ws.cell(row=row_index, column=AMOUNT_COLUMN - 1, value=label)
        value_cell = ws.cell(row=row_index, column=AMOUNT_COLUMN, value=value)
        label_cell.font = value_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="right")
        label_cell.fill = value_cell.fill = PatternFill("solid", fgColor=_BRAND_SOFT)
        if isinstance(value, Decimal):
            value_cell.number_format = "#,##0"
        row_index += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def _csv_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return str(value.quantize(Decimal("1")))
    return str(value)


def export_csv(report: ScholarshipReport) -> str:
    buffer = io.StringIO()
    buffer.write("﻿")  # BOM so Excel opens Cyrillic correctly
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow([title for title, _ in COLUMNS])
    for row in report.rows:
        writer.writerow([_csv_value(v) for v in _values(report, row)])
    writer.writerow([])
    for label, value in summary_lines(report):
        writer.writerow([label, _csv_value(value)])
    return buffer.getvalue()
