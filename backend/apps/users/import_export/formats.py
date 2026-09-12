"""CSV / XLSX read & write helpers shared by every Import/Export service.

Model-agnostic on purpose: `read_rows()` turns an uploaded file into a list
of ``{header: value}`` dicts (every value coerced to a plain string), and
`build_export_response()` does the reverse for a download. Anything that
knows about Student/Teacher fields belongs in the caller, not here.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

CSV_EXTENSIONS = (".csv",)
XLSX_EXTENSIONS = (".xlsx",)

TRUE_VALUES = {"1", "true", "yes", "y", "да", "истина", "on"}
FALSE_VALUES = {"0", "false", "no", "n", "нет", "ложь", "off"}

_PHONE_RE = re.compile(r"^[0-9+()\-\s]{5,30}$")


class UnsupportedFileFormat(Exception):
    """Raised for an unrecognized upload extension or export `format` value."""


def read_rows(uploaded_file) -> list[dict[str, str]]:
    """Parse an uploaded CSV or XLSX file into a list of {header: value} rows.

    The first row is always treated as the header. Every cell is coerced to
    a stripped string so downstream validators only ever deal with one
    type. Fully blank rows are skipped (common trailing-rows artifact from
    spreadsheet exports).
    """
    name = (getattr(uploaded_file, "name", "") or "").lower()
    if name.endswith(XLSX_EXTENSIONS):
        return _read_xlsx(uploaded_file)
    if name.endswith(CSV_EXTENSIONS):
        return _read_csv(uploaded_file)
    raise UnsupportedFileFormat("Неподдерживаемый формат файла. Используйте CSV или XLSX.")


def _read_csv(uploaded_file) -> list[dict[str, str]]:
    raw = uploaded_file.read()
    text = raw.decode("utf-8-sig") if isinstance(raw, bytes) else raw
    all_rows = list(csv.reader(io.StringIO(text)))
    if not all_rows:
        return []

    header = [cell.strip() for cell in all_rows[0]]
    rows = []
    for raw_row in all_rows[1:]:
        if not any(cell.strip() for cell in raw_row):
            continue
        row = {key: (raw_row[i].strip() if i < len(raw_row) else "") for i, key in enumerate(header)}
        rows.append(row)
    return rows


def _read_xlsx(uploaded_file) -> list[dict[str, str]]:
    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []
    header = [str(cell).strip() if cell is not None else "" for cell in header_row]

    rows = []
    for raw_row in rows_iter:
        if raw_row is None or not any(cell not in (None, "") for cell in raw_row):
            continue
        row = {key: _cell_to_str(raw_row[i] if i < len(raw_row) else None) for i, key in enumerate(header)}
        rows.append(row)
    return rows


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_bool(value: str, *, default: bool | None = None) -> bool:
    """Parse a human-entered boolean. Raises ValueError if it can't be read."""
    normalized = (value or "").strip().lower()
    if not normalized:
        if default is None:
            raise ValueError("empty boolean value")
        return default
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    raise ValueError(f"invalid boolean: {value!r}")


def is_valid_phone(value: str) -> bool:
    return bool(_PHONE_RE.match(value))


def build_export_response(
    rows: list[dict],
    fieldnames: list[str],
    fmt: str,
    base_filename: str,
    headers: dict[str, str] | None = None,
) -> HttpResponse:
    """Build a CSV/XLSX download from ``rows`` (dicts keyed by ``fieldnames``).

    ``headers`` optionally maps a field name to the column label shown in the
    file (e.g. ``{"name": "Название курса"}``); fields missing from the map
    fall back to their raw name, and omitting ``headers`` entirely keeps the
    original behaviour of using ``fieldnames`` as the header row verbatim.
    """
    fmt = (fmt or "csv").strip().lower()
    header_labels = [(headers or {}).get(field, field) for field in fieldnames]
    timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
    if fmt == "xlsx":
        return _build_xlsx_response(rows, fieldnames, header_labels, f"{base_filename}_{timestamp}.xlsx")
    if fmt == "csv":
        return _build_csv_response(rows, fieldnames, header_labels, f"{base_filename}_{timestamp}.csv")
    raise UnsupportedFileFormat("Параметр format должен быть 'csv' или 'xlsx'.")


def _build_csv_response(rows: list[dict], fieldnames: list[str], header_labels: list[str], filename: str) -> HttpResponse:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header_labels)
    for row in rows:
        writer.writerow([row.get(field, "") for field in fieldnames])

    # Leading BOM keeps Excel from mangling Cyrillic content on Windows.
    response = HttpResponse("﻿" + buffer.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _build_xlsx_response(rows: list[dict], fieldnames: list[str], header_labels: list[str], filename: str) -> HttpResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header_labels)
    for row in rows:
        sheet.append([row.get(field, "") for field in fieldnames])

    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
