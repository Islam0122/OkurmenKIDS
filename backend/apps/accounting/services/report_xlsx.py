"""Excel-отчёты бухгалтерии: числовые суммы, автофильтр, строка итогов."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .report_service import IndividualReport, PayrollReport

MONEY_FMT = '#,##0.00'
HEAD_FILL = PatternFill("solid", fgColor="2F8F5B")
TOTAL_FILL = PatternFill("solid", fgColor="E7F3EC")


def _sheet(ws, title: str, subtitle: str, head: list[tuple[str, int]], rows: list[list], money_cols: set[int], total: list | None):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle])
    ws["A2"].font = Font(color="68736C", size=10)
    ws.append([])
    ws.append([h for h, _ in head])
    head_row = ws.max_row
    for i, (_, width) in enumerate(head, start=1):
        cell = ws.cell(row=head_row, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in rows:
        ws.append(row)
    last = ws.max_row
    for r in range(head_row + 1, last + 1):
        for c in money_cols:
            ws.cell(row=r, column=c).number_format = MONEY_FMT
    if rows:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(head))}{last}"
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    if total is not None:
        ws.append(total)
        tr = ws.max_row
        for c in range(1, len(head) + 1):
            cell = ws.cell(row=tr, column=c)
            cell.font = Font(bold=True)
            cell.fill = TOTAL_FILL
            if c in money_cols:
                cell.number_format = MONEY_FMT
    return head_row


def render_period_xlsx(report: PayrollReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Начисления"
    t = report.totals
    _sheet(
        ws, report.title, f"Период: {report.range_label} · KGS · сформировано {report.generated_at:%d.%m.%Y %H:%M}",
        [("Сотрудник", 30), ("Должность", 20), ("Тип оплаты", 26), ("Период", 24), ("Начислено", 14),
         ("Корректировки", 14), ("Итого", 14), ("Выплачено", 14), ("Остаток", 14), ("Статус", 20)],
        [[r.employee, r.position, r.salary_type, r.period, r.accrued, r.adjustments, r.total, r.paid, r.due, r.status]
         for r in report.rows],
        {5, 6, 7, 8, 9},
        ["Итого", "", "", "", t.accrued, t.adjustments, t.total, t.paid, t.due, ""],
    )
    if report.by_group:
        groups = wb.create_sheet("По группам")
        _sheet(groups, "Итоги по группам", report.range_label, [("Группа", 30), ("Программа", 30), ("Сумма", 16)],
               [[g, c, v] for g, c, v in report.by_group], {3}, ["Итого", "", sum(v for _, _, v in report.by_group)])
        programs = wb.create_sheet("По программам")
        _sheet(programs, "Итоги по программам", report.range_label, [("Программа", 36), ("Сумма", 16)],
               [[n, v] for n, v in report.by_program], {2}, ["Итого", sum(v for _, v in report.by_program)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_individual_xlsx(report: IndividualReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Начисления"
    sub = (f"{report.employee} · {report.position or '—'} · {report.salary_type or '—'} · период {report.period} · "
           f"сформировано {report.generated_at:%d.%m.%Y %H:%M}")
    _sheet(
        ws, "Расчётный лист", sub,
        [("Начисление", 60), ("Тип", 22), ("Кол-во", 10), ("Ставка", 14), ("%", 8), ("База", 14), ("Сумма", 14)],
        [[l.description, l.get_line_type_display(), l.quantity, l.rate, l.percentage, l.base_amount, l.amount]
         for l in report.lines],
        {4, 6, 7},
        ["Начислено", "", "", "", "", "", report.accrued],
    )
    adj = wb.create_sheet("Корректировки")
    _sheet(adj, "Корректировки", sub, [("Тип", 22), ("Причина", 60), ("Сумма", 14)],
           [[a.get_kind_display(), a.reason, a.amount] for a in report.adjustments], {3},
           ["Итого", "", report.adjustments_total])
    pay = wb.create_sheet("Выплаты")
    _sheet(pay, "Выплаты", sub, [("Дата", 14), ("Способ", 22), ("Аванс", 8), ("Документ", 20), ("Сумма", 14)],
           [[p.payment_date, p.get_payment_method_display(), "да" if p.is_advance else "", p.reference, p.amount]
            for p in report.payments], {5}, ["Выплачено", "", "", "", report.paid])
    pay.append([])
    pay.append(["Остаток к выплате", "", "", "", report.due])
    pay.cell(row=pay.max_row, column=5).number_format = MONEY_FMT
    pay.cell(row=pay.max_row, column=1).font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
