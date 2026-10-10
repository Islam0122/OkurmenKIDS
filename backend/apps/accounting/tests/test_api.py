"""API бухгалтерии: права доступа, полный сценарий, отчёты (ТЗ §9–§12)."""
from __future__ import annotations

import io

from django.test import override_settings
from openpyxl import load_workbook

from apps.accounting.models import Payroll, PayrollAuditLog, PayrollPayment, SalaryRule
from apps.accounting.services.approval_service import approve_payroll
from apps.accounting.services.payment_service import register_payment, void_payment
from apps.accounting.tests.base import FIRST, SECOND, D, AccountingFixture, day, make_user
from apps.users.models import User

API = "/api/v1/accounting"
R = SalaryRule.RuleType


class PermissionTests(AccountingFixture):
    def setUp(self):
        super().setUp()
        self.rule(self.profile(), R.FIXED, amount=D("40000"))
        self.payroll = self.calc(FIRST)
        self.other = make_user("other")
        self.rule(self.profile(self.other), R.FIXED, amount=D("10000"))
        self.other_payroll = self.calc(FIRST, self.other)

    def test_accountant_cannot_become_director_or_admin(self):
        client = self.client_for(self.accountant)
        self.assertEqual(client.patch("/api/v1/auth/me/", {"role": "director"}, format="json").status_code, 405)
        self.assertEqual(client.post("/api/v1/trainers/", {}, format="json").status_code, 403)
        # Ни одна форма или скрипт не превратит бухгалтера в сотрудника Django admin.
        self.accountant.is_staff = self.accountant.is_superuser = True
        self.accountant.save()
        self.accountant.refresh_from_db()
        self.assertFalse(self.accountant.is_staff or self.accountant.is_superuser)
        self.assertEqual(self.client.get("/admin/").status_code, 302)
        self.assertEqual(User.objects.get(pk=self.accountant.pk).role, User.Role.ACCOUNTANT)

    def test_employee_cannot_see_other_salaries(self):
        client = self.client_for(self.trainer_user)
        for url in (f"{API}/payrolls/", f"{API}/payrolls/{self.other_payroll.pk}/", f"{API}/dashboard/",
                    f"{API}/employees/", f"{API}/audit-log/", f"{API}/student-payments/"):
            self.assertEqual(client.get(url).status_code, 403, url)
        approve_payroll(self.payroll, self.director)
        approve_payroll(self.other_payroll, self.director)
        own = client.get(f"{API}/my/payrolls/").json()
        self.assertEqual([p["id"] for p in own["results"]], [self.payroll.pk])
        self.assertNotIn("audit", own["results"][0])
        self.assertEqual(client.get(f"{API}/my/payrolls/{self.other_payroll.pk}/").status_code, 404)
        self.assertEqual(client.get(f"{API}/my/payrolls/{self.other_payroll.pk}/report.pdf/").status_code, 404)
        self.assertEqual(client.get(f"{API}/my/payrolls/{self.payroll.pk}/report.pdf/").status_code, 200)

    def test_accountant_cannot_approve(self):
        client = self.client_for(self.accountant)
        response = client.post(f"{API}/payrolls/{self.payroll.pk}/approve/")
        self.assertEqual(response.status_code, 403)
        self.payroll.refresh_from_db()
        self.assertEqual(self.payroll.status, Payroll.Status.CALCULATED)

    @override_settings(ACCOUNTING_ACCOUNTANT_CAN_APPROVE=True)
    def test_policy_can_allow_accountant_to_approve(self):
        response = self.client_for(self.accountant).post(f"{API}/payrolls/{self.payroll.pk}/approve/")
        self.assertEqual(response.status_code, 200)

    def test_only_authorised_user_approves(self):
        for user in (self.trainer_user, self.admin, make_user("tl", User.Role.TEAM_LEAD)):
            self.assertEqual(self.client_for(user).post(f"{API}/payrolls/{self.payroll.pk}/approve/").status_code, 403)
        response = self.client_for(self.director).post(f"{API}/payrolls/{self.payroll.pk}/approve/")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["status"], "APPROVED")

    def test_audit_log_is_read_only(self):
        entry = PayrollAuditLog.objects.first()
        for user in (self.accountant, self.director, self.admin, self.trainer_user):
            client = self.client_for(user)
            self.assertIn(client.post(f"{API}/audit-log/", {"action": "x"}, format="json").status_code, (403, 405))
            self.assertIn(client.patch(f"{API}/audit-log/{entry.pk}/", {"action": "x"}, format="json").status_code,
                          (403, 405))
            self.assertIn(client.delete(f"{API}/audit-log/{entry.pk}/").status_code, (403, 405))
        self.assertEqual(PayrollAuditLog.objects.get(pk=entry.pk).action, entry.action)

    def test_admin_reads_but_does_not_operate(self):
        client = self.client_for(self.admin)
        self.assertEqual(client.get(f"{API}/payrolls/").status_code, 200)
        self.assertEqual(client.post(f"{API}/payrolls/{self.payroll.pk}/recalculate/").status_code, 403)

    def test_director_cannot_register_payments(self):
        approve_payroll(self.payroll, self.director)
        response = self.client_for(self.director).post(
            f"{API}/payrolls/{self.payroll.pk}/payments/", {"amount": "100", "payment_date": "2026-09-20"}, format="json",
        )
        self.assertEqual(response.status_code, 403)

    def test_export_is_restricted(self):
        for url in (f"{API}/reports/payroll.pdf?year=2026&month=9", f"{API}/reports/payroll.xlsx?year=2026&month=9",
                    f"{API}/payrolls/{self.other_payroll.pk}/report.pdf/"):
            self.assertEqual(self.client_for(self.trainer_user).get(url).status_code, 403, url)
            self.assertEqual(self.client.get(url).status_code, 401, url)


class FlowTests(AccountingFixture):
    """Бухгалтер рассчитывает → директор утверждает → бухгалтер платит."""

    def test_full_cycle(self):
        acc, boss = self.client_for(self.accountant), self.client_for(self.director)
        r = acc.post(f"{API}/salary-profiles/", {"employee": self.trainer_user.pk, "salary_type": "PER_STUDENT",
                                                  "effective_from": "2026-01-01"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        r = acc.post(f"{API}/salary-rules/", {"employee_profile": r.json()["id"], "rule_type": "PER_STUDENT",
                                               "amount": "11000", "effective_from": "2026-01-01"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["calculation_method"], "STUDENT_DAYS")
        for _ in range(3):
            self.student()
        r = acc.post(f"{API}/periods/", {"year": 2026, "month": 9, "period_type": FIRST}, format="json")
        self.assertEqual(r.status_code, 201)
        period_id = r.json()["id"]
        self.assertEqual(r.json()["start_date"], "2026-09-01")
        self.assertEqual(r.json()["end_date"], "2026-09-15")
        again = acc.post(f"{API}/periods/", {"year": 2026, "month": 9, "period_type": FIRST}, format="json")
        self.assertEqual((again.status_code, again.json()["id"]), (200, period_id))

        r = acc.post(f"{API}/periods/{period_id}/calculate/")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(len(r.json()["calculated"]), 1)
        payroll_id = r.json()["calculated"][0]["id"]
        self.assertEqual(D(r.json()["calculated"][0]["total_accrued"]), D("16500"))
        acc.post(f"{API}/periods/{period_id}/calculate/")
        self.assertEqual(Payroll.objects.count(), 1)

        detail = acc.get(f"{API}/payrolls/{payroll_id}/").json()
        self.assertEqual(len(detail["lines"]), 1)
        self.assertEqual(len(detail["lines"][0]["metadata"]["students"]), 3)
        self.assertTrue(detail["audit"])

        r = acc.post(f"{API}/payrolls/{payroll_id}/payments/", {"amount": "100", "payment_date": "2026-09-16"},
                     format="json")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["code"], "not_approved")

        r = boss.post(f"{API}/periods/{period_id}/approve/")
        self.assertEqual(r.json()["approved"], [payroll_id])
        self.assertEqual(r.json()["period"]["status"], "APPROVED")
        self.assertEqual(acc.post(f"{API}/payrolls/{payroll_id}/recalculate/").status_code, 400)

        body = {"amount": "6500", "payment_date": "2026-09-16", "idempotency_key": "k-1", "is_advance": True}
        first = acc.post(f"{API}/payrolls/{payroll_id}/payments/", body, format="json")
        second = acc.post(f"{API}/payrolls/{payroll_id}/payments/", body, format="json")
        self.assertEqual((first.status_code, second.status_code), (201, 200))
        self.assertEqual(PayrollPayment.objects.count(), 1)
        too_much = acc.post(f"{API}/payrolls/{payroll_id}/payments/", {"amount": "10000.01", "payment_date": "2026-09-17"},
                            format="json")
        self.assertEqual(too_much.json()["code"], "exceeds_due")
        r = acc.post(f"{API}/payrolls/{payroll_id}/payments/", {"amount": "10000", "payment_date": "2026-09-17"},
                     HTTP_IDEMPOTENCY_KEY="k-2", format="json")
        self.assertEqual(r.status_code, 201)
        payroll = acc.get(f"{API}/payrolls/{payroll_id}/").json()
        self.assertEqual(payroll["status"], "PAID")
        self.assertEqual(D(payroll["amount_due"]), D("0"))

        dash = boss.get(f"{API}/dashboard/?period={period_id}").json()
        self.assertEqual(D(dash["total_accrued"]), D("16500"))
        self.assertEqual(D(dash["total_paid"]), D("16500"))
        self.assertEqual(D(dash["total_due"]), D("0"))
        self.assertEqual(dash["employees_with_accruals"], 1)

        employees = acc.get(f"{API}/employees/?period={period_id}&search=trainer").json()["results"]
        self.assertEqual(len(employees), 1)
        self.assertEqual(employees[0]["status"], "PAID")
        self.assertEqual(employees[0]["active_students"], 3)

        r = boss.post(f"{API}/periods/{period_id}/close/")
        self.assertEqual(r.json()["status"], "CLOSED")
        self.assertEqual(acc.post(f"{API}/periods/{period_id}/calculate/").status_code, 400)

    def test_student_payment_and_refund_api(self):
        acc = self.client_for(self.accountant)
        student = self.student()
        r = acc.post(f"{API}/student-payments/", {"student": student.pk, "amount": "10000", "received_date": "2026-09-03",
                                                   "service_start": "2026-09-01", "service_end": "2026-09-30"},
                     format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["group"], self.group_a.pk)
        payment_id = r.json()["id"]
        r = acc.post(f"{API}/student-payments/", {"student": student.pk, "amount": "20000", "received_date": "2026-09-20",
                                                   "refund_of": payment_id}, format="json")
        self.assertEqual(r.status_code, 400)
        r = acc.post(f"{API}/student-payments/", {"student": student.pk, "amount": "4000", "received_date": "2026-09-20",
                                                   "refund_of": payment_id}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["kind"], "refund")
        self.assertEqual(acc.delete(f"{API}/student-payments/{payment_id}/").status_code, 405)
        r = acc.post(f"{API}/student-payments/{payment_id}/void/", {"reason": "ошибка"}, format="json")
        self.assertEqual(r.status_code, 400)  # сначала возвраты

    def test_salary_rule_cannot_be_edited_in_place(self):
        acc = self.client_for(self.accountant)
        rule = self.rule(self.profile(), R.FIXED, amount=D("30000"))
        self.assertEqual(acc.patch(f"{API}/salary-rules/{rule.pk}/", {"amount": "1"}, format="json").status_code, 405)
        self.assertEqual(acc.delete(f"{API}/salary-rules/{rule.pk}/").status_code, 405)
        r = acc.post(f"{API}/salary-rules/{rule.pk}/new-version/", {"effective_from": "2026-10-01", "amount": "35000"},
                     format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["previous_version"], rule.pk)
        rule.refresh_from_db()
        self.assertEqual(str(rule.effective_to), "2026-09-30")


class ReportTests(AccountingFixture):
    def setUp(self):
        super().setUp()
        self.rule(self.profile(), R.FIXED, amount=D("40000"))
        self.other = make_user("other", User.Role.ASSISTANT)
        self.rule(self.profile(self.other), R.REVENUE_PERCENT, percentage=D("10"), group=self.group_a)
        self.pay(self.student(), 30000, day(9, 4))
        self.a = self.calc(FIRST)
        self.b = self.calc(FIRST, self.other)
        self.c = self.calc(SECOND)
        for p in (self.a, self.b, self.c):
            approve_payroll(p, self.director)
        register_payment(self.a, amount=D("5000"), payment_date=day(9, 16), actor=self.accountant)
        voided, _ = register_payment(self.a, amount=D("7000"), payment_date=day(9, 16), actor=self.accountant)
        void_payment(voided, actor=self.accountant, reason="дубль")

    def xlsx(self, query):
        r = self.client_for(self.director).get(f"{API}/reports/payroll.xlsx?{query}")
        self.assertEqual(r.status_code, 200, r.content)
        ws = load_workbook(io.BytesIO(r.content)).active
        rows = [row for row in ws.iter_rows(values_only=True)]
        return ws, rows

    def test_excel_totals_and_voided_payments_excluded(self):
        ws, rows = self.xlsx("year=2026&month=9&period_type=FIRST_HALF")
        total = rows[-1]
        self.assertEqual(total[0], "Итого")
        self.assertEqual(D(str(total[4])), D("23000"))  # 20000 оклад + 3000 процент
        self.assertEqual(D(str(total[7])), D("5000"))  # отменённая выплата 7000 не учтена
        self.assertEqual(D(str(total[8])), D("18000"))
        self.assertIsNotNone(ws.auto_filter.ref)
        self.assertIsInstance(rows[4][4], (int, float))

    def test_period_filter(self):
        _, first = self.xlsx("year=2026&month=9&period_type=FIRST_HALF")
        _, month = self.xlsx("year=2026&month=9")
        self.assertEqual(len(first) - 5, 2)
        self.assertEqual(len(month) - 5, 3)
        self.assertEqual(D(str(month[-1][4])), D("43000"))

    def test_pdf_report_totals(self):
        from apps.accounting.services.report_service import build_report, resolve_periods

        periods, label = resolve_periods(year=2026, month=9, period_type=FIRST)
        report = build_report(periods, label=label)
        self.assertEqual(report.totals.accrued, D("23000"))
        self.assertEqual(report.totals.paid, D("5000"))
        self.assertEqual(report.totals.due, D("18000"))
        self.assertEqual(report.by_program, [("Prog SOFT", D("3000.00"))])
        r = self.client_for(self.accountant).get(f"{API}/reports/payroll.pdf?year=2026&month=9&period_type={FIRST}")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))
        r = self.client_for(self.accountant).get(f"{API}/payrolls/{self.a.pk}/report.pdf/")
        self.assertTrue(r.content.startswith(b"%PDF"))
        r = self.client_for(self.accountant).get(f"{API}/payrolls/{self.a.pk}/report.xlsx/")
        wb = load_workbook(io.BytesIO(r.content))
        pay_rows = list(wb["Выплаты"].iter_rows(values_only=True))
        self.assertEqual(D(str(pay_rows[-1][4])), D("15000"))  # остаток

    def test_export_access(self):
        self.assertEqual(self.client_for(self.other).get(f"{API}/reports/payroll.xlsx?year=2026&month=9").status_code, 403)
        self.assertEqual(self.client_for(self.admin).get(f"{API}/reports/payroll.xlsx?year=2026&month=9").status_code, 200)
