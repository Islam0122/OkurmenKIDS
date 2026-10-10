"""Модуль Payroll & Accounting: предварительная зарплата, плановые даты
выплат, история тарифов, уроки блоков, спорные начисления, аналитика,
отчёт по сотрудникам, автоматизация, сверка и перенос данных."""
from __future__ import annotations

import datetime as dt
import io
import threading
import unittest

from django.apps import apps as django_apps
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db import IntegrityError, connection, transaction
from django.test import TransactionTestCase, override_settings
from openpyxl import load_workbook

from apps.academy.models import GroupTeacher, Lesson, Student
from apps.accounting.models import (
    CourseCycle,
    CoursePriceVersion,
    CycleAccrual,
    CycleLesson,
    Department,
    EmployeeSalaryProfile,
    Payroll,
    PayrollAuditLog,
    PayrollPayment,
    PayrollPeriod,
    SalaryRule,
    SalaryType,
)
from apps.accounting.services import AccountingError, analytics, estimates, reconcile
from apps.accounting.services.approval_service import approve_payroll
from apps.accounting.services.cycles import review_accrual, sync_all
from apps.accounting.services.payment_service import register_payment, void_payment
from apps.accounting.services.payout import fixed_payday, planned_date_for_completion, planned_date_for_period
from apps.accounting.services.pricing import set_price
from apps.accounting.tests.base import FIRST, MONTH, SECOND, D, AccountingFixture, day, make_user
from apps.users.models import Subject, Teacher, User

R = SalaryRule.RuleType
API = "/api/v1/accounting"
OCT = (2026, 10)


def days(month, first, last, year=2026):
    return [dt.date(year, month, d) for d in range(first, last + 1)]


class PercentFixture(AccountingFixture):
    """10 студентов × 10 000 сом × 10 % — пример из ТЗ."""

    lessons_in_block = 12

    def setUp(self):
        super().setUp()
        self.settings = self.course_settings(price="10000", lessons=self.lessons_in_block, start=None)
        self.percent_profile = self.profile(salary_type=SalaryType.PERCENT)
        self.percent_rule = self.rule(self.percent_profile, R.PERCENT, percentage=D("10"))
        self.kids = [self.student(self.group_a) for _ in range(10)]

    def add_lessons(self, dates, group=None, **kw):
        with self.captureOnCommitCallbacks(execute=True):
            self.lessons(group or self.group_a, dates, **kw)

    def my_estimates(self, today=day(10, 10)):
        return estimates.for_employee(self.trainer_user, today=today)


# ---------------------------------------------------------------------------
# Плановые даты выплат
# ---------------------------------------------------------------------------

class PayoutDateTests(AccountingFixture):
    def test_trainer_payout_dates(self):
        cases = {
            day(10, 10): day(10, 15), day(10, 15): day(10, 15), day(10, 16): day(11, 1), day(10, 17): day(11, 1),
            day(10, 31): day(11, 1), day(10, 1): day(10, 15), day(12, 31): dt.date(2027, 1, 1),
            dt.date(2028, 2, 29): dt.date(2028, 3, 1),
        }
        for completed, expected in cases.items():
            self.assertEqual(planned_date_for_completion(completed), expected, completed)

    def test_period_planned_dates(self):
        self.assertEqual(planned_date_for_period(self.period(OCT, FIRST)), day(10, 15))
        self.assertEqual(planned_date_for_period(self.period((2026, 12), SECOND)), dt.date(2027, 1, 1))

    def test_fixed_salary_calendar_is_configured_not_invented(self):
        self.assertIsNone(fixed_payday(2026, 10))
        with override_settings(ACCOUNTING_FIXED_PAYDAY_DAY=31, ACCOUNTING_FIXED_PAYDAY_NEXT_MONTH=True):
            self.assertEqual(fixed_payday(2026, 10), day(11, 30))
            self.assertEqual(fixed_payday(2026, 12), dt.date(2027, 1, 31))
        with override_settings(ACCOUNTING_FIXED_PAYDAY_DAY=25, ACCOUNTING_FIXED_PAYDAY_NEXT_MONTH=False):
            self.assertEqual(planned_date_for_period(self.period(OCT, MONTH)), day(10, 25))

    def test_reaching_payout_date_does_not_mark_paid(self):
        self.rule(self.profile(), R.FIXED, amount=D("20000"))
        payroll = approve_payroll(self.calc(MONTH), self.director)
        payroll.refresh_from_db()
        self.assertEqual(payroll.status, Payroll.Status.APPROVED)
        self.assertEqual(payroll.total_paid, D("0"))


# ---------------------------------------------------------------------------
# Предварительная зарплата
# ---------------------------------------------------------------------------

class EstimateTests(PercentFixture):
    def test_formula_example_from_the_spec(self):
        self.add_lessons(days(10, 1, 8))
        [row] = self.my_estimates()
        self.assertEqual((row["student_count"], row["price_per_student"], row["percentage"], row["expected_amount"]),
                         (10, D("10000"), D("10"), D("10000.00")))
        self.assertEqual((row["lessons_done"], row["required_lessons"], row["lessons_remaining"]), (8, 12, 4))
        self.assertEqual(row["status"], estimates.BLOCK_IN_PROGRESS)
        self.assertEqual(row["status_display"], "Блок в процессе")
        self.assertFalse(CycleAccrual.objects.exists())

    def test_estimate_is_not_an_accrual_or_a_debt(self):
        self.add_lessons(days(10, 1, 8))
        data = self.client_for(self.trainer_user).get(f"{API}/my/salary/").json()
        self.assertEqual(data["totals"]["estimated"], "10000.00")
        self.assertEqual((data["totals"]["accrued"], data["totals"]["due"]), ("0.00", "0.00"))
        self.assertEqual(data["estimates"][0]["lessons_done"], 8)
        self.assertFalse(Payroll.objects.exists())
        summary = analytics.summary(2026, 10)
        self.assertEqual((summary["accrued"], summary["outstanding"]), (D("0"), D("0")))

    def test_projected_completion_and_payout_from_scheduled_lessons(self):
        self.add_lessons(days(10, 1, 8))
        self.add_lessons([day(10, 12), day(10, 14), day(10, 16), day(10, 19)], status=Lesson.Status.SCHEDULED)
        [row] = self.my_estimates(today=day(10, 10))
        self.assertEqual(row["projected_completion_date"], day(10, 19))
        self.assertEqual(row["expected_payment_date"], day(11, 1))

    def test_no_projection_without_enough_scheduled_lessons(self):
        self.add_lessons(days(10, 1, 8))
        [row] = self.my_estimates()
        self.assertIsNone(row["expected_payment_date"])

    def test_estimate_follows_group_price_and_rate_changes(self):
        self.add_lessons(days(9, 1, 3))
        self.leave(self.kids[0], day(10, 5))
        self.assertEqual(self.my_estimates()[0]["expected_amount"], D("9000.00"))
        set_price(course=self.course, price_per_student=D("12000"), effective_from=day(10, 1), reason="Новый прайс",
                  actor=self.accountant)
        self.assertEqual(self.my_estimates()[0]["expected_amount"], D("10800.00"))
        self.percent_rule.effective_to = day(9, 30)
        self.percent_rule.save()
        self.rule(self.percent_profile, R.PERCENT, percentage=D("15"), start=day(10, 1))
        self.assertEqual(self.my_estimates()[0]["expected_amount"], D("16200.00"))

    def test_trainer_change_shown_in_estimate(self):
        other = Teacher.objects.create(user=make_user("substitute"))
        self.add_lessons(days(10, 1, 4))
        program = GroupTeacher.objects.get(group=self.group_a)
        with self.captureOnCommitCallbacks(execute=True):
            Lesson.objects.create(group=self.group_a, group_teacher=program, teacher=other, subject=self.python,
                                  lesson_number=99, date=day(10, 5), start_time=dt.time(10), end_time=dt.time(11),
                                  status=Lesson.Status.COMPLETED)
        self.assertIn("разные тренеры", self.my_estimates()[0]["warnings"][0])

    def test_employee_sees_only_own_estimates(self):
        other = make_user("other")
        Teacher.objects.create(user=other)
        self.add_lessons(days(10, 1, 8))
        self.assertEqual(estimates.for_employee(other), [])
        r = self.client_for(other).get(f"{API}/my/estimates/?employee={self.trainer_user.pk}")
        self.assertEqual((r.status_code, r.json()["results"]), (200, []))
        mine = self.client_for(self.trainer_user).get(f"{API}/my/estimates/").json()["results"]
        self.assertEqual(mine[0]["expected_amount"], "10000.00")
        self.assertEqual(self.client_for(other).get(f"{API}/estimates/").status_code, 403)

    def test_accounting_sees_open_blocks_with_parameters(self):
        self.add_lessons(days(10, 1, 8))
        rows = self.client_for(self.accountant).get(f"{API}/estimates/").json()["results"]
        row = next(r for r in rows if r["group_name"] == "Group A")
        self.assertEqual((row["student_count"], row["price_per_student"], row["expected_amount"]),
                         (10, "10000.00", "10000.00"))
        b = next(r for r in rows if r["group_name"] == "Group B")
        self.assertEqual(b["status"], estimates.ESTIMATED)


class TwentyLessonEstimateTests(PercentFixture):
    lessons_in_block = 20

    def test_progress_out_of_twenty(self):
        self.add_lessons(days(10, 1, 8))
        [row] = self.my_estimates()
        self.assertEqual((row["lessons_done"], row["required_lessons"], row["lessons_remaining"]), (8, 20, 12))
        self.add_lessons(days(10, 9, 20))
        accrual = CycleAccrual.objects.get()
        self.assertEqual((accrual.lessons, accrual.amount, accrual.completed_on), (20, D("10000.00"), day(10, 20)))
        self.assertEqual(accrual.planned_payment_date, day(11, 1))


# ---------------------------------------------------------------------------
# Блоки: уроки, предметы, спорные начисления
# ---------------------------------------------------------------------------

class BlockTests(PercentFixture):
    def test_completed_block_links_exact_lessons(self):
        self.add_lessons(days(10, 1, 13))
        cycle = CourseCycle.objects.get(status="COMPLETED")
        links = list(cycle.cycle_lessons.order_by("position"))
        self.assertEqual(len(links), 12)
        self.assertEqual((links[0].lesson_date, links[-1].lesson_date), (day(10, 1), day(10, 12)))
        self.assertEqual(links[-1].lesson_ref, cycle.last_lesson_id)
        self.assertTrue(all(link.teacher_id == self.trainer.pk for link in links))

    def test_lesson_cannot_be_in_two_live_blocks(self):
        self.add_lessons(days(10, 1, 12))
        cycle = CourseCycle.objects.get(status="COMPLETED")
        link = cycle.cycle_lessons.first()
        other = CourseCycle.objects.create(group=self.group_b, course=self.course, number=1, status="COMPLETED",
                                           required_lessons=12)
        with self.assertRaises(IntegrityError), transaction.atomic():
            CycleLesson.objects.create(cycle=other, lesson_ref=link.lesson_ref, lesson_date=link.lesson_date,
                                       position=1)

    def test_cancelled_lessons_are_not_counted_and_corrections_keep_history(self):
        self.add_lessons(days(10, 1, 11))
        self.add_lessons([day(10, 12)], status=Lesson.Status.CANCELLED)
        self.assertFalse(CourseCycle.objects.filter(status="COMPLETED").exists())
        self.add_lessons([day(10, 13)])
        cycle = CourseCycle.objects.get(status="COMPLETED")
        lesson = Lesson.objects.filter(group=self.group_a, date=day(10, 5)).get()
        with self.captureOnCommitCallbacks(execute=True):
            lesson.status = Lesson.Status.CANCELLED
            lesson.save()
        cycle.refresh_from_db()
        self.assertEqual(cycle.status, CourseCycle.Status.INVALIDATED)
        self.assertFalse(cycle.cycle_lessons.filter(is_live=True).exists())
        self.assertEqual(cycle.cycle_lessons.count(), 12)  # история сохранена
        self.assertEqual(CycleAccrual.objects.get().status, CycleAccrual.Status.CANCELLED)

    def test_only_counted_subjects_form_the_block(self):
        english = Subject.objects.get_or_create(name="English")[0]
        self.settings.counted_subjects.set([self.python])
        program = GroupTeacher.objects.get(group=self.group_a)
        with self.captureOnCommitCallbacks(execute=True):
            for i, d in enumerate(days(10, 1, 6)):
                Lesson.objects.create(group=self.group_a, group_teacher=program, teacher=self.trainer,
                                      subject=english, lesson_number=200 + i, date=d, start_time=dt.time(12),
                                      end_time=dt.time(13), status=Lesson.Status.COMPLETED)
        self.add_lessons(days(10, 1, 11))
        self.assertFalse(CourseCycle.objects.filter(status="COMPLETED").exists())
        self.assertEqual(CourseCycle.objects.get(group=self.group_a).lessons_done, 11)

    def test_trainer_change_inside_block_requires_review(self):
        substitute_user = make_user("substitute")
        substitute = Teacher.objects.create(user=substitute_user)
        program = GroupTeacher.objects.get(group=self.group_a)
        self.add_lessons(days(9, 1, 6))
        with self.captureOnCommitCallbacks(execute=True):
            for i, d in enumerate(days(9, 7, 12)):
                Lesson.objects.create(group=self.group_a, group_teacher=program, teacher=substitute,
                                      subject=self.python, lesson_number=300 + i, date=d, start_time=dt.time(10),
                                      end_time=dt.time(11), status=Lesson.Status.COMPLETED)
        accrual = CycleAccrual.objects.get()
        self.assertEqual(accrual.status, CycleAccrual.Status.REVIEW_REQUIRED)
        self.assertIn("Смена тренера", accrual.review_reasons[0])
        payroll = self.calc(FIRST)
        self.assertEqual(payroll.total_accrued, D("0"))
        self.assertTrue(any("требует проверки" in e for e in payroll.errors))
        with self.assertRaises(AccountingError):
            approve_payroll(payroll, self.director)
        # Решение бухгалтера с основанием — начисление входит в расчёт.
        with self.assertRaises(AccountingError):
            review_accrual(accrual, confirm=True, reason="", actor=self.accountant)
        review_accrual(accrual, confirm=True, reason="Замена на 6 уроков согласована, оплата основному тренеру",
                       actor=self.accountant)
        payroll = self.calc(FIRST)
        self.assertEqual((payroll.total_accrued, payroll.errors), (D("10000.00"), []))
        self.assertTrue(PayrollAuditLog.objects.filter(action="review_confirm", reason__startswith="Замена").exists())

    def test_review_cancel_keeps_history_and_api_rights(self):
        self.add_lessons(days(9, 1, 24))  # второй блок 24.09 — тот же месяц
        second = CycleAccrual.objects.get(cycle__number=2)
        self.assertEqual(second.status, CycleAccrual.Status.REVIEW_REQUIRED)
        director = self.client_for(self.director)
        r = director.post(f"{API}/cycle-accruals/{second.pk}/review/", {"confirm": False, "reason": "x"}, format="json")
        self.assertEqual(r.status_code, 403)
        listed = self.client_for(self.accountant).get(f"{API}/cycle-accruals/?status=REVIEW_REQUIRED").json()
        self.assertEqual([a["id"] for a in listed["results"]], [second.pk])
        r = self.client_for(self.accountant).post(f"{API}/cycle-accruals/{second.pk}/review/",
                                                  {"confirm": False, "reason": "Один оплаченный месяц"}, format="json")
        self.assertEqual((r.status_code, r.json()["status"]), (200, CycleAccrual.Status.CANCELLED))
        self.assertEqual(self.calc(SECOND).total_accrued, D("0"))
        self.assertEqual(CycleAccrual.objects.count(), 2)

    def test_cannot_review_own_accrual(self):
        self.add_lessons(days(9, 1, 24))
        second = CycleAccrual.objects.get(cycle__number=2)
        with self.assertRaises(AccountingError) as ctx:
            review_accrual(second, confirm=True, reason="сам себе", actor=self.trainer_user)
        self.assertEqual(ctx.exception.code, "own_payroll")

    def test_rerunning_sync_does_not_duplicate(self):
        self.add_lessons(days(10, 1, 12))
        for _ in range(3):
            sync_all()
            call_command("payroll_autorun", "--calculate", "--date", "2026-10-13", stdout=io.StringIO())
        self.assertEqual((CourseCycle.objects.filter(status="COMPLETED").count(), CycleAccrual.objects.count(),
                          CycleLesson.objects.count()), (1, 1, 12))
        self.assertEqual(Payroll.objects.filter(employee=self.trainer_user).count(), 1)
        self.assertEqual(PayrollPeriod.objects.filter(year=2026, month=10).count(), 2)


# ---------------------------------------------------------------------------
# Тарифы
# ---------------------------------------------------------------------------

class PricingTests(PercentFixture):
    def test_new_price_applies_to_new_blocks_only(self):
        self.add_lessons(days(9, 1, 12))
        payroll = approve_payroll(self.calc(FIRST), self.director)
        set_price(course=self.course, price_per_student=D("15000"), effective_from=day(10, 1), reason="Повышение",
                  actor=self.accountant)
        self.add_lessons(days(10, 1, 12))
        old, new = CourseCycle.objects.filter(status="COMPLETED").order_by("number")
        self.assertEqual((old.course_price, new.course_price), (D("10000"), D("15000")))
        self.assertEqual(CycleAccrual.objects.get(cycle=new).amount, D("15000.00"))
        payroll.refresh_from_db()
        self.assertEqual(payroll.total_accrued, D("10000.00"))

    def test_history_is_kept_and_versions_are_immutable(self):
        first = set_price(course=self.course, price_per_student=D("10000"), effective_from=day(1, 1),
                          reason="Начальная", actor=self.accountant)
        second = set_price(course=self.course, price_per_student=D("11000"), effective_from=day(6, 1),
                           reason="Индексация", actor=self.accountant)
        first.refresh_from_db()
        self.assertEqual((first.effective_to, second.previous_version_id, second.created_by),
                         (day(5, 31), first.pk, self.accountant))
        with self.assertRaises(ValidationError):
            first.delete()
        with self.assertRaises(ValidationError):
            CoursePriceVersion.objects.all().delete()
        self.assertTrue(PayrollAuditLog.objects.filter(action="price_change", entity_id=second.pk,
                                                       reason="Индексация").exists())

    def test_rules(self):
        set_price(course=self.course, price_per_student=D("10000"), effective_from=day(1, 1), reason="a",
                  actor=self.accountant)
        with self.assertRaises(AccountingError):
            set_price(course=self.course, price_per_student=D("1"), effective_from=day(2, 1), reason=" ",
                      actor=self.accountant)
        with self.assertRaises(AccountingError):
            set_price(course=self.course, price_per_student=D("12000"), effective_from=day(1, 1), reason="a",
                      actor=self.accountant)
        self.add_lessons(days(9, 1, 12))
        with self.assertRaises(AccountingError) as ctx:  # задним числом на завершённый блок
            set_price(course=self.course, price_per_student=D("12000"), effective_from=day(9, 10), reason="b",
                      actor=self.accountant)
        self.assertEqual(ctx.exception.code, "retroactive")

    def test_api(self):
        acc = self.client_for(self.accountant)
        r = acc.post(f"{API}/pricing/", {"course": self.course.pk, "price_per_student": "12500",
                                         "effective_from": "2026-11-01", "reason": "Новый учебный год"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["created_by_name"], "Buh Test")
        r = acc.patch(f"{API}/course-settings/{self.settings.pk}/", {"price_per_student": "1"}, format="json")
        self.assertEqual((r.status_code, r.json()["code"]), (400, "use_pricing"))
        r = acc.patch(f"{API}/course-settings/{self.settings.pk}/", {"required_lessons": 20}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        history = acc.get(f"{API}/pricing/?course={self.course.pk}").json()["results"]
        self.assertEqual([v["price_per_student"] for v in history], ["12500.00"])
        self.assertEqual(self.client_for(self.director).get(f"{API}/pricing/").status_code, 200)
        self.assertEqual(self.client_for(self.director).post(f"{API}/pricing/", {}, format="json").status_code, 403)
        self.assertEqual(self.client_for(self.trainer_user).get(f"{API}/pricing/").status_code, 403)

    def test_creating_course_settings_starts_price_history(self):
        acc = self.client_for(self.accountant)
        r = acc.post(f"{API}/course-settings/", {"course": self.other_course.pk, "price_per_student": "9000",
                                                 "required_lessons": 20, "price_effective_from": "2026-01-01"},
                     format="json")
        self.assertEqual(r.status_code, 201, r.content)
        version = CoursePriceVersion.objects.get(course=self.other_course)
        self.assertEqual((version.price_per_student, version.effective_from), (D("9000"), day(1, 1)))


# ---------------------------------------------------------------------------
# Частичные выплаты (пример из ТЗ) и конкуренция
# ---------------------------------------------------------------------------

class PartialPaymentTests(AccountingFixture):
    def setUp(self):
        super().setUp()
        self.rule(self.profile(), R.FIXED, amount=D("15000"))
        self.payroll = approve_payroll(self.calc(MONTH), self.director)

    def pay(self, amount, **kw):
        return register_payment(self.payroll, amount=D(amount), payment_date=day(10, 1), actor=self.accountant, **kw)[0]

    def test_example_from_the_spec(self):
        self.pay("10000")
        second = self.pay("3000")
        self.payroll.refresh_from_db()
        self.assertEqual((self.payroll.amount_due, self.payroll.status), (D("2000"), Payroll.Status.PARTIALLY_PAID))
        with self.assertRaises(AccountingError):
            self.pay("2000.01")
        self.pay("2000")
        self.payroll.refresh_from_db()
        self.assertEqual((self.payroll.amount_due, self.payroll.status), (D("0"), Payroll.Status.PAID))
        void_payment(second, actor=self.accountant, reason="Ошибочный перевод")
        self.payroll.refresh_from_db()
        self.assertEqual((self.payroll.amount_due, self.payroll.status), (D("3000"), Payroll.Status.PARTIALLY_PAID))
        self.assertEqual(PayrollPayment.objects.count(), 3)  # отменённая выплата осталась в истории
        self.assertEqual(PayrollPayment.objects.get(pk=second.pk).status, PayrollPayment.Status.VOID)

    def test_idempotency_key_over_api(self):
        acc = self.client_for(self.accountant)
        body = {"amount": "5000", "payment_date": "2026-10-01", "payment_method": "bank"}
        a = acc.post(f"{API}/payrolls/{self.payroll.pk}/payments/", body, format="json", HTTP_IDEMPOTENCY_KEY="k-1")
        b = acc.post(f"{API}/payrolls/{self.payroll.pk}/payments/", body, format="json", HTTP_IDEMPOTENCY_KEY="k-1")
        self.assertEqual((a.status_code, b.status_code, a.json()["id"]), (201, 200, b.json()["id"]))
        self.assertEqual(PayrollPayment.objects.count(), 1)
        neg = acc.post(f"{API}/payrolls/{self.payroll.pk}/payments/", {**body, "amount": "-1"}, format="json")
        self.assertEqual(neg.status_code, 400)


@unittest.skipUnless(connection.vendor == "postgresql", "нужны настоящие параллельные транзакции (PostgreSQL)")
class ConcurrentPaymentTests(TransactionTestCase):
    """Параллельные выплаты не превышают остаток (select_for_update)."""

    def test_parallel_payments_cannot_exceed_due(self):
        accountant = make_user("buh_c", User.Role.ACCOUNTANT)
        director = make_user("boss_c", User.Role.DIRECTOR)
        employee = make_user("emp_c", User.Role.ASSISTANT)
        profile = EmployeeSalaryProfile.objects.create(employee=employee, salary_type=SalaryType.FIXED,
                                                       effective_from=day(1, 1))
        SalaryRule.objects.create(employee_profile=profile, rule_type=R.FIXED, amount=D("10000"),
                                  calculation_method="MONTHLY", effective_from=day(1, 1))
        from apps.accounting.services.payroll_calculator import calculate_payroll
        from apps.accounting.services.periods import get_or_create_period

        period = get_or_create_period(2026, 9, MONTH, accountant)[0]
        payroll = approve_payroll(calculate_payroll(period, employee, accountant), director)
        barrier = threading.Barrier(4)
        results = []

        def attempt(i):
            try:
                barrier.wait()
                register_payment(payroll, amount=D("4000"), payment_date=day(9, 30), actor=accountant,
                                 reference=f"tx-{i}")
                results.append("ok")
            except AccountingError as exc:
                results.append(exc.code)
            finally:
                connection.close()

        threads = [threading.Thread(target=attempt, args=(i,)) for i in range(4)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        payroll.refresh_from_db()
        self.assertEqual(results.count("ok"), 2)
        self.assertEqual(results.count("exceeds_due"), 2)
        self.assertEqual((payroll.total_paid, payroll.amount_due), (D("8000"), D("2000")))


# ---------------------------------------------------------------------------
# Аналитика и отчёт по сотрудникам
# ---------------------------------------------------------------------------

class AnalyticsTests(PercentFixture):
    def setUp(self):
        super().setUp()
        self.percent_profile.department = Department.IT
        self.percent_profile.save()
        self.assistant = make_user("helper", User.Role.ASSISTANT)
        self.rule(self.profile(self.assistant), R.FIXED, amount=D("20000"))
        self.add_lessons(days(9, 1, 12))
        self.trainer_payroll = approve_payroll(self.calc(FIRST), self.director)
        self.assistant_payroll = approve_payroll(self.calc(MONTH, user=self.assistant), self.director)
        register_payment(self.trainer_payroll, amount=D("4000"), payment_date=day(9, 30), actor=self.accountant)
        register_payment(self.assistant_payroll, amount=D("20000"), payment_date=day(10, 2), actor=self.accountant)
        self.add_lessons(days(10, 1, 8))  # открытый блок — оценка, не начисление

    def test_definitions(self):
        sep = analytics.summary(2026, 9)
        self.assertEqual((sep["accrued"], sep["paid"], sep["outstanding"], sep["employees"]),
                         (D("30000.00"), D("24000"), D("6000.00"), 2))
        self.assertEqual(sep["cash_paid_in_month"], D("4000"))  # перевод 02.10 — в октябре
        octo = analytics.summary(2026, 10)
        self.assertEqual((octo["accrued"], octo["cash_paid_in_month"]), (D("0"), D("20000")))
        self.assertEqual(octo["previous_month"]["accrued"], D("30000.00"))
        self.assertEqual(octo["accrued_change"], D("-30000.00"))
        deps = {d["department"]: d for d in sep["by_department"]}
        self.assertEqual((deps["IT"]["accrued"], deps["ASSISTANT"]["accrued"], deps["ENGLISH"]["accrued"]),
                         (D("10000.00"), D("20000.00"), D("0")))
        self.assertEqual(len(sep["by_department"]), 6)
        self.assertEqual([r["payroll_id"] for r in sep["upcoming_payments"]], [self.trainer_payroll.pk])
        self.assertEqual(sep["upcoming_payments"][0]["planned_payment_date"], day(9, 15))

    def test_filters_and_unapproved_kept_apart(self):
        self.calc(SECOND)  # рассчитан, не утверждён — не входит в «начислено»
        it = analytics.summary(2026, 9, {"department": "IT"})
        self.assertEqual((it["accrued"], it["employees"]), (D("10000.00"), 1))
        self.assertEqual(analytics.summary(2026, 9, {"salary_type": "FIXED"})["accrued"], D("20000.00"))
        self.assertEqual(analytics.summary(2026, 9, {"group": self.group_a.pk})["accrued"], D("10000.00"))

    def test_api_and_rights(self):
        for user in (self.director, self.accountant, self.admin):
            r = self.client_for(user).get(f"{API}/analytics/summary/?year=2026&month=9")
            self.assertEqual((r.status_code, r.json()["accrued"]), (200, "30000.00"))
        self.assertEqual(self.client_for(self.director).get(
            f"{API}/analytics/by-department/?year=2026&month=9").json()["results"][0]["department"], "IT")
        for path in ("analytics/summary/", "analytics/by-department/", "reports/teachers/", "reports/teachers.xlsx"):
            self.assertEqual(self.client_for(self.trainer_user).get(f"{API}/{path}").status_code, 403, path)
            self.assertEqual(self.client_class().get(f"{API}/{path}").status_code, 401, path)
        bad = self.client_for(self.director).get(f"{API}/analytics/summary/?department=SPACE")
        self.assertEqual(bad.status_code, 400)

    def test_teacher_report_does_not_double_count_payments(self):
        report = analytics.teacher_report(2026, 9)
        self.assertEqual((report["totals"]["total"], report["totals"]["paid"], report["totals"]["due"]),
                         (D("30000.00"), D("24000"), D("6000.00")))
        trainer = next(p for p in report["payrolls"] if p["employee"] == self.trainer_user.pk)
        line = trainer["lines"][0]
        self.assertEqual((line["students"], line["price_per_student"], line["percentage"], line["lessons_done"],
                          line["target_lessons"], line["accrued"], line["group_name"]),
                         (10, D("10000"), D("10"), 12, 12, D("10000.00"), "Group A"))
        self.assertEqual(trainer["planned_payment_date"], day(9, 15))
        deps = {d["department"]: d for d in report["by_department"]}
        self.assertEqual(sum((d["paid"] for d in deps.values()), D(0)), D("24000"))

    def test_teacher_report_api_and_exports(self):
        director = self.client_for(self.director)
        r = director.get(f"{API}/reports/teachers/?year=2026&month=9&page_size=1")
        body = r.json()
        self.assertEqual((r.status_code, body["count"], len(body["results"])), (200, 2, 1))
        self.assertEqual(body["totals"]["total"], "30000.00")
        x = director.get(f"{API}/reports/teachers.xlsx?year=2026&month=9")
        self.assertEqual(x.status_code, 200)
        ws = load_workbook(io.BytesIO(x.content))["Сотрудники"]
        totals = [row for row in ws.iter_rows(values_only=True) if row and row[0] == "Итого"]
        self.assertEqual([float(v) for v in totals[0][10:13]], [30000.0, 24000.0, 6000.0])
        p = director.get(f"{API}/reports/teachers.pdf?year=2026&month=9")
        self.assertEqual((p.status_code, p["Content-Type"], p.content[:4]), (200, "application/pdf", b"%PDF"))

    def test_audit_log_filters(self):
        r = self.client_for(self.director).get(f"{API}/audit-log/?action=approve&date_from=2000-01-01")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(all(row["action"] == "approve" for row in r.json()["results"]))
        self.assertEqual(self.client_for(self.director).delete(f"{API}/audit-log/{r.json()['results'][0]['id']}/")
                         .status_code, 405)

    def test_payroll_shows_department_and_planned_date(self):
        row = self.client_for(self.accountant).get(f"{API}/payrolls/{self.trainer_payroll.pk}/").json()
        self.assertEqual((row["department"], row["planned_payment_date"]), ("IT", "2026-09-15"))
        assistant = self.client_for(self.accountant).get(f"{API}/employees/?year=2026&month=9&department=ASSISTANT")
        self.assertEqual([r["employee"] for r in assistant.json()["results"]], [self.assistant.pk])


# ---------------------------------------------------------------------------
# Сверка и перенос данных
# ---------------------------------------------------------------------------

class ReconcileAndMigrationTests(PercentFixture):
    def test_reconcile_clean_and_detects_mismatch(self):
        self.add_lessons(days(9, 1, 12))
        payroll = approve_payroll(self.calc(FIRST), self.director)
        report = reconcile.run()
        self.assertEqual(reconcile.problems_count(report), 0, report)
        self.assertEqual(report["totals"]["accrued"], "10000.00")
        Payroll.objects.filter(pk=payroll.pk).update(total_paid=D("1"))
        self.assertEqual(len(reconcile.run()["payroll_totals_mismatch"]), 1)
        with self.assertRaises(SystemExit):
            call_command("payroll_reconcile", stdout=io.StringIO())

    def _migration(self):
        import importlib

        return importlib.import_module("apps.accounting.migrations.0006_payroll_backfill")

    def test_backfill_is_repeatable_and_keeps_ids(self):
        mig = self._migration()
        self.add_lessons(days(9, 1, 12))
        cycle = CourseCycle.objects.get(status="COMPLETED")
        CycleLesson.objects.filter(cycle=cycle).update(is_live=False)  # как будто связей не было
        CycleLesson.objects.all()._raw_delete(CycleLesson.objects.db)
        for _ in range(2):
            mig.backfill_prices(django_apps, None)
            mig.backfill_cycle_lessons(django_apps, None)
        version = CoursePriceVersion.objects.get(course=self.course)
        self.assertEqual((version.price_per_student, version.is_migrated), (D("10000"), True))
        links = CycleLesson.objects.filter(cycle=cycle)
        self.assertEqual((links.count(), links.filter(backfilled=True).count()), (12, 12))
        self.assertEqual(links.order_by("-position").first().lesson_ref, cycle.last_lesson_id)
        self.assertEqual(reconcile.run()["migrated_prices_to_verify"][0]["version_id"], version.pk)

    def test_ambiguous_history_is_reported_not_guessed(self):
        mig = self._migration()
        self.add_lessons(days(9, 1, 12))
        CycleLesson.objects.all()._raw_delete(CycleLesson.objects.db)
        Lesson.objects.filter(group=self.group_a, date=day(9, 3)).update(status=Lesson.Status.CANCELLED)
        mig.backfill_cycle_lessons(django_apps, None)
        self.assertFalse(CycleLesson.objects.exists())
        self.assertEqual(len(reconcile.run()["completed_cycles_without_lesson_links"]), 1)


class DepartmentTests(AccountingFixture):
    def test_default_department_follows_role(self):
        lead = make_user("lead", User.Role.TEAM_LEAD)
        self.assertEqual(self.profile(lead).effective_department, Department.TEAM_LEAD)
        self.assertEqual(self.profile().effective_department, Department.OTHER)

    def test_payroll_keeps_department_snapshot(self):
        profile = self.profile()
        profile.department = Department.ENGLISH
        profile.save()
        self.rule(profile, R.FIXED, amount=D("1000"))
        payroll = self.calc(MONTH)
        profile.department = Department.IT
        profile.save()
        payroll.refresh_from_db()
        self.assertEqual(payroll.department, Department.ENGLISH)

    def test_student_counting_ignores_departed_and_duplicates(self):
        self.course_settings(price="10000", lessons=12, start=None)
        self.rule(self.profile(salary_type=SalaryType.PERCENT), R.PERCENT, percentage=D("10"))
        kids = [self.student(self.group_a) for _ in range(3)]
        self.leave(kids[0], day(9, 5))
        self.transfer(kids[1], self.group_b, day(9, 5))
        Student.objects.create(first_name="Late", last_name="Test", group=self.group_a, enrollment_date=day(11, 1))
        with self.captureOnCommitCallbacks(execute=True):
            self.lessons(self.group_a, days(9, 1, 12))
        self.assertEqual(CourseCycle.objects.get(status="COMPLETED").student_count, 1)
