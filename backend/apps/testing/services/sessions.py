"""Sessions as a business entity: schedule sync for lists, phase filters,
result rows and their Excel/PDF exports. Used by the admin «Сессии»
section (session_admin_views.py) and the teacher API (teacher_api.py)."""
from __future__ import annotations

import io
from dataclasses import dataclass
from xml.sax.saxutils import escape
from datetime import datetime

from django.db.models import Q, QuerySet
from django.utils import timezone

from ..models import (
    AttemptStatus,
    SessionPhase,
    SessionStatus,
    StudentAttempt,
    TestSession,
)
from .grading import attempt_score

PHASE_FILTERS: dict[str, Q] = {
    SessionPhase.DRAFT: Q(status=SessionStatus.CREATED, scheduled_start__isnull=True),
    SessionPhase.SCHEDULED: Q(status=SessionStatus.CREATED, scheduled_start__isnull=False),
    SessionPhase.ACTIVE: Q(status__in=[SessionStatus.RUNNING, SessionStatus.PAUSED]),
    SessionPhase.FINISHED: Q(status__in=[SessionStatus.FINISHED, SessionStatus.EXPIRED]),
    SessionPhase.CANCELLED: Q(status=SessionStatus.CANCELLED),
}


def sync_due_sessions(queryset: QuerySet | None = None, now=None) -> None:
    """Bring stored statuses up to date before filtering by them: start
    scheduled sessions whose time has come, close those past their end,
    persist lazy expiry. Only rows that need it are touched."""
    now = now or timezone.now()
    sessions = queryset if queryset is not None else TestSession.objects.all()
    due = sessions.filter(
        Q(status=SessionStatus.CREATED, scheduled_start__lte=now)
        | Q(status__in=[SessionStatus.RUNNING, SessionStatus.PAUSED], scheduled_end__lte=now)
        | Q(status=SessionStatus.RUNNING, session_type="exam", expires_at__lte=now)
    ).select_related("test")
    for session in due:
        session.sync_schedule(now)
        if session.status == SessionStatus.RUNNING and session.effective_status_at(now) == SessionStatus.EXPIRED:
            session.expire()


def filter_by_phase(queryset: QuerySet, phase: str) -> QuerySet:
    return queryset.filter(PHASE_FILTERS[phase]) if phase in PHASE_FILTERS else queryset


def session_date(session: TestSession) -> datetime | None:
    """When the session is (scheduled to be) held."""
    return session.scheduled_start or session.started_at or session.created_at


def display_title(session: TestSession) -> str:
    return session.title or session.test.title


# ---------------------------------------------------------------------------
# Results
# ---------------------------------------------------------------------------

@dataclass
class ResultRow:
    attempt: StudentAttempt
    student_name: str
    status: str
    status_label: str
    percent: float | None
    earned: int | None
    possible: int | None
    correct: int | None
    wrong: int | None
    pending: int | None
    duration_seconds: int | None
    finished_at: datetime | None
    passed: bool | None


ATTEMPT_STATUS_LABELS = {
    AttemptStatus.ACTIVE: "Проходит",
    AttemptStatus.FINISHED: "Завершил",
    AttemptStatus.EXPIRED: "Время истекло",
}


def result_rows(session: TestSession) -> list[ResultRow]:
    """Every attempt of the session (newest first per student), with its
    score breakdown. Active attempts show no score (answers stay hidden
    while the exam is running)."""
    rows = []
    attempts = session.attempts.select_related("student").prefetch_related("answers").order_by("student_name", "-started_at")
    passing = session.test.passing_score
    for attempt in attempts:
        finished = attempt.status == AttemptStatus.FINISHED
        score = attempt_score(attempt) if finished and attempt.question_ids else None
        if score:
            wrong = score.total_questions - score.correct - score.pending
            percent = score.percent
        else:
            wrong = None
            percent = attempt.score if finished else None
        rows.append(ResultRow(
            attempt=attempt,
            student_name=attempt.student_name,
            status=attempt.status,
            status_label=ATTEMPT_STATUS_LABELS.get(attempt.status, attempt.status),
            percent=percent,
            earned=score.earned if score else None,
            possible=score.possible if score else None,
            correct=score.correct if score else None,
            wrong=wrong,
            pending=score.pending if score else None,
            duration_seconds=int(attempt.duration_seconds) if attempt.duration_seconds is not None else None,
            finished_at=attempt.finished_at,
            passed=(None if percent is None or (score and score.pending) else percent >= passing),
        ))
    return rows


def format_duration(seconds: int | None) -> str:
    if seconds is None:
        return "—"
    minutes, sec = divmod(int(seconds), 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours}:{minutes:02d}:{sec:02d}" if hours else f"{minutes}:{sec:02d}"


def _export_rows(session: TestSession) -> list[list[str]]:
    rows = []
    for number, r in enumerate(result_rows(session), start=1):
        rows.append([
            str(number),
            r.student_name,
            r.status_label,
            "—" if r.percent is None else f"{r.percent:g}%",
            "—" if r.earned is None else f"{r.earned} / {r.possible}",
            "—" if r.correct is None else str(r.correct),
            "—" if r.wrong is None else str(r.wrong),
            format_duration(r.duration_seconds),
            timezone.localtime(r.finished_at).strftime("%d.%m.%Y %H:%M") if r.finished_at else "—",
        ])
    return rows


EXPORT_HEADERS = ["№", "Студент", "Статус", "Процент", "Баллы", "Верно", "Неверно", "Время", "Завершил"]


def _subtitle(session: TestSession) -> str:
    parts = [f"Тест: {session.test.title}"]
    if session.group_id:
        parts.append(f"Группа: {session.group.name}")
    when = session_date(session)
    if when:
        parts.append(timezone.localtime(when).strftime("%d.%m.%Y %H:%M"))
    parts.append(f"Ключ: {session.key}")
    return " · ".join(parts)


def export_results_excel(session: TestSession) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    ws["A1"] = f"Результаты сессии: {display_title(session)}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = _subtitle(session)
    ws["A2"].font = Font(color="68736C")
    header_row = 4
    fill = PatternFill("solid", fgColor="E7F3EC")
    for col, title in enumerate(EXPORT_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True, color="1F5F3C")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(_export_rows(session), start=header_row + 1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    for col, width in enumerate([5, 32, 16, 10, 10, 8, 9, 10, 18], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_results_pdf(session: TestSession) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from apps.academy.services.monthly_report_pdf import (
        _BOLD, _REGULAR, BORDER, BRAND_DARK, BRAND_SOFT, INK, INK_SECONDARY, _ensure_fonts,
    )

    _ensure_fonts()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36,
                            title=f"Результаты: {display_title(session)}")
    title_style = ParagraphStyle("t", fontName=_BOLD, fontSize=16, leading=20, textColor=INK)
    sub_style = ParagraphStyle("s", fontName=_REGULAR, fontSize=9.5, leading=13, textColor=INK_SECONDARY)
    cell_style = ParagraphStyle("c", fontName=_REGULAR, fontSize=9, leading=11, textColor=INK)
    rows = _export_rows(session)
    data = [EXPORT_HEADERS] + [[Paragraph(escape(v), cell_style) if i == 1 else v for i, v in enumerate(row)] for row in rows]
    table = Table(data, repeatRows=1, colWidths=[28, 210, 90, 60, 60, 50, 55, 60, 100])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), _BOLD),
        ("FONTNAME", (0, 1), (-1, -1), _REGULAR),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (-1, 0), BRAND_DARK),
        ("BACKGROUND", (0, 0), (-1, 0), BRAND_SOFT),
        ("LINEBELOW", (0, 0), (-1, -1), 0.5, BORDER),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story = [
        Paragraph(escape(f"Результаты сессии: {display_title(session)}"), title_style),
        Spacer(1, 4),
        Paragraph(escape(_subtitle(session)), sub_style),
        Spacer(1, 14),
        table if rows else Paragraph("Попыток пока нет.", sub_style),
    ]
    doc.build(story)
    return buffer.getvalue()
