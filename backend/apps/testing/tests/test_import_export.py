"""«Тесты»: import / export of tests and questions (services/import_export.py,
io_admin_views.py), «Дублировать», and the list page's toolbar and cards."""
from __future__ import annotations

import csv
import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.testing.models import Question, QuestionType, Test, TestLevel, TestStatus
from apps.testing.services import questions as svc
from apps.testing.services import import_export
from apps.testing.services.question_rules import OptionData, QuestionData
from apps.testing.services.question_selector import _serialize_question
from apps.users.models import Subject, User


def csv_file(rows: list[list], name="data.csv") -> SimpleUploadedFile:
    buffer = io.StringIO()
    csv.writer(buffer).writerows(rows)
    return SimpleUploadedFile(name, buffer.getvalue().encode("utf-8-sig"), content_type="text/csv")


def xlsx_file(rows: list[list], name="data.xlsx") -> SimpleUploadedFile:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


def read_csv(response) -> list[dict]:
    return list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))


Q_HEADER = import_export.question_columns()


def q_row(**values) -> list:
    return [values.get(column, "") for column in Q_HEADER]


class Fixture(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(username="root", email="root@okurmen.kg", password="x")
        self.client.force_login(self.admin)
        self.python = Subject.objects.get_or_create(name="Python")[0]
        self.test = Test.objects.create(title="Python Basics", subject=self.python, level=TestLevel.EASY)

    def url(self, name, *args):
        return reverse(f"admin:{name}", args=args)

    def add_choice(self, text="Что такое Python?", test=None):
        return svc.save_question(test or self.test, QuestionData(
            QuestionType.SINGLE_CHOICE, text, options=[OptionData("Язык", True), OptionData("Змея")],
        ))


class TestsExportTests(Fixture):
    def setUp(self):
        super().setUp()
        self.add_choice()
        self.html = Test.objects.create(title="HTML", status=TestStatus.ARCHIVED, level=TestLevel.HARD)

    def export(self, **params):
        response = self.client.get(self.url("testing_test_export"), {"format": "csv", **params})
        self.assertEqual(response.status_code, 200)
        return read_csv(response)

    def test_all_selected_and_filtered(self):
        rows = self.export(scope="all")
        self.assertEqual([r["title"] for r in rows], ["HTML", "Python Basics"])
        python = rows[1]
        self.assertEqual((python["subject"], python["level"], python["questions_count"]), ("Python", "easy", "1"))
        self.assertEqual([r["title"] for r in self.export(scope="selected", ids=[str(self.html.pk)])], ["HTML"])
        self.assertEqual([r["title"] for r in self.export(scope="filtered", status="archived")], ["HTML"])
        self.assertEqual([r["title"] for r in self.export(scope="filtered", subject=str(self.python.pk))], ["Python Basics"])

    def test_selected_without_ids_goes_back_with_a_message(self):
        response = self.client.get(self.url("testing_test_export"), {"format": "csv", "scope": "selected"})
        self.assertRedirects(response, self.url("testing_test_changelist"))

    def test_xlsx_and_form_page(self):
        response = self.client.get(self.url("testing_test_export"), {"format": "xlsx"})
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet["A1"].value, "title")
        self.assertEqual(sheet.max_row, 3)
        page = self.client.get(self.url("testing_test_export"))
        self.assertContains(page, "Экспорт тестов")
        self.assertNotContains(page, "Только выбранные")


class TestsImportTests(Fixture):
    HEADER = ["title", "subject", "level", "status", "description", "image_url", "time_limit_minutes", "max_attempts", "passing_score"]

    def run_import(self, rows, **kw):
        return import_export.import_tests(csv_file([self.HEADER, *rows]), **kw)

    def test_creates_valid_rows_and_reports_bad_ones(self):
        report = self.run_import([
            ["JavaScript", "", "Средний", "draft", "Основы", "https://example.com/js.png", "30", "", "70"],
            ["", "", "", "", "", "", "", "", ""],                 # blank row: skipped by the reader
            ["No title row", "Химия", "", "", "", "", "", "", ""],  # unknown subject
            ["Bad level", "", "super", "", "", "", "", "", ""],
            ["Bad image", "", "", "", "", "javascript:alert(1)", "", "", ""],
            ["Too long", "", "", "", "", "", "9999", "", ""],
            ["Live", "", "", "active", "", "", "", "", ""],       # active without questions
            ["JavaScript", "", "", "", "", "", "", "", ""],       # twice in the file
        ])
        self.assertEqual((report.created, report.updated, report.error_rows), (1, 0, 6))
        js = Test.objects.get(title="JavaScript")
        self.assertEqual((js.level, js.status, js.time_limit_minutes, js.passing_score), ("medium", "draft", 30, 70))
        messages = {row.row: " ".join(row.errors) for row in report.errors}
        self.assertIn("Химия", messages[4])
        self.assertIn("super", messages[5])
        self.assertIn("http", messages[6])
        self.assertIn("Время", messages[7])
        self.assertIn("активным", messages[8])
        self.assertIn("строка 2", messages[9])
        self.assertFalse(Test.objects.filter(title__in=["Bad level", "Live"]).exists())

    def test_existing_title_is_skipped_unless_update_is_allowed(self):
        report = self.run_import([["python basics", "", "hard", "", "Новое", "", "", "", ""]])
        self.assertEqual((report.created, len(report.skipped)), (0, 1))
        self.assertEqual(Test.objects.filter(title__iexact="python basics").count(), 1)
        self.test.refresh_from_db()
        self.assertEqual(self.test.level, "easy")
        report = self.run_import([["python basics", "", "hard", "", "Новое", "", "", "", ""]], update_existing=True)
        self.assertEqual(report.updated, 1)
        self.test.refresh_from_db()
        self.assertEqual((self.test.level, self.test.description, self.test.title), ("hard", "Новое", "Python Basics"))

    def test_xlsx_and_missing_title_column(self):
        report = import_export.import_tests(xlsx_file([["Title", "Level"], ["Git", "easy"]]))
        self.assertEqual(report.created, 1)
        report = import_export.import_tests(csv_file([["name"], ["Git"]]))
        self.assertEqual((report.created, report.errors[0].row), (0, 1))

    def test_import_view_shows_the_report(self):
        upload = csv_file([self.HEADER, ["Docker", "", "", "", "", "", "", "", ""], ["Bad", "", "x", "", "", "", "", "", ""]])
        response = self.client.post(self.url("testing_test_import"), {"file": upload})
        self.assertEqual(response.status_code, 200)
        self.assertEqual((response.context["report"].created, response.context["report"].error_rows), (1, 1))
        self.assertContains(response, "Результат импорта")
        self.assertContains(response, "Строк с ошибками")

    def test_import_view_rejects_wrong_files(self):
        response = self.client.post(self.url("testing_test_import"), {"file": SimpleUploadedFile("tests.txt", b"x")})
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "CSV или XLSX", status_code=400)
        broken = SimpleUploadedFile("tests.xlsx", b"not a zip")
        self.assertContains(self.client.post(self.url("testing_test_import"), {"file": broken}), "Не удалось прочитать", status_code=400)
        self.assertEqual(self.client.post(self.url("testing_test_import"), {}).status_code, 400)

    def test_templates(self):
        response = self.client.get(self.url("testing_test_import_template"), {"format": "xlsx"})
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Тесты", "Пример", "Справка"])
        self.assertEqual([c.value for c in workbook["Тесты"][1]], self.HEADER)
        self.assertEqual(workbook["Тесты"].max_row, 1)
        response = self.client.get(self.url("testing_test_import_template"), {"format": "csv"})
        self.assertEqual(response.content.decode("utf-8-sig").strip(), ",".join(self.HEADER))


class QuestionsImportTests(Fixture):
    def run_import(self, rows, **kw):
        return import_export.import_questions(self.test, csv_file([Q_HEADER, *rows]), **kw)

    def test_all_question_types(self):
        report = self.run_import([
            q_row(question_type="single_choice", text="2 ** 3?", option_1="6", option_2="8", correct_answer="2",
                  points="2", explanation="2 ** 3 = 8", image_url="https://example.com/q.png"),
            q_row(question_type="Несколько вариантов", text="Изменяемые типы?", option_1="list", option_2="tuple",
                  option_3="dict", correct_answer="1;3"),
            q_row(question_type="multiple_choice", text="По тексту?", option_1="A", option_2="B", correct_answer="b"),
            q_row(test="Python Basics", question_type="text", text="Функция вывода?", correct_answer="print;print()"),
            q_row(question_type="code", text="sum(a, b)", language="python", correct_answer="3", hint="return"),
        ])
        self.assertEqual((report.created, report.error_rows), (5, 0), report.errors)
        single = Question.objects.get(text="2 ** 3?")
        self.assertEqual([(o.text, o.is_correct) for o in single.options.order_by("order")], [("6", False), ("8", True)])
        self.assertEqual((single.points, single.image_url, single.metadata["explanation"]), (2, "https://example.com/q.png", "2 ** 3 = 8"))
        self.assertNotIn("metadata", _serialize_question(single))  # the explanation is not sent to students
        multi = Question.objects.get(text="Изменяемые типы?")
        self.assertEqual([o.text for o in multi.options.filter(is_correct=True).order_by("order")], ["list", "dict"])
        self.assertEqual(Question.objects.get(text="Функция вывода?").correct_answers, ["print", "print()"])
        code = Question.objects.get(text="sum(a, b)")
        self.assertEqual((code.language, code.code_tests, code.hint), ("python", [{"input": "", "expected_output": "3"}], "return"))
        self.assertEqual(list(self.test.questions.order_by("order").values_list("order", flat=True)), [1, 2, 3, 4, 5])

    def test_invalid_rows_are_reported_and_not_saved(self):
        report = self.run_import([
            q_row(question_type="essay", text="?"),
            q_row(question_type="single_choice", text=""),
            q_row(question_type="single_choice", text="Без правильного", option_1="A", option_2="B"),
            q_row(question_type="single_choice", text="Два правильных", option_1="A", option_2="B", correct_answer="1;2"),
            q_row(question_type="single_choice", text="Нет варианта 5", option_1="A", option_2="B", correct_answer="5"),
            q_row(question_type="single_choice", text="Один вариант", option_1="A", correct_answer="1"),
            q_row(question_type="text", text="Без ответа"),
            q_row(question_type="code", text="Без языка"),
            q_row(question_type="text", text="Баллы", correct_answer="x", points="500"),
            q_row(test="HTML", question_type="text", text="Чужой", correct_answer="x"),
            q_row(question_type="text", text="Картинка", correct_answer="x", image_url="ftp://x"),
            q_row(question_type="text", text="Варианты у текста", correct_answer="x", option_1="A"),
        ])
        self.assertEqual((report.created, report.error_rows), (0, 12))
        self.assertFalse(self.test.questions.exists())
        joined = {row.row: " ".join(row.errors) for row in report.errors}
        self.assertIn("essay", joined[2])
        self.assertIn("текст", joined[3])
        self.assertIn("правильный", joined[4])
        self.assertIn("только один", joined[5])
        self.assertIn("№5", joined[6])
        self.assertIn("два варианта", joined[7])
        self.assertIn("правильный ответ", joined[8])
        self.assertIn("язык", joined[9])
        self.assertIn("1 до 100", joined[10])
        self.assertIn("HTML", joined[11])
        self.assertIn("http", joined[12])
        self.assertIn("option_N", joined[13])

    def test_duplicates_are_skipped_unless_update_is_allowed(self):
        self.add_choice("Что такое Python?")
        rows = [
            q_row(question_type="single_choice", text="Что такое Python?", option_1="Язык", option_2="Змея", option_3="Остров", correct_answer="1"),
            q_row(question_type="text", text="Новый", correct_answer="да"),
            q_row(question_type="text", text="Новый", correct_answer="да"),
        ]
        report = self.run_import(rows)
        self.assertEqual((report.created, len(report.skipped), report.error_rows), (1, 1, 1))
        self.assertEqual(self.test.questions.count(), 2)
        report = self.run_import(rows[:1], update_existing=True)
        self.assertEqual(report.updated, 1)
        self.assertEqual(self.test.questions.get(text="Что такое Python?").options.count(), 3)

    def test_round_trip_export_edit_import(self):
        question = self.add_choice()
        question.options.filter(text="Змея").update(image_url="https://example.com/snake.png")
        svc.save_question(self.test, QuestionData(QuestionType.TEXT, "Вывод?", correct_answers=["print"]))
        response = self.client.get(self.url("testing_questions_export", self.test.pk), {"format": "csv"})
        rows = read_csv(response)
        self.assertEqual([r["question"] for r in rows][0], str(question.pk))
        self.assertEqual((rows[0]["correct_answer"], rows[1]["correct_answer"], rows[1]["order"]), ("1", "print", "2"))
        rows[0]["text"] = "Что такое Python? (ред.)"
        rows[0]["points"] = "5"
        upload = csv_file([Q_HEADER, *[[r.get(c, "") for c in Q_HEADER] for r in rows]])
        report = import_export.import_questions(self.test, upload, update_existing=True)
        self.assertEqual((report.created, report.updated, report.error_rows), (0, 2, 0), report.errors)
        question.refresh_from_db()
        self.assertEqual((question.text, question.points), ("Что такое Python? (ред.)", 5))
        self.assertEqual(question.options.get(text="Змея").image_url, "https://example.com/snake.png")
        upload = csv_file([Q_HEADER, q_row(question="00000000-0000-0000-0000-000000000000", question_type="text", text="X", correct_answer="x")])
        self.assertIn("не найден", import_export.import_questions(self.test, upload, update_existing=True).errors[0].errors[0])

    def test_order_from_the_file(self):
        self.add_choice("Первый"), self.add_choice("Второй")
        self.run_import([q_row(question_type="text", text="Новый первый", correct_answer="x", order="1")])
        self.assertEqual(list(self.test.questions.order_by("order").values_list("text", flat=True)), ["Новый первый", "Первый", "Второй"])

    def test_views_and_template(self):
        page = self.client.get(self.url("testing_questions_import", self.test.pk))
        self.assertContains(page, "Тест: <strong data-okt-io-test>Python Basics</strong>", html=False)
        upload = csv_file([Q_HEADER, q_row(question_type="text", text="Вывод?", correct_answer="print")])
        response = self.client.post(self.url("testing_questions_import", self.test.pk), {"file": upload})
        self.assertEqual(response.context["report"].created, 1)
        self.assertContains(response, "Создано вопросов")
        workbook = load_workbook(io.BytesIO(self.client.get(self.url("testing_questions_import_template", self.test.pk)).content))
        self.assertEqual([c.value for c in workbook.active[1]][:4], ["test", "question", "question_type", "text"])
        examples = list(workbook["Пример"].iter_rows(min_row=2, values_only=True))
        self.assertEqual({row[2] for row in examples}, {"single_choice", "multiple_choice", "text", "code"})
        # The template's own examples import cleanly.
        example_file = xlsx_file([[c.value for c in workbook["Пример"][1]], *examples])
        report = import_export.import_questions(self.test, example_file)
        self.assertEqual((report.created, report.error_rows), (4, 0), report.errors)
        self.assertContains(self.client.get(self.url("testing_questions_export", self.test.pk)), "Экспорт вопросов")


class DuplicateAndListTests(Fixture):
    def test_duplicate_copies_questions_as_a_draft(self):
        question = self.add_choice()
        Test.objects.filter(pk=self.test.pk).update(status=TestStatus.ACTIVE)
        self.assertEqual(self.client.get(self.url("testing_test_duplicate", self.test.pk)).status_code, 405)
        response = self.client.post(self.url("testing_test_duplicate", self.test.pk))
        copy = Test.objects.get(title="Python Basics (копия)")
        self.assertRedirects(response, self.url("testing_test_change", copy.pk))
        self.assertEqual((copy.status, copy.subject, copy.level), (TestStatus.DRAFT, self.python, TestLevel.EASY))
        copied = copy.questions.get()
        self.assertNotEqual(copied.pk, question.pk)
        self.assertEqual([(o.text, o.is_correct) for o in copied.options.order_by("order")], [("Язык", True), ("Змея", False)])
        self.assertEqual(self.test.questions.get().options.count(), 2)
        self.client.post(self.url("testing_test_duplicate", self.test.pk))
        self.assertTrue(Test.objects.filter(title="Python Basics (копия 2)").exists())

    def test_list_toolbar_filters_cards_and_dialogs(self):
        Test.objects.create(title="HTML", level=TestLevel.HARD, image_url="https://example.com/html.png")
        response = self.client.get(self.url("testing_test_changelist"))
        for text in ("Создать тест", "Импорт", "Экспорт", "Все предметы", "Все статусы", "Любая сложность",
                     'id="okt-tests-import"', 'id="okt-tests-export"', 'id="okt-questions-import"',
                     "Управление вопросами", "Импорт вопросов", "Экспорт вопросов", "Дублировать", "В архив",
                     'src="https://example.com/html.png"', "okt-cover__placeholder", 'form="okt-export-form"'):
            self.assertContains(response, text)
        self.assertNotContains(response, "{#")  # a multi-line {# #} comment would leak into the page
        titles = lambda **p: [t.title for t in self.client.get(self.url("testing_test_changelist"), p).context["tests"]]
        self.assertEqual(titles(subject=str(self.python.pk)), ["Python Basics"])
        self.assertEqual(titles(subject="none"), ["HTML"])
        self.assertEqual(titles(level="hard"), ["HTML"])
        self.assertEqual(sorted(titles(subject="junk", level="junk")), ["HTML", "Python Basics"])
        filtered = self.client.get(self.url("testing_test_changelist"), {"level": "hard"})
        self.assertContains(filtered, '<input type="hidden" name="level" value="hard">', html=True)

    def test_test_page_has_questions_import_export(self):
        self.add_choice()
        response = self.client.get(self.url("testing_test_change", self.test.pk))
        for text in ("Назад к тестам", "Добавить вопрос", "Импорт вопросов", "Экспорт вопросов",
                     reverse("admin:testing_questions_import_template", args=[self.test.pk])):
            self.assertContains(response, text)
        self.assertNotContains(response, "{#")
        self.assertNotContains(self.client.get(self.url("testing_questions_import", self.test.pk)), "{#")

    def test_question_image_thumb(self):
        question = self.add_choice()
        Question.objects.filter(pk=question.pk).update(image_url="https://example.com/q.png")
        response = self.client.get(self.url("testing_test_change", self.test.pk))
        self.assertContains(response, 'class="okt-qthumb"')

    def test_teacher_staff_has_no_access(self):
        teacher = User.objects.create_user(
            username="t1", email="t1@okurmen.kg", password="x", role=User.Role.TEACHER, is_staff=True,
        )
        self.client.force_login(teacher)
        for url in (
            self.url("testing_test_export"), self.url("testing_test_import"), self.url("testing_test_import_template"),
            self.url("testing_questions_export", self.test.pk), self.url("testing_questions_import", self.test.pk),
        ):
            self.assertIn(self.client.get(url).status_code, (302, 403), url)
        self.assertIn(self.client.post(self.url("testing_test_duplicate", self.test.pk)).status_code, (302, 403))
